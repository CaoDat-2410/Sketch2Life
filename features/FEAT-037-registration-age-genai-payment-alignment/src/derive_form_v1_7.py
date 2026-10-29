from __future__ import annotations

import hashlib
from pathlib import Path
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[3]
FEATURE = ROOT / "features/FEAT-037-registration-age-genai-payment-alignment"
SOURCE = FEATURE / "artifacts/Phieu_FA26SE225_v1.5.docx"
OUTPUT = FEATURE / "artifacts/Phieu_FA26SE225_v1.7.docx"
EXPECTED_SOURCE_SHA256 = "47F23FD4486A4605208BB44A7B3696081F973DF74D4F1B8D56F7D35C34DF6956"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}

# Paragraph index -> (expected old prefix, replacement text or run-preserving pieces).
updates: dict[int, tuple[str, str | tuple[str, ...]]] = {
    13: ("3.1.1. English:", ("3.1.1. English: ", "Sketch2Life: Collaborative Creative Learning with AI for Children")),
    14: ("3.1.2. Vietnamese:", ("3.1.2. Vietnamese: ", "Sketch2Life: Học tập sáng tạo cộng tác với AI cho trẻ em")),
    18: ("Sketch2Life is an adult-supervised", "Sketch2Life is a collaborative creative-learning system for children ages 3 to 12. Children draw individually or together in an adult-led classroom session, explore age-adapted AI suggestions, share their work, and connect it to reviewed knowledge, video, and an off-screen group activity. Teachers manage class and group progress and make the final decisions about AI content and learning activities."),
    19: ("-\tA child's drawing", "-\tChildren use drawing to express ideas, explore a topic, and learn with classmates; artistic similarity to a model is not the learning goal."),
    20: ("-\tMontessori activities", "-\tThe product supports one-school classroom pilots for children from 36 through 155 completed months, with experience bands of ages 3–5, 6–8, and 9–12."),
    21: ("-\tChildren's drawings", "-\tPersonal and shared drawing needs clear group progress, contribution context, reliable saving, and recovery after connection loss."),
    22: ("-\tA screen-based result", "-\tAI can offer a sketch suggestion, but the Teacher reviews each suggestion before children can use it; children may hide or decline it."),
    23: ("-\tGenerated stories", "-\tKnowledge, video, and off-screen activities need age-appropriate content, review, safety checks, and Teacher oversight."),
    24: ("-\tDrawings, voice recordings", "-\tChildren's artwork and session records require school-mediated consent, purpose-limited access, retention, deletion, and audit controls."),
    27: ("Build Sketch2Life as an adult-supervised flow", "Build Sketch2Life around class sessions with managed child profiles, QR or session-code admission, groups, and individual or shared Android drawing. AI may help children explore ideas through reviewed sketch suggestions. The Teacher directs session stages, reviews each suggestion and video before child access, chooses activities for the class or a group, and guides reflection and off-screen work. Groups can progress independently while the Teacher controls shared class stages."),
    28: ("-\tDrawing Capture and Child Narration:", "-\tClass and Profile Management: A Teacher manages assigned classes, enrollments, and adult-confirmed child profiles; a child joins a session through an approved QR code or session code without a separate login account."),
    29: ("-\tMultimodal Understanding and Gate A:", "-\tSession and Group Management: A Teacher configures a class session, admits participants, assigns groups, and follows each group's progress without automatically advancing the whole class when one group is ready."),
    30: ("-\tMontessori Mapping and Gate B:", "-\tCollaborative Drawing: Children draw on personal or shared canvases with age-adapted tools. Shared-device turn selection records the selected contributor for new strokes and does not rewrite earlier contributions."),
    31: ("-\tReviewed Knowledge and Story Planning:", "-\tAI Vision and Sketch Assistance: The system analyzes a drawing in its class and group context, preserves uncertainty, and prepares a suggestion for Teacher review. Children can view, hide, or decline an approved suggestion while continuing to draw."),
    32: ("-\tPersonalized Drawing Exploration:", "-\tArtwork Gallery: A Teacher selects and orders artwork for sharing within the class session; children can preview and explain their own or permitted group work."),
    33: ("-\tIllustrated Story Video:", "-\tKnowledge and Video: A Teacher creates or selects reviewed knowledge and video, then approves the exact item before playback. The class sees the true generation state and waits while a video is being prepared. If permitted retries are exhausted, the Teacher chooses whether to retry, skip the video, or end the session."),
    34: ("-\tHands-On Activity Delivery:", "-\tOff-Screen Activities: A Teacher chooses or adapts a reviewed library activity for the class or a group, guides children through the physical activity, and records its session outcome."),
    35: ("-\tSafety, Fallback and Recovery:", "-\tReflection and Assessment: Children reflect on their work when the session allows it; Teachers record descriptive observations and progress over time without ranking children."),
    36: ("-\tObservation and History:", "-\tPrivacy and Recovery: The system records consent purpose and access outcomes, preserves artwork provenance, supports session and canvas recovery, and applies retention and deletion controls."),
    39: ("-\tChild, with an accompanying adult:", "-\tChild participant, using a Teacher-managed profile:"),
    40: ("•\tHave the original drawing captured", "•\tJoin an admitted class session on an Android tablet or phone using a QR code or session code; no separate child login account is required."),
    41: ("•\tExplore a source-derived PixiJS", "•\tDraw individually or collaboratively with age-adapted tools, take turns on a shared device, request AI help, and hide or decline a reviewed suggestion."),
    42: ("•\tReceive one approved, concrete", "•\tPreview and explain permitted artwork, view Teacher-approved knowledge or video, take part in an off-screen activity, and reflect on the session."),
    44: ("-\tParent / Owner Caregiver:", "-\tTeacher, within the assigned school and class scope:"),
    45: ("•\tCreate and manage a ChildProfile", "•\tManage assigned classes, enrollments, and child profiles; use adult-confirmed age context and never infer a child's age from artwork or media."),
    46: ("•\tReview Gate A understanding", "•\tSet up class sessions with topic, age range, drawing tools, groups, and presets; admit children who join by QR code or session code."),
    47: ("•\tReview and approve the activity", "•\tPause, resume, advance, or end a session; monitor group progress, provide help, and keep one group's completion from advancing the whole class automatically."),
    48: ("•\tMark an activity as completed", "•\tReview each AI sketch suggestion, edit, approve, or reject it before child access; respect a child's choice to hide or decline a suggestion."),
    49: ("•\tView interest history", "•\tSelect and present artwork within the class scope; review and approve the exact knowledge or video item before playback."),
    50: ("•\tAssign or revoke Guides", "•\tChoose or adapt an off-screen activity for the class or a group, guide the activity, record observations, and manage session records within authorized privacy procedures."),
    52: ("-\tMontessori Guide, within an active assignment:", "-\tTeacher with content review permissions, within the assigned scope:"),
    53: ("•\tManage the curriculum knowledge base", "•\tCreate or submit reviewed learning content and off-screen activities through the permitted review and publication process."),
    54: ("•\tReview and correct mappings", "•\tReview and correct learning content, age suitability, safety guidance, and activity instructions with version history."),
    55: ("•\tView assigned-child observation records", "•\tReview assigned class and group progress and keep direct observation separate from AI suggestions and child responses."),
    56: ("•\tCurate versioned story and activity", "•\tRecord descriptive learning observations and submit correction requests without treating unobserved behavior as a lack of ability."),
    58: ("-\tAI Processing and Media Engine:", "-\tAI and Media Services, under application policy:"),
    59: ("•\tValidate media, transcribe narration", "•\tAnalyze drawing snapshots with authorized class, group, topic, and age context; preserve uncertainty and the source version used."),
    60: ("•\tPreserve source conflicts and uncertainty", "•\tPrepare adaptive sketch suggestions as separate content; do not write over children's original strokes or bypass Teacher review."),
    61: ("•\tReturn all reviewed matches", "•\tSupport Teacher-reviewed knowledge and video generation or library selection, with accurate status and versioned review outcomes."),
    62: ("•\tRetrieve reviewed knowledge", "•\tKeep factual knowledge tied to reviewed sources and make clear when content is a creative suggestion rather than a factual claim."),
    63: ("•\tAnimate the original artwork", "•\tPreserve original artwork and track versions and provenance for approved derivatives and shared outputs."),
    64: ("•\tSupport quick and free-form story edits", "•\tApply applicable safety and consent policies before a child can access a suggestion or media item; do not approve content on the Teacher's behalf."),
    65: ("•\tApply content, activity, audio, video", "•\tReport typed generation, review, and recovery states; do not mark failed or unreviewed content as ready or automatically skip a class decision."),
    67: ("-\tSystem Admin:", "-\tSuper Admin:"),
    68: ("•\tManage adult accounts", "•\tManage adult accounts, roles, school/class scopes, and Teacher assignments under authorized permissions."),
    69: ("•\tConfigure approved model profiles", "•\tManage reviewed learning content, AI policies, and operating settings through versioned and auditable controls."),
    70: ("•\tManage consent, retention", "•\tOversee consent records, data requests, retention, deletion, and audit processes according to their separate purposes and policies."),
    71: ("•\tMonitor processing jobs", "•\tMonitor service health and processing jobs; access child-related records only for an authorized purpose and within the assigned scope."),
    74: ("-\tOff-Screen Orientation:", "-\tTeacher-Led Pacing: Teachers control session stages, can pause or end a session, and guide children from digital drawing toward a physical class or group activity."),
    75: ("-\tOwnership and Provenance", "-\tContribution History: Preserve original artwork and accepted contributions. On shared devices, turn selection records the contributor chosen for new strokes and does not change earlier attribution."),
    76: ("-\tSafe Handling of Misrecognition:", "-\tChild Agency and AI Review: AI suggestions remain separate from children's drawings. A Teacher reviews each suggestion before use, and children may hide or decline it."),
    77: ("-\tContent and Activity Safety:", "-\tReviewed Content: Knowledge, video, and physical activity guidance must pass the applicable content and safety review before use with children."),
    78: ("-\tPedagogical Correctness:", "-\tAge-Appropriate Interaction: The product supports ages 3–12 across the experience bands 3–5, 6–8, and 9–12, using Teacher-managed age context."),
    79: ("-\tPrivacy of Children's Data:", "-\tConsent and Purpose: The school collects legal-representative consent; authorized Teachers or Super Admins record its evidence and permitted purposes before policy-limited processing."),
    80: ("-\tResponsiveness and Recovery:", "-\tRetention and Privacy: Session and artwork data have a default retention period of 90 days after the session ends; other data classes follow separate policies."),
    81: ("-\tOffline Tolerance:", "-\tRecovery and Capacity: Sessions and canvases support safe reconnect and recovery. The one-school pilot targets one class of up to 40 children; this target is not a measured capacity result."),
    85: ("The project applies Montessori pedagogy", "The project applies child-centered collaborative learning, age-adapted drawing experiences, Montessori-informed physical activities, teacher-led AI review, classroom group management, and privacy-aware recovery. The product does not claim to be a pure Montessori program."),
    86: ("-\tMontessori Method", ("-\t", "Montessori-Informed Learning:", " Use reviewed learning activities to connect children's questions and classroom topics to hands-on exploration.")),
    87: ("-\tMontessori Material Sequence:", ("-\t", "Collaborative Creative Learning:", " Children express ideas through individual and shared drawing, with group participation and process valued over copying a model.")),
    88: ("-\tReality Orientation", ("-\t", "Age-Adapted Interaction:", " Drawing tools and classroom experiences should suit the supported ages 3–5, 6–8, and 9–12.")),
    89: ("-\tStages of Children's Artistic Development:", ("-\t", "Children's Drawing and Expression:", " Interpret children's artwork as a form of expression; do not judge learning by resemblance to an expected picture.")),
    90: ("-\tSketch and Abstract Image Understanding:", ("-\t", "Sketch Understanding and Uncertainty:", " Drawing analysis must account for children's visual expression and communicate uncertainty for Teacher review.")),
    91: ("-\tMultimodal Grounding", ("-\t", "Human Review and Child Agency:", " Teacher review keeps AI suggestions within classroom policy, while children retain the choice to hide or decline them.")),
    92: ("-\tAutomatic Rigging and Animation", ("-\t", "Classroom Collaboration:", " Group progress, shared drawing, and turn selection require clear contribution history across devices and reconnects.")),
    93: ("-\tTopic-and-Age-Grounded Activity Discovery:", ("-\t", "Learning Assessment:", " Descriptive Teacher observations and child reflection capture process over time without ranking children.")),
    94: ("-\tSafety Alignment", ("-\t", "Child Data Protection:", " School-mediated consent, purpose limits, audit, provenance, retention, and deletion govern classroom records and media.")),
    95: ("-\tJoint Media Engagement:", ("-\t", "Teacher-Facilitated Media Use:", " Teachers choose and review classroom media, then guide discussion and shared exploration.")),
    98: ("These concepts are applied by building and validating the target workflow", "These concepts are applied by designing and evaluating the classroom workflow with representative age bands, teacher-reviewed content, and recovery scenarios. Any evaluation involving children follows school consent and approved research and data-handling procedures."),
    99: ("-\tBuild a reviewed Montessori curriculum", "-\tModel class, profile, enrollment, session, group, and participant-admission requirements for a one-school pilot."),
    100: ("-\tCollect and annotate research data", "-\tDesign personal and shared drawing experiences for Android tablets and phones, including age-adapted tools, group progress, and shared-device turn selection."),
    101: ("-\tBuild the multimodal understanding engine", "-\tEvaluate AI Vision and adaptive sketch suggestions with uncertainty, Teacher review, and child hide or decline choices."),
    102: ("-\tImplement PixiJS exploration", "-\tDesign class gallery, reviewed knowledge and video, off-screen activities, reflection, and Teacher observation workflows."),
    103: ("-\tImplement age-aware reviewed knowledge", "-\tSpecify role and purpose controls for profiles, consent, artwork, session records, portfolio, audit, retention, and deletion."),
    104: ("-\tImplement complete topic-and-age discovery", "-\tEvaluate save, reconnect, pause, group progress, and recovery behavior with synthetic scenarios before any classroom data collection."),
    105: ("-\tEncode and validate activity, content", "-\tReview age suitability, safety, Teacher decision points, accessibility, and contribution attribution across the supported experience bands."),
    106: ("-\tBuild the adult-supervised session flow", "-\tBuild the Child Android application, Teacher Desktop, Super Admin Console, and FastAPI services as separately planned product workstreams."),
    107: ("-\tValidate mappings and media quality", "-\tEvaluate classroom usability and service reliability without claiming learning gains or pilot capacity that have not been measured."),
    110: ("-\tParent and Child Android Mobile Application", "-\tChild Android Application: React Native experience for tablets and phones, with managed-profile session admission, individual and collaborative drawing, AI assistance, and session recovery."),
    111: ("-\tResponsive Parent Web", "-\tTeacher Desktop and Super Admin Console for class and group setup, session control, content review, learning records, consent, and operations."),
    112: ("-\tMontessori Curriculum Knowledge Base:", "-\tClassroom and Group Services: Managed child profiles, class enrollment, session admission, group membership, independent progress, and teacher assignments."),
    113: ("-\tMultimodal Drawing Understanding Engine:", "-\tCollaborative Drawing Canvas: Personal or shared drawing, age-adapted tools, contribution context, gallery presentation, save, and recovery."),
    114: ("-\tPersonalized Drawing Exploration Engine:", "-\tAI Vision and Adaptive Sketch Service: Contextual analysis, uncertainty, separate suggestions, Teacher review, and child response."),
    115: ("-\tAge-Aware Story Experience Service:", "-\tKnowledge and Video Service: Reviewed learning content, video creation or library selection, Teacher approval, playback, and failure disposition."),
    116: ("-\tTopic-and-Age Activity Discovery:", "-\tOff-Screen Activity and Learning Records: Class or group activities, reflection, Teacher observations, and portfolio records."),
    117: ("-\tIllustrated Story Video and Narration Pipeline:", "-\tPrivacy and Recovery Services: Purpose-scoped consent records, audit, retention, data-request handling, and session/canvas recovery."),
    118: ("-\tResearch and Evaluation Outputs:", "-\tEvaluation Materials: Age-band interaction scenarios, collaborative drawing and AI-review evaluation, classroom recovery cases, and documented study limitations."),
    121: ("The work is organised into four coordinated Work Packages", "The work is organised into four product work packages. They describe related deliverables; assignment of people and sprint tasks is managed separately."),
    122: ("-\tWP1 - Project Management", "-\tWP1 - Product, Classroom Learning and Evaluation: Requirements, age-appropriate interaction, learning activity review, safety, and evaluation planning."),
    123: ("-\tWP2 - Backend, Identity", "-\tWP2 - Backend and Classroom Services: Adult identity, school and class scope, managed profiles, consent, sessions, groups, permissions, persistence, and recovery."),
    124: ("-\tWP3 - Understanding, Story and Media:", "-\tWP3 - AI and Learning Media: Drawing understanding, adaptive sketch suggestions, reviewed knowledge and video, safety controls, and provenance."),
    125: ("-\tWP4 - Mobile, Parent Web", "-\tWP4 - Product Surfaces and Integration: Child Android, Teacher Desktop, Super Admin workflows, collaborative drawing, session stages, and classroom usability."),
    129: ("The study asks whether combining a child's drawing", "The study examines how children ages 3–12 use individual and collaborative drawing in a classroom session, and whether reviewed AI sketch suggestions support exploration while preserving children's control of their work. It also examines how Teacher-led knowledge, video, and off-screen activities fit into group progress and session recovery."),
    132: ("-\tConstruct a reviewed Montessori curriculum", "-\tDescribe and evaluate class, group, and session workflows for children joining with managed profiles and QR or session codes."),
    133: ("-\tMeasure subject identification accuracy", "-\tEvaluate age-adapted drawing tools and collaborative canvas usability across the experience bands 3–5, 6–8, and 9–12."),
    134: ("-\tHave trained Montessori guides rate", "-\tAssess AI sketch suggestions for usefulness and safety through Teacher review, and document child decisions to view, hide, or decline suggestions."),
    135: ("-\tMeasure whether returned activities", "-\tEvaluate group progress, artwork sharing, approved knowledge or video, off-screen activity handoff, and descriptive learning records."),
    136: ("-\tMeasure off-screen completion", "-\tEvaluate session and canvas saving, shared-device contribution context, and reconnect recovery using representative synthetic scenarios."),
    137: ("-\tMeasure adult corrections", "-\tDocument Teacher review decisions and children's choices to view, hide, or decline AI sketch suggestions without treating those choices as measures of ability."),
    140: ("The mixed-methods study covers children under 9", "The evaluation combines requirement-based classroom scenarios with age-band usability and content review. Scenarios cover admission, personal and shared drawing, group progress, Teacher-reviewed suggestions, artwork sharing, approved video, off-screen activities, reflection, and recovery. Any study that involves children uses school-mediated consent and approved research and data-handling procedures. Before recruiting children, a study protocol will define participant count, duration, measures, and analysis."),
    143: ("The study tests whether an adult-supervised digital experience", "The project contributes an evaluated design for collaborative creative learning, age-adapted drawing, Teacher-controlled AI assistance, and off-screen classroom activity for children ages 3–12. Findings will describe usability, safety, and reliability limits; learning outcomes are not claimed without supporting evidence."),
    146: ("-\tMontessori pedagogical literature", ("-\t", "Research on child-centered and collaborative learning, classroom participation, and teacher-guided activity.")),
    147: ("-\tLowenfeld's stages", ("-\t", "Research on children's drawing, visual expression, and age-appropriate interpretation from early childhood through pre-adolescence.")),
    148: ("-\tResearch on sketch and abstract", ("-\t", "Research on sketch understanding, AI assistance, uncertainty, and human review for child-directed systems.")),
    149: ("-\tPublished work on automatically", ("-\t", "Research on shared digital drawing, collaborative classroom systems, contribution attribution, and recovery across devices.")),
    150: ("-\tMultimodal image-and-text grounding", ("-\t", "Research on age-adapted interaction, accessibility, and children's agency when using AI suggestions.")),
    151: ("-\tResearch on educational recommender", ("-\t", "Research on classroom activity selection, reviewed educational knowledge, and teacher-directed learning.")),
    152: ("-\tWork on safety filtering", ("-\t", "Research on privacy, consent, safety review, and responsible AI use in child-facing educational settings.")),
    153: ("-\tResearch on young children's screen media", "-\tResearch on collaborative learning, teacher-facilitated media use, and transitions from digital activity to hands-on classroom work."),
    155: ("-\tThe Montessori constraint", "-\tDrawing supports children's expression and collaboration; the classroom experience also includes reviewed learning content, teacher-led reflection, and a practical off-screen activity."),
    156: ("-\tRecognition must fail gracefully", "-\tTeachers retain the final decision about AI suggestions, video, and activities. Children can continue drawing without an AI suggestion and may hide or decline one that has been approved for their use."),
    157: ("-\tCapstone-trial plans for adults:", "-\tThe initial pilot targets one school and one class of up to 40 children ages 36–155 completed months. The school collects legal-representative consent, and authorized Teachers or Super Admins record evidence and permitted purposes. Session and artwork data have a default retention period of 90 days after session end; other data classes follow separate policies."),
}

source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest().upper()
if source_hash != EXPECTED_SOURCE_SHA256:
    raise SystemExit(f"Unexpected v1.5 source hash: {source_hash}")
source_doc = Document(SOURCE)
if len(source_doc.paragraphs) != 159 or len(source_doc.tables) != 3:
    raise SystemExit("Unexpected source form structure")
source_paragraphs = [p.text for p in source_doc.paragraphs]
source_tables = [[[(cell.text) for cell in row.cells] for row in table.rows] for table in source_doc.tables]
if set(updates) - set(range(len(source_paragraphs))):
    raise SystemExit("Replacement index is outside the source paragraph list")

for index, (expected_prefix, replacement) in updates.items():
    old_text = source_paragraphs[index]
    if not old_text.startswith(expected_prefix):
        raise SystemExit(f"Source paragraph {index} changed unexpectedly: {old_text[:100]!r}")
    paragraph = source_doc.paragraphs[index]
    if isinstance(replacement, tuple):
        if len(paragraph.runs) != len(replacement):
            raise SystemExit(f"Paragraph {index} has {len(paragraph.runs)} runs; expected {len(replacement)}")
        pieces = replacement
    else:
        if not paragraph.runs:
            raise SystemExit(f"Paragraph {index} has no text run")
        pieces = (replacement,) + tuple("" for _ in paragraph.runs[1:])
    for run, piece in zip(paragraph.runs, pieces, strict=True):
        run.text = piece

# Use python-docx to preserve run-level tabs and formatting, then copy only its
# updated document.xml into the source package so every other package part stays exact.
with tempfile.NamedTemporaryFile(prefix="fa26-v1-7-", suffix=".docx", dir=FEATURE, delete=False) as handle:
    staging_path = Path(handle.name)
try:
    source_doc.save(staging_path)
    with ZipFile(staging_path, "r") as stage:
        updated_xml = stage.read("word/document.xml")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(SOURCE, "r") as zin, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            payload = updated_xml if info.filename == "word/document.xml" else zin.read(info.filename)
            zout.writestr(info, payload)
finally:
    staging_path.unlink(missing_ok=True)

# Verify package structure, table content, required current scope, and obsolete claims.
result = Document(OUTPUT)
result_paragraphs = [p.text for p in result.paragraphs]
result_tables = [[[(cell.text) for cell in row.cells] for row in table.rows] for table in result.tables]
if len(result.paragraphs) != 159 or len(result.tables) != 3:
    raise SystemExit("Output form structure changed")
if source_tables != result_tables:
    raise SystemExit("Output changed original table content")
for index, (_, replacement) in updates.items():
    expected_text = "".join(replacement) if isinstance(replacement, tuple) else replacement
    if result_paragraphs[index] != expected_text:
        raise SystemExit(f"Paragraph {index} did not round-trip as expected")
text = "\n".join(result_paragraphs)
required = ["36 through 155 completed months", "up to 40 children", "90 days after the session ends", "Super Admin", "collaborative", "school collects legal-representative consent"]
for phrase in required:
    if phrase.casefold() not in text.casefold():
        raise SystemExit(f"Required current-scope phrase is missing: {phrase}")
obsolete = ["under 9 years", "0–107", "Parent Web", "Parent /", "Owner Caregiver", "parent or guide", "99,000 VND", "499,000 VND", "monthly credits", "40-60 second", "40–60 second", "10 minutes", "Gate A", "Gate B", "B33", "household study", "screen time"]
for phrase in obsolete:
    if phrase.casefold() in text.casefold():
        raise SystemExit(f"Superseded or internal phrase remains: {phrase}")

with ZipFile(SOURCE, "r") as before, ZipFile(OUTPUT, "r") as after:
    if before.namelist() != after.namelist():
        raise SystemExit("DOCX package members changed")
    changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
    if changed != ["word/document.xml"]:
        raise SystemExit(f"Unexpected package-part changes: {changed}")

print(f"Source: {SOURCE}")
print(f"Source SHA-256: {source_hash}")
print(f"Output: {OUTPUT}")
print(f"Output SHA-256: {hashlib.sha256(OUTPUT.read_bytes()).hexdigest().upper()}")
print(f"Updated paragraphs: {len(updates)}; paragraphs retained: {len(result.paragraphs)}; tables retained: {len(result.tables)}")
print("DOCX package check: only word/document.xml changed; original tables and all other package parts are unchanged")
