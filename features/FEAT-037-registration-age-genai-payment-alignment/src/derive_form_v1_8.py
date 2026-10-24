from __future__ import annotations

import hashlib
from pathlib import Path
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document

ROOT = Path(__file__).resolve().parents[3]
FEATURE = ROOT / "features/FEAT-037-registration-age-genai-payment-alignment"
SOURCE = FEATURE / "artifacts/Phieu_FA26SE225_v1.2.docx"
CONTENT_REFERENCE = FEATURE / "artifacts/Phieu_FA26SE225_v1.7.docx"
OUTPUT = FEATURE / "artifacts/Phieu_FA26SE225_v1.8.docx"
EXPECTED_SOURCE_SHA256 = "C19D0AC8154BECA2F6CB850C1272A4577048DD0D81A5FF3273C01861B9260959"
EXPECTED_CONTENT_REFERENCE_SHA256 = "003AA55B24BB612C8E8FA27CE241D96C3CAA8FCD09925A3A2E1C8336E41F75E3"
EXPECTED_TABLE_SHAPES = [(2, 5), (5, 6), (1, 2)]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def table_snapshot(doc):
    return [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in doc.tables
    ]


def apply_text_preserving_runs(paragraph, target: str) -> None:
    runs = paragraph.runs
    if not runs:
        raise ValueError("Changed paragraph has no runs")

    # Preserve the two title runs, including the source's bold italic title style.
    if target.startswith("3.1.1. English: ") or target.startswith("3.1.2. Vietnamese: "):
        label = runs[0].text
        if not target.startswith(label) or len(runs) < 2:
            raise ValueError(f"Unexpected title run layout: {paragraph.text!r}")
        runs[0].text = label
        runs[1].text = target[len(label):]
        for run in runs[2:]:
            run.text = ""
        return

    # Preserve source list markers, tab stops, and the heading/body run split.
    for marker in ("-\t", "•\t", "–\t"):
        if paragraph.text.startswith(marker) and target.startswith(marker):
            body = target[len(marker):]
            if len(runs) >= 4 and runs[0].text == marker[0] and runs[1].text == "\t" and runs[2].text.endswith(":"):
                if ": " not in body:
                    raise ValueError(f"Expected a heading and body in target: {target!r}")
                heading, description = body.split(": ", 1)
                runs[0].text = marker[0]
                runs[1].text = "\t"
                runs[2].text = heading + ":"
                runs[3].text = " " + description
                for run in runs[4:]:
                    run.text = ""
            elif len(runs) >= 2 and runs[0].text == marker[0]:
                runs[0].text = marker[0]
                runs[1].text = "\t" + body
                for run in runs[2:]:
                    run.text = ""
            elif len(runs) == 1:
                runs[0].text = target
            else:
                raise ValueError(f"Unexpected list run layout at {paragraph.text[:80]!r}")
            return

    # For ordinary prose, all source runs must have the same direct character formatting.
    signatures = []
    for run in runs:
        font = run.font
        signatures.append((
            run.bold, run.italic, run.underline,
            font.name, font.size, font.color.rgb if font.color and font.color.type else None,
            run.style.name if run.style else None,
        ))
    if len(set(signatures)) > 1:
        raise ValueError(f"Unexpected mixed formatting in paragraph: {paragraph.text[:120]!r}")
    runs[0].text = target
    for run in runs[1:]:
        run.text = ""


source_hash = sha256(SOURCE)
reference_hash = sha256(CONTENT_REFERENCE)
if source_hash != EXPECTED_SOURCE_SHA256:
    raise SystemExit(f"Selected v1.2 source hash changed: {source_hash}")
if reference_hash != EXPECTED_CONTENT_REFERENCE_SHA256:
    raise SystemExit(f"Reviewed content reference hash changed: {reference_hash}")
if OUTPUT.exists():
    raise SystemExit(f"Refusing to replace existing output: {OUTPUT}")

source_doc = Document(SOURCE)
content_doc = Document(CONTENT_REFERENCE)
if len(source_doc.paragraphs) != 159 or len(content_doc.paragraphs) != 159:
    raise SystemExit("Expected 159 paragraphs in both source and content reference")
source_tables = table_snapshot(source_doc)
if [(len(table.rows), len(table.columns)) for table in source_doc.tables] != EXPECTED_TABLE_SHAPES:
    raise SystemExit("Selected v1.2 table shapes changed")

changed = []
for index, (source_paragraph, content_paragraph) in enumerate(zip(source_doc.paragraphs, content_doc.paragraphs, strict=True)):
    old_text = source_paragraph.text
    desired_text = content_paragraph.text
    # Preserve v1.2's original whitespace-only tab in its empty spacer paragraph.
    if old_text.strip() == desired_text.strip():
        continue
    if old_text != desired_text:
        apply_text_preserving_runs(source_paragraph, desired_text)
        if source_paragraph.text != desired_text:
            raise SystemExit(f"Paragraph {index} did not match target text after run-preserving edit")
        changed.append(index)

if not changed:
    raise SystemExit("No paragraphs needed updating")

with tempfile.NamedTemporaryFile(prefix="fa26-v1-8-", suffix=".docx", dir=FEATURE, delete=False) as handle:
    staging_path = Path(handle.name)
try:
    source_doc.save(staging_path)
    with ZipFile(staging_path, "r") as stage:
        updated_xml = stage.read("word/document.xml")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(SOURCE, "r") as original, ZipFile(OUTPUT, "w", ZIP_DEFLATED) as derived:
        for info in original.infolist():
            payload = updated_xml if info.filename == "word/document.xml" else original.read(info.filename)
            derived.writestr(info, payload)
finally:
    staging_path.unlink(missing_ok=True)

result = Document(OUTPUT)
if len(result.paragraphs) != 159 or len(result.tables) != 3:
    raise SystemExit("v1.8 structure differs from selected v1.2 source")
if table_snapshot(result) != source_tables:
    raise SystemExit("v1.2 table contents or row/cell shape changed")
for index, (_, desired) in enumerate(zip(source_doc.paragraphs, content_doc.paragraphs, strict=True)):
    if index in changed and result.paragraphs[index].text != desired.text:
        raise SystemExit(f"Final paragraph {index} failed text comparison")

text = "\n".join(paragraph.text for paragraph in result.paragraphs)
required = [
    "36 through 155 completed months", "up to 40 children",
    "90 days after the session ends", "Super Admin", "collaborative",
    "school collects legal-representative consent",
]
for phrase in required:
    if phrase.casefold() not in text.casefold():
        raise SystemExit(f"Required current-scope phrase missing: {phrase}")
obsolete = [
    "under 9 years", "0–107", "Parent Web", "Parent /", "Owner Caregiver",
    "parent or guide", "99,000 VND", "499,000 VND", "monthly credits",
    "40-60 second", "40–60 second", "10 minutes", "Gate A", "Gate B",
    "B33", "household study", "screen time",
]
for phrase in obsolete:
    if phrase.casefold() in text.casefold():
        raise SystemExit(f"Superseded/internal phrase remains: {phrase}")

with ZipFile(SOURCE, "r") as original, ZipFile(OUTPUT, "r") as derived:
    if original.namelist() != derived.namelist():
        raise SystemExit("DOCX package member list changed")
    different = [name for name in original.namelist() if original.read(name) != derived.read(name)]
    if different != ["word/document.xml"]:
        raise SystemExit(f"Unexpected package changes: {different}")

print(f"Source v1.2 SHA-256: {source_hash}")
print(f"Reviewed text reference v1.7 SHA-256: {reference_hash}")
print(f"Output v1.8 SHA-256: {sha256(OUTPUT)}")
print(f"Updated paragraphs: {len(changed)}; source paragraphs/tables preserved: {len(result.paragraphs)}/{len(result.tables)}")
print(f"Table shapes preserved from v1.2: {EXPECTED_TABLE_SHAPES}")
print("DOCX package check: only word/document.xml changed")

