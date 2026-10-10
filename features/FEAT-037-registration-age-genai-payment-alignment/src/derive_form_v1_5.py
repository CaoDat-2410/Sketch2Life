from __future__ import annotations

import argparse
import hashlib
import tempfile
import zipfile
from pathlib import Path

from docx import Document


SOURCE_SHA256 = "21739477916A1FC2E03CC0E22B6F8DD383A03F6831F2EA4EEB1593747CABB548"

REPLACEMENTS = {
    "Sketch2Life is an adult-supervised system for children aged 0-12 that turns a child's drawing and narration into a personalized Montessori learning experience. The drawing is treated as an interest signal, not as an end product. The digital experience must remain a short, safe bridge to a concrete activity that the child performs away from the screen with a parent, caregiver or Montessori guide.": (
        "Sketch2Life is an adult-supervised system for children under 9 years (0–107 completed months) "
        "that turns a child's drawing and narration into a personalized Montessori learning experience. "
        "The drawing is an interest signal, not an end product. The short, safe digital experience bridges "
        "to an off-screen activity with a parent, caregiver or Montessori guide. Ages 9–12 remain in the "
        "catalog but outside supported product sessions."
    ),
    "-\tMontessori activities follow age, readiness and prerequisite sequences, so topical similarity alone is not enough for recommendation.": (
        "-\tMontessori activities are discovered by confirmed topic and exact product age. Readiness, "
        "prior work, prerequisite sequence and material availability may inform adult context but do not "
        "filter the otherwise eligible activity set."
    ),
    "•\tCreate and manage a ChildProfile with age/readiness context, consent and retention settings.": (
        "•\tCreate and manage a ChildProfile with adult-supplied age/readiness context, consent and retention settings."
    ),
    "•\tRetrieve reviewed knowledge with claim/source references and create age/readiness-constrained story drafts.": (
        "•\tRetrieve reviewed knowledge with source citations and create age/readiness-constrained story "
        "drafts from adult/profile-supplied age; never infer age from a drawing or narration."
    ),
    "Build Sketch2Life as an adult-supervised flow in which a drawing and the child's own description are transformed into a Montessori activity. The system validates the media, combines image and narration, opens an adult understanding review, filters activities by age, readiness, prerequisite, material and safety rules, and asks an adult to approve the activity and learning objective. It then provides a source-derived PixiJS exploration and an age-aware illustrated story short, before closing the digital experience and handing the child off to the approved physical activity.": (
        "Build Sketch2Life as an adult-supervised flow that transforms a child's drawing and description "
        "into a Montessori activity. It validates media, combines image and narration, and asks an adult "
        "to confirm understanding. It returns the complete reviewed activity set matching the confirmed "
        "topic and exact product age, applying active-catalog, safety/policy and required-supervision "
        "constraints. Adult-confirmed interests may order or explain matches; readiness, prior work, "
        "sequence and material availability do not exclude eligible matches. An adult approves the "
        "activity and objective before source-derived PixiJS exploration and an age-aware illustrated story; "
        "each session hands off to the approved physical activity."
    ),
    "-\tMontessori Mapping and Gate B: Maps confirmed meaning to curriculum areas, materials, activities and learning objectives, then applies hard age, readiness, prerequisite, supervision and material rules before ranking.": (
        "-\tMontessori Mapping and Gate B: Maps confirmed meaning to curriculum activities and objectives. "
        "Discovery returns the complete reviewed topic-and-age set, with active-catalog, safety/policy and "
        "supervision filters. Adult-confirmed interests order or explain matches; readiness, prior work, "
        "sequence and materials do not exclude matches. An adult approves the selected activity/objective."
    ),
    "-\tMontessori activities follow age, readiness and prerequisite sequences, so topical similarity alone is not enough for recommendation.": (
        "-\tMontessori activities are discovered by confirmed topic and exact product age. Readiness, "
        "prior work, prerequisite sequence and material availability may inform adult context but do not "
        "filter the otherwise eligible activity set."
    ),
    "•\tMap confirmed meaning to catalog activities/objectives and apply hard filters before ranking.": (
        "•\tReturn all reviewed matches for the confirmed topic and exact product age; apply active-catalog, "
        "safety/policy and required-supervision filters, then order by adult-confirmed interests."
    ),
    "-\tPedagogical Correctness: Recommendations must respect the Montessori material sequence, readiness and hard safety rules, and remain reviewable and overridable by an authorized guide.": (
        "-\tPedagogical Correctness: Return the complete reviewed topic-and-age set, applying catalog, "
        "safety/policy and supervision constraints. Guides may review recommendations; readiness and "
        "material sequence inform context but do not filter matches."
    ),
    "The project applies Montessori pedagogy, developmental psychology, sketch and multimodal understanding, constrained recommendation, child-directed safety, provenance and human-in-the-loop review.": (
        "The project applies Montessori pedagogy, developmental psychology, sketch and multimodal "
        "understanding, topic-and-age-grounded activity discovery with interest-based ordering, "
        "child-directed safety, provenance and human-in-the-loop review."
    ),
    " Recommenders in which hard constraints — ordering, readiness, safety — filter the candidate set before ranking.": (
        " Return reviewed matches for the confirmed topic and exact age, filtering by active catalog, "
        "safety/policy and supervision, then ordering by adult-confirmed interests."
    ),
    "-\tImplement the constrained recommender over the prerequisite graph, incorporating completed-activity history, readiness and safety rules.": (
        "-\tImplement complete topic-and-age discovery with catalog, safety and supervision filters and "
        "interest-based ordering."
    ),
    "-\tPrerequisite-Constrained Activity Recommender: Candidate filtering by age, readiness, prior work, sequence, materials and safety before ranking.": (
        "-\tTopic-and-Age Activity Discovery: Return reviewed topic-and-age matches with catalog, safety "
        "and supervision filtering and interest-based ordering."
    ),
    "The work is organised into four coordinated Work Packages. They describe project traceability and collaboration; exact sprint allocation, implementation approval and contract ownership follow the approved project plan.": (
        "The work is organised into four coordinated Work Packages that describe the project scope and "
        "team collaboration."
    ),
    "Systems that turn children's drawings into generated stories are often evaluated mainly by engagement, while Sketch2Life is intended to use digital media as a short bridge to hands-on Montessori work. The study therefore asks whether combining the drawing with the child's own description and applying age, readiness, prior-work and prerequisite constraints produces recommendations that trained Montessori guides judge as appropriate, compared with image-only and unconstrained baselines. It also asks whether the complete digital flow, including original-art exploration and story/video presentation, leads to off-screen activity completion without adding unnecessary screen time.": (
        "The study asks whether combining a child's drawing and description improves understanding, and "
        "whether complete activity discovery for a confirmed topic and exact product age, filtered for "
        "safety/supervision and ordered by adult-confirmed interests, helps Montessori guides select "
        "appropriate work. It compares this approach with image-only and similarity baselines and examines "
        "whether the full digital experience leads to off-screen activity without unnecessary screen time."
    ),
    "-\tMeasure subject identification accuracy under three input conditions - image only, description only and both combined - reported by age band and curriculum area.": (
        "-\tMeasure subject identification accuracy under image-only, description-only and combined input "
        "conditions, reported across the product age bands (0–3, 3–6 and 6–9) and curriculum areas."
    ),
    "-\tHave trained Montessori guides blind-rate the developmental appropriateness of constrained recommendations against an unconstrained similarity baseline and report inter-rater agreement.": (
        "-\tHave trained Montessori guides rate the complete confirmed-topic/exact-age activity set and its "
        "interest-based ordering against an unconstrained-similarity baseline; report inter-rater agreement."
    ),
    "-\tMeasure the rate at which constrained and unconstrained recommenders propose activities that violate the material sequence or safety rules.": (
        "-\tMeasure whether returned activities meet active-catalog, safety/policy and required-supervision "
        "constraints; report Guide ratings separately."
    ),
    "-\tMeasure off-screen activity completion and per-session digital screen time in a household study comparing the full flow with delivery of the same activity without animation; exact design and sample remain subject to approval.": (
        "-\tMeasure off-screen completion and per-session screen time in a household study comparing the "
        "full flow with delivery of the same activity without animation; define design and sample in its protocol."
    ),
    "A mixed-methods study is bounded to early-childhood use cases and spontaneously produced drawings, while the product target supports ages 0-12 through catalog age bands and readiness rules. Understanding is evaluated on held-out data under the three input conditions. Recommendation quality is rated by trained Montessori guides, sequence and safety violations are counted, and a household study may measure session count, screen time, adult-recorded completion and interviews. Consent, deletion, data minimization, sample size, assignment design, duration and acceptance thresholds must be approved before collection; the study does not claim long-term learning improvement.": (
        "The mixed-methods study covers children under 9 (0–107 completed months) and spontaneous drawings. "
        "Product bands are 0–35, 36–71 and 72–107 months; 9–12 catalog records are outside product and "
        "research scope. Understanding uses held-out data; trained Guides rate complete topic-and-age "
        "activity sets with catalog, safety and supervision checks. A household study may measure sessions, "
        "screen time, completion and interviews. Its protocol defines consent, deletion, data minimization, "
        "sample, design, duration and thresholds before collection. No long-term learning benefit is claimed."
    ),
    "The expected contribution is evidence about whether a carefully constrained, adult-supervised digital artifact can act as a bridge to hands-on work rather than displacing it. The study also evaluates multimodal understanding, prerequisite-constrained recommendation and adult correction. It must report limitations and must not treat engagement, short-term completion or a small household trial as proof of long-term educational benefit.": (
        "The study tests whether an adult-supervised digital experience bridges to hands-on activity. It "
        "evaluates multimodal understanding, topic-and-age discovery and adult correction, reports limits, "
        "and makes no claim of long-term educational benefit."
    ),
    "-\tResearch and Evaluation Outputs: Annotated Children's Drawing Dataset and Evaluation Report, released only under approved consent, protocol, privacy and scope decisions.": (
        "-\tResearch and Evaluation Outputs: Annotated Children's Drawing Dataset and Evaluation Report, "
        "governed by participant consent, research protocol, privacy and scope safeguards."
    ),
    "-\tThe target product scope includes the knowledge base, multimodal understanding, constrained recommender, activity delivery, Guide Console, Parent Web, PixiJS exploration and the age-aware illustrated story-video flow. Dataset release, production provider selection, performance thresholds and household research remain subject to consent, evidence and separate approval gates.": (
        "-\tCapstone-trial plans for adults: Free, 0 VND (10 monthly credits; 1 child); Gia đình, "
        "99,000 VND/month (30 pooled credits; up to 3 children); Lớp học, 499,000 VND/class/month "
        "(120 pooled credits; one Guide; up to 25 children). Top-ups: 10 credits/49,000 VND; "
        "30/129,000 VND; 60/239,000 VND. One credit covers a full adult-approved drawing/story experience "
        "through off-screen handoff: reserve at session start, charge on successful handoff, release on "
        "failure/cancellation; same-session retries are idempotent. Paid plans renew manually; backend "
        "payment verification precedes entitlement or credit grant."
    ),
}


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def set_paragraph_text(paragraph, text: str) -> None:
    if len(paragraph.runs) != 1:
        raise SystemExit(f"Expected one run for paragraph edit, found {len(paragraph.runs)}: {paragraph.text[:90]!r}")
    paragraph.runs[0].text = text


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    source_bytes = args.input.read_bytes()
    source_hash = sha256(source_bytes)
    if source_hash != SOURCE_SHA256:
        raise SystemExit(f"Unexpected v1.1 source SHA-256: {source_hash}")

    document = Document(args.input)
    if len(document.paragraphs) != 159 or len(document.tables) != 3:
        raise SystemExit(f"Unexpected source structure: {len(document.paragraphs)} paragraphs, {len(document.tables)} tables")

    special_prereq = " Recommenders in which hard constraints — ordering, readiness, safety — filter the candidate set before ranking."
    for old, new in REPLACEMENTS.items():
        if old == special_prereq:
            matches = [p for p in document.paragraphs if p.text.endswith(old)]
        else:
            matches = [p for p in document.paragraphs if p.text == old]
        if len(matches) != 1:
            raise SystemExit(f"Expected one source paragraph matching {old[:100]!r}; found {len(matches)}")
        paragraph = matches[0]
        if old == special_prereq:
            if len(paragraph.runs) != 3 or paragraph.runs[1].text != "Prerequisite-Constrained Recommendation:":
                raise SystemExit("Unexpected formatting in the original recommendation heading")
            paragraph.runs[1].text = "Topic-and-Age-Grounded Activity Discovery:"
            paragraph.runs[2].text = new
        else:
            set_paragraph_text(paragraph, new)

    if len(document.tables) != 3 or len(document.paragraphs) != 159:
        raise SystemExit("The form structure changed unexpectedly during editing")
    original_tables = Document(args.input).tables
    for index in (0, 1, 2):
        old_table = [[cell.text for cell in row.cells] for row in original_tables[index].rows]
        new_table = [[cell.text for cell in row.cells] for row in document.tables[index].rows]
        if old_table != new_table:
            raise SystemExit(f"Original form table {index} was changed")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=args.output.parent) as tmp:
        temp_path = Path(tmp.name)
    try:
        document.save(temp_path)
        with zipfile.ZipFile(args.input, "r") as source_zip, zipfile.ZipFile(temp_path, "r") as edited_zip:
            edited_xml = edited_zip.read("word/document.xml")
            with zipfile.ZipFile(args.output, "w", compression=zipfile.ZIP_DEFLATED) as output_zip:
                for item in source_zip.infolist():
                    content = edited_xml if item.filename == "word/document.xml" else source_zip.read(item.filename)
                    output_zip.writestr(item, content)
    finally:
        temp_path.unlink(missing_ok=True)

    with zipfile.ZipFile(args.input, "r") as source_zip, zipfile.ZipFile(args.output, "r") as output_zip:
        if source_zip.namelist() != output_zip.namelist():
            raise SystemExit("DOCX package part list changed")
        changed_parts = [name for name in source_zip.namelist() if source_zip.read(name) != output_zip.read(name)]
        if changed_parts != ["word/document.xml"]:
            raise SystemExit(f"Unexpected DOCX package parts changed: {changed_parts}")

    result = Document(args.output)
    text = "\n".join(p.text for p in result.paragraphs)
    for expected in (
        "under 9 years (0–107 completed months)",
        "product age bands (0–3, 3–6 and 6–9)",
        "Product bands are 0–35, 36–71 and 72–107 months",
        "adult-supplied age/readiness context",
        "never infer age from a drawing or narration",
        "99,000 VND/month",
        "499,000 VND/class/month",
        "10 credits/49,000 VND",
        "30/129,000 VND",
        "60/239,000 VND",
        "backend payment verification",
    ):
        if expected.casefold() not in text.casefold():
            raise SystemExit(f"Required content missing: {expected}")
    for unwanted in (
        "aged 0-12",
        "subject to approval",
        "must be approved before collection",
        "approval gates",
        "Owner Review Draft",
        "REVIEW DRAFT",
        "provisional hypotheses",
        "waiting for owner approval",
    ):
        if unwanted.casefold() in text.casefold():
            raise SystemExit(f"Internal or outdated wording remains: {unwanted}")

    print(f"Output: {args.output}")
    print(f"Source SHA-256: {source_hash}")
    print(f"Output SHA-256: {sha256(args.output.read_bytes())}")
    print(f"Paragraphs: {len(result.paragraphs)}; tables: {len(result.tables)}")
    print("Changed DOCX package part: word/document.xml")


if __name__ == "__main__":
    main()
