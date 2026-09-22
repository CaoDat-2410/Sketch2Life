from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOCX = Path(__file__).resolve().parents[1] / "artifacts" / "Sketch2Life_SRS_Form.docx"
doc = Document(DOCX)
for table in doc.tables:
    if table.rows and table.cell(0, 0).text.startswith("Trạng thái"):
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        break
doc.save(DOCX)
print(DOCX)
