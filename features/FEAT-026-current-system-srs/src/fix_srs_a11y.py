from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX = Path(__file__).resolve().parents[1] / "artifacts" / "Sketch2Life_SRS_Form.docx"
doc = Document(DOCX)

for table in doc.tables:
    if table.rows and table.cell(0, 0).text.startswith("Implemented / boundary"):
        row = table.add_row()
        labels = ["Trạng thái", "Ý nghĩa trong tài liệu"]
        for cell, label in zip(row.cells, labels):
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(label)
            run.font.name = "Arial"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string("FFFFFF")
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "163A5F")
            cell._tc.get_or_add_tcPr().append(shd)
        tr = row._tr
        table._tbl.remove(tr)
        first_tr = table.rows[0]._tr
        table._tbl.insert(table._tbl.index(first_tr), tr)
        break

doc.save(DOCX)
print(DOCX)
