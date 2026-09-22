from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK


DOCX = Path(__file__).resolve().parents[1] / "artifacts" / "Sketch2Life_SRS_Form.docx"
doc = Document(DOCX)
for paragraph in doc.paragraphs:
    if paragraph.text.startswith("2.2 Ngoài phạm vi"):
        before = paragraph.insert_paragraph_before()
        before.add_run().add_break(WD_BREAK.PAGE)
        break
doc.save(DOCX)
print(DOCX)
