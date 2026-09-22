from pathlib import Path
from docx import Document


DOCX = Path(__file__).resolve().parents[1] / "artifacts" / "Sketch2Life_SRS_Form.docx"
doc = Document(DOCX)
for index, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.startswith("3 Hiện trạng hệ thống") and index > 0:
        previous = doc.paragraphs[index - 1]
        if not previous.text:
            previous._element.getparent().remove(previous._element)
        break
doc.save(DOCX)
print(DOCX)
