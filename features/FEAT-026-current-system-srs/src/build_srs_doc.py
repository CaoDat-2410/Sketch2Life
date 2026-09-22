from __future__ import annotations

import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FEATURE_ROOT = Path(__file__).resolve().parents[1]
ARTIFACT_DIR = FEATURE_ROOT / "artifacts"
FIGURE_DIR = ARTIFACT_DIR / "figures"
OUTPUT_DOCX = ARTIFACT_DIR / "Sketch2Life_SRS_Form.docx"

FIGURE_DIR.mkdir(parents=True, exist_ok=True)

FONT_REGULAR = Path("C:/Windows/Fonts/arial.ttf")
FONT_BOLD = Path("C:/Windows/Fonts/arialbd.ttf")

NAVY = "163A5F"
BLUE = "2F6B8A"
TEAL = "2A9D8F"
GREEN = "2E7D5B"
AMBER = "B7791F"
RED = "A94442"
LIGHT_BLUE = "EAF2F8"
LIGHT_TEAL = "E9F6F4"
LIGHT_AMBER = "FFF4D6"
LIGHT_RED = "FBEAEA"
LIGHT_GRAY = "F5F7F9"
MID_GRAY = "D9E0E7"
TEXT = "1F2933"
BLACK = "000000"


def pil_font(size: int, bold: bool = False):
    path = FONT_BOLD if bold and FONT_BOLD.exists() else FONT_REGULAR
    return ImageFont.truetype(str(path), size)


def wrap_lines(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, font, fill: str, max_lines: int = 4):
    x1, y1, x2, y2 = box
    lines = wrap_lines(draw, text, font, max(20, x2 - x1 - 28))[:max_lines]
    if len(lines) < max_lines and len(" ".join(lines)) < len(text):
        lines[-1] = lines[-1].rstrip(".") + "..."
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    line_gap = 5
    total = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + max(6, (y2 - y1 - total) // 2)
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) // 2, y), line, font=font, fill=fill)
        y += height + line_gap


def rounded_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, fill: str, outline: str = "FFFFFF", text_fill: str = "FFFFFF", font_size: int = 25):
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=3)
    centered_text(draw, box, text, pil_font(font_size, bold=True), text_fill)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str = NAVY, width: int = 5):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 >= x1 else -1
        tip = (x2, y2)
        p1 = (x2 - direction * 18, y2 - 10)
        p2 = (x2 - direction * 18, y2 + 10)
    else:
        direction = 1 if y2 >= y1 else -1
        tip = (x2, y2)
        p1 = (x2 - 10, y2 - direction * 18)
        p2 = (x2 + 10, y2 - direction * 18)
    draw.polygon([tip, p1, p2], fill=fill)


def label(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, size: int = 22, fill: str = TEXT, bold: bool = False):
    draw.text(xy, text, font=pil_font(size, bold), fill=fill)


def save_system_topology() -> Path:
    path = FIGURE_DIR / "system_topology.png"
    img = Image.new("RGB", (1900, 1120), "F7FAFC")
    draw = ImageDraw.Draw(img)
    label(draw, (55, 38), "Sketch2Life current and target boundary topology", 34, NAVY, True)
    label(draw, (55, 85), "The solid path is the intended application boundary; provider and persistence wiring are not fully live in the current checkout.", 20, TEXT)

    rounded_box(draw, (60, 190, 360, 335), "Parent / Guide\naccount", BLUE, font_size=25)
    rounded_box(draw, (60, 405, 360, 550), "Child\nsupervised mode", TEAL, font_size=25)
    rounded_box(draw, (465, 255, 850, 505), "Android React Native\nrole-aware app\nFixtureFlow + AuthSessionPort", NAVY, font_size=25)
    rounded_box(draw, (465, 625, 850, 825), "PixiJS + GSAP\ncontrolled bridge\noriginal-art renderer", "4B5563", font_size=25)
    rounded_box(draw, (1010, 250, 1400, 535), "FastAPI modular monolith\nApplication + Domain\nversioned contracts\nGate A / P1 / Gate B", BLUE, font_size=25)
    rounded_box(draw, (1510, 105, 1840, 260), "PostgreSQL\nplanned adapter", "64748B", font_size=24)
    rounded_box(draw, (1510, 310, 1840, 465), "Redis / RQ\nplanned jobs", "64748B", font_size=24)
    rounded_box(draw, (1510, 515, 1840, 670), "S3-compatible\nobject storage", "64748B", font_size=24)
    rounded_box(draw, (1510, 720, 1840, 875), "AiGateway\nbackend-only", "4B5563", font_size=24)
    rounded_box(draw, (1265, 905, 1530, 1050), "Lightning\ndev / fixture", AMBER, text_fill="1F2933", font_size=23)
    rounded_box(draw, (1580, 905, 1840, 1050), "Runpod\nproduction target", RED, font_size=23)

    arrow(draw, (360, 262), (465, 340), BLUE)
    arrow(draw, (360, 477), (465, 420), TEAL)
    arrow(draw, (850, 380), (1010, 380), NAVY)
    arrow(draw, (660, 505), (660, 625), "64748B")
    arrow(draw, (1400, 300), (1510, 185), "64748B")
    arrow(draw, (1400, 360), (1510, 390), "64748B")
    arrow(draw, (1400, 425), (1510, 590), "64748B")
    arrow(draw, (1400, 495), (1510, 795), "4B5563")
    arrow(draw, (1650, 875), (1410, 905), AMBER)
    arrow(draw, (1750, 875), (1710, 905), RED)
    label(draw, (890, 305), "HTTPS + Firebase ID token", 19, NAVY, True)
    label(draw, (865, 555), "renderer intent boundary", 19, "4B5563")
    label(draw, (1430, 705), "backend-owned", 18, TEXT)
    img.save(path)
    return path


def save_end_to_end_flow() -> Path:
    path = FIGURE_DIR / "end_to_end_flow.png"
    img = Image.new("RGB", (1850, 1220), "F7FAFC")
    draw = ImageDraw.Draw(img)
    label(draw, (50, 30), "Sketch2Life end-to-end experience flow", 34, NAVY, True)
    label(draw, (50, 78), "The current fixture flow proves the sequence; capture, persistence, auth, provider jobs and full mobile wiring remain open for integration.", 20, TEXT)
    boxes = [
        (60, 180, 360, 340, "1  Capture\ndrawing + narration", TEAL),
        (455, 180, 755, 340, "2  Validate\nimage + audio", BLUE),
        (850, 180, 1170, 340, "3  ASR + Vision\nproposal", NAVY),
        (1265, 180, 1585, 340, "4  Gate A\nconfirm meaning", "4B5563"),
        (160, 520, 500, 690, "5  P1\nhard rules + fit", BLUE),
        (700, 520, 1040, 690, "6  Gate B\napprove identity", "4B5563"),
        (1240, 520, 1610, 690, "7  Experience\nP3 art + P4 media", TEAL),
        (360, 890, 700, 1060, "8  Activity handoff\noff-screen", GREEN),
        (980, 890, 1320, 1060, "9  Feedback\nrecorded", GREEN),
    ]
    for x1, y1, x2, y2, text, color in boxes:
        rounded_box(draw, (x1, y1, x2, y2), text, color, font_size=24)
    arrow(draw, (360, 260), (455, 260))
    arrow(draw, (755, 260), (850, 260))
    arrow(draw, (1170, 260), (1265, 260))
    arrow(draw, (1425, 340), (330, 520))
    arrow(draw, (500, 605), (700, 605))
    arrow(draw, (1040, 605), (1240, 605))
    arrow(draw, (1425, 690), (530, 890))
    arrow(draw, (700, 975), (980, 975))
    label(draw, (160, 380), "recapture if invalid", 19, RED, True)
    arrow(draw, (250, 340), (250, 470), RED, width=4)
    rounded_box(draw, (60, 520, 500, 690), "Recapture\nno AI call", RED, font_size=25)
    arrow(draw, (500, 605), (700, 605), "FFFFFF", width=0)
    label(draw, (80, 1125), "Failure rule: media/provider failure may change media status or apply typed fallback, but must preserve the off-screen activity handoff.", 20, TEXT)
    img.save(path)
    return path


def save_state_machine() -> Path:
    path = FIGURE_DIR / "state_machine.png"
    img = Image.new("RGB", (1900, 1040), "F7FAFC")
    draw = ImageDraw.Draw(img)
    label(draw, (50, 30), "FEAT-016 session state model", 34, NAVY, True)
    label(draw, (50, 78), "Application-owned state is the intended source of truth; the mobile reducer is a local fixture representation.", 20, TEXT)
    boxes = {
        "CREATED": (80, 220, 360, 340, BLUE),
        "MEDIA_RECAPTURE": (80, 520, 390, 650, RED),
        "GATE_A_PENDING": (520, 220, 850, 340, "4B5563"),
        "UNDERSTANDING_PROPOSED": (960, 220, 1370, 340, NAVY),
        "CONTEXT_REQUIRED": (960, 520, 1370, 650, AMBER),
        "GATE_B_PENDING": (1480, 220, 1810, 340, "4B5563"),
        "EXPERIENCE_READY": (1480, 520, 1810, 650, TEAL),
        "HANDOFF_READY": (870, 820, 1190, 950, GREEN),
        "FEEDBACK_RECORDED": (1410, 820, 1800, 950, GREEN),
    }
    for key, (x1, y1, x2, y2, color) in boxes.items():
        rounded_box(draw, (x1, y1, x2, y2), key.replace("_", " "), color, font_size=22)
    arrow(draw, (360, 280), (520, 280))
    arrow(draw, (850, 280), (960, 280))
    arrow(draw, (1370, 280), (1480, 280))
    arrow(draw, (1200, 340), (1200, 520), AMBER)
    arrow(draw, (1370, 585), (1480, 585), AMBER)
    arrow(draw, (1645, 340), (1645, 520), TEAL)
    arrow(draw, (1645, 650), (1030, 820), GREEN)
    arrow(draw, (1190, 885), (1410, 885), GREEN)
    arrow(draw, (220, 340), (220, 520), RED)
    label(draw, (395, 245), "valid media", 18, GREEN, True)
    label(draw, (235, 415), "invalid media", 18, RED, True)
    label(draw, (1220, 405), "missing context", 18, AMBER, True)
    label(draw, (1658, 410), "Gate B approved", 18, TEAL, True)
    label(draw, (70, 990), "All commands carry command_id, session_id, expected version, actor reference and timestamp. Duplicate commands replay; stale versions are rejected.", 20, TEXT)
    img.save(path)
    return path


def save_maturity_matrix() -> Path:
    path = FIGURE_DIR / "maturity_matrix.png"
    img = Image.new("RGB", (1900, 1260), "FFFFFF")
    draw = ImageDraw.Draw(img)
    label(draw, (55, 35), "Current system maturity matrix", 34, NAVY, True)
    label(draw, (55, 83), "Qualitative snapshot from the repository current-system report; this is not a production-completion percentage.", 20, TEXT)
    x0, y0 = 55, 165
    row_h = 92
    area_w = 760
    col_w = 250
    headers = ["System area", "Boundary / local", "Offline / fixture", "Not wired production"]
    for i, head in enumerate(headers):
        x1 = x0 + (0 if i == 0 else area_w + (i - 1) * col_w)
        x2 = x1 + (area_w if i == 0 else col_w)
        draw.rectangle((x1, y0, x2, y0 + 72), fill=NAVY, outline="FFFFFF", width=2)
        centered_text(draw, (x1, y0, x2, y0 + 72), head, pil_font(22, True), "FFFFFF", max_lines=2)
    rows = [
        ("Governance / context / security", 1),
        ("Backend contract boundaries", 1),
        ("Mobile fixture UI", 2),
        ("P1 Montessori compiler", 2),
        ("P2 understanding", 2),
        ("P3 art renderer", 2),
        ("P4 learning media", 2),
        ("Runtime integration", 2),
        ("Production infra / release", 3),
    ]
    colors = [LIGHT_BLUE, LIGHT_TEAL, LIGHT_AMBER, LIGHT_RED]
    for r, (name, selected) in enumerate(rows):
        y1 = y0 + 72 + r * row_h
        y2 = y1 + row_h
        fill = "FFFFFF" if r % 2 == 0 else "F8FAFC"
        draw.rectangle((x0, y1, x0 + area_w, y2), fill=fill, outline=MID_GRAY, width=2)
        centered_text(draw, (x0 + 16, y1, x0 + area_w - 16, y2), name, pil_font(22, True), TEXT, max_lines=2)
        for c in range(1, 4):
            x1 = x0 + area_w + (c - 1) * col_w
            x2 = x1 + col_w
            active = c == selected
            draw.rectangle((x1, y1, x2, y2), fill=colors[c] if active else fill, outline=MID_GRAY, width=2)
            if active:
                centered_text(draw, (x1, y1, x2, y2), "●", pil_font(35, True), NAVY, max_lines=1)
    label(draw, (55, 1085), "Interpretation: the project has strong contract and fixture evidence, while the production vertical slice is still a roadmap item.", 20, TEXT)
    img.save(path)
    return path


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=MID_GRAY, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_layout(table, fixed=True):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed" if fixed else "autofit")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=10.5, color=TEXT, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def format_paragraph(paragraph, space_before=0, space_after=5, line=1.08, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def clear_cell(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    format_paragraph(p, 0, 0, 1.0)
    return p


def cell_text(cell, text, bold=False, color=TEXT, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    p = clear_cell(cell)
    p.alignment = align
    run = p.add_run(str(text))
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_borders(cell)
    return p


def set_cell_width(cell, width_inches: float):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None, header_fill=NAVY, font_size=9.2, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_layout(table)
    header = table.rows[0]
    repeat_header(header)
    for i, h in enumerate(headers):
        cell_text(header.cells[i], h, bold=True, color="FFFFFF", size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(header.cells[i], header_fill)
        if widths:
            set_cell_width(header.cells[i], widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            alignment = WD_ALIGN_PARAGRAPH.CENTER if i == status_col else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(cells[i], value, size=font_size, align=alignment)
            if widths:
                set_cell_width(cells[i], widths[i])
            set_cell_shading(cells[i], "FFFFFF" if ridx % 2 == 0 else LIGHT_BLUE)
    for row in table.rows:
        for cell in row.cells:
            set_cell_borders(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 12.5, color=BLACK, bold=True)
    return p


def add_body(doc, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="Normal")
    format_paragraph(p, 0, 6, 1.12)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=10.5, color=TEXT, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.5, color=TEXT)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=TEXT)
    return p


def add_bullets(doc, items: list[str], level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.22 + 0.18 * level)
        format_paragraph(p, 0, 2, 1.05)
        r = p.add_run(item)
        set_run_font(r, size=10.3, color=TEXT)


def add_caption(doc, text: str):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, 3, 4, 1.0, keep=True)
    r = p.add_run(text)
    set_run_font(r, size=9.2, color="5B6770", italic=True)
    return p


def add_figure(doc, path: Path, caption: str, width=6.9):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, 2, 2, 1.0, keep=True)
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    inline._inline.docPr.set("descr", caption)
    add_caption(doc, caption)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=8.5, color="5B6770")


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_paragraph(p, 0, 0, 1.0)
    r = p.add_run("Sketch2Life  |  SRS form")
    set_run_font(r, size=8.5, color="5B6770", bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, 0, 0, 1.0)
    r = p.add_run("Internal project document  |  Synthetic and fixture data only  |  Trang ")
    set_run_font(r, size=8.5, color="5B6770")
    add_page_field(p)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11.2)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Title"].paragraph_format.space_after = Pt(7)
    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    caption.font.size = Pt(9.2)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string("5B6770")


def add_status_legend(doc):
    rows = [
        ("Implemented / boundary", "Có code hoặc boundary đang chạy local; chưa mặc định là production.", GREEN),
        ("Offline / fixture-only", "Có harness, fixture, in-memory hoặc standalone package để kiểm thử.", AMBER),
        ("Accepted architecture", "Đã là quyết định kiến trúc nhưng cần integration/evidence trước khi triển khai.", BLUE),
        ("Planned / open", "Chưa wire, còn phụ thuộc approval, benchmark, môi trường hoặc quyết định tiếp theo.", RED),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_layout(table)
    for title, desc, color in rows:
        cells = table.add_row().cells
        cell_text(cells[0], title, bold=True, size=9.5, color=BLACK)
        cell_text(cells[1], desc, size=9.5)
        set_cell_shading(cells[0], color)
        set_cell_shading(cells[1], "FFFFFF")
        set_cell_width(cells[0], 1.65)
        set_cell_width(cells[1], 5.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document():
    topology = save_system_topology()
    flow = save_end_to_end_flow()
    state = save_state_machine()
    matrix = save_maturity_matrix()

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    add_header_footer(section)
    doc.core_properties.title = "Sketch2Life Software Requirements Specification"
    doc.core_properties.subject = "Current system SRS form"
    doc.core_properties.author = "Sketch2Life project"
    doc.core_properties.comments = "Based on repository current-system snapshot; synthetic and fixture data only."

    # Cover and opening
    p = doc.add_paragraph(style="Title")
    r = p.add_run("Đặc tả yêu cầu phần mềm Sketch2Life")
    set_run_font(r, size=24, color=BLACK, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p2, 0, 18, 1.0)
    r = p2.add_run("Biểu mẫu SRS hiện trạng hệ thống")
    set_run_font(r, size=13, color=BLUE, bold=True)

    control_rows = [
        ("Mã tài liệu", "SRS-S2L-001"),
        ("Phiên bản", "0.1  |  working baseline for review"),
        ("Ngày lập", "15/09/2026"),
        ("Trạng thái", "Cần review và xác nhận phạm vi"),
        ("Snapshot nguồn", "docs/CURRENT_SYSTEM_STATE.md  |  11/09/2026"),
        ("Checkout tham chiếu", "codex/feat-018-contract-plan"),
        ("Đối tượng đọc", "Project owner, team, reviewer, người tích hợp"),
        ("Phân loại", "Tài liệu nội bộ  |  synthetic và fixture data only"),
    ]
    add_table(doc, ["Trường", "Giá trị"], control_rows, widths=[1.75, 5.15], font_size=9.7)
    add_body(doc, "Tài liệu này chuyển snapshot hiện trạng Sketch2Life thành một biểu mẫu SRS có thể chỉnh sửa. Nó mô tả product boundary, requirements, contract, flow và tiêu chí nghiệm thu theo những gì repository hiện có; các phần chưa được nối thành production vertical slice được gắn trạng thái rõ ràng.")
    add_body(doc, "Mục đích sử dụng: dùng làm baseline để review phạm vi, phân biệt evidence fixture với production readiness, và ghi nhận các quyết định cần phê duyệt trước khi triển khai tiếp.", bold_prefix="Mục đích sử dụng:")
    doc.add_page_break()

    # 1
    add_heading(doc, "1 Tóm tắt và mục đích", 1)
    add_body(doc, "Sketch2Life biến bản vẽ và lời kể của trẻ thành một trải nghiệm học tập ngắn, sau đó chuyển trẻ sang một hoạt động Montessori ngoài màn hình và thu feedback từ người lớn. Hệ thống không phải generic image generator: bản vẽ gốc phải được bảo toàn, suy luận AI phải có human review, activity phải qua các luật xác định, và Gate B phải khóa identity/version trước khi nội dung số hoặc activity vật lý được sử dụng.")
    add_body(doc, "Kết luận hiện trạng: repository đang ở giai đoạn foundation với nhiều boundary, contract, fixture và standalone workstream đã được harden; chưa phải sản phẩm end-to-end production-ready.", bold_prefix="Kết luận hiện trạng:")
    add_status_legend(doc)

    add_heading(doc, "2 Phạm vi sản phẩm và tác nhân", 1)
    actor_rows = [
        ("Child", "Không có account độc lập; dùng supervised session để vẽ, kể, xem experience và bắt đầu activity.", "Trong fixture flow"),
        ("Parent", "Xác thực, consent/supervision, xác nhận hoặc sửa Gate A, duyệt Gate B và gửi feedback.", "Account model đã quyết định; UI/auth chưa wire"),
        ("Guide", "Review Montessori, activity/objective identity/version, supervision và feedback.", "Account model đã quyết định; UI/auth chưa wire"),
        ("Backend", "Giữ business truth, session/artifact/job state, validation, authorization, provenance và provider adapters.", "Boundary/fixture hiện có; persistence còn mở"),
        ("AI providers", "Lightning cho dev/fixture; Runpod là production target qua AiGateway.", "Provider boundary; model profile còn benchmark-gated"),
    ]
    add_table(doc, ["Tác nhân", "Trách nhiệm", "Trạng thái"], actor_rows, widths=[1.05, 4.15, 1.7], font_size=9.1, status_col=2)

    add_heading(doc, "2.1 Trong phạm vi", 2)
    add_bullets(doc, [
        "Android-only React Native app với child, parent và guide modes trong một application.",
        "Backend modular monolith theo clean architecture: application/domain/contracts/interfaces/infrastructure.",
        "Multimodal understanding, Gate A, P1 deterministic filter, Gate B, original-art renderer, learning-media cache/fallback, handoff và feedback.",
        "Versioned contracts, immutable source/derived artifact lineage, stale-version và idempotency semantics.",
        "Fixture/contract-driven evidence làm baseline cho Integration Sprint.",
    ])
    add_heading(doc, "2.2 Ngoài phạm vi hoặc chưa sẵn sàng production", 2)
    add_bullets(doc, [
        "Firebase ID-token verification middleware, role/authorization persistence và parent/guide sign-in UI.",
        "Camera/image picker/microphone capture, consent UX, backend session repository, migrations và persistence adapter.",
        "S3/MinIO adapter, Redis/RQ worker, job resource API, retention/deletion job và observability production.",
        "Runpod production adapter, model-profile freeze, live video generation, signed APK/AAB và Play Console release.",
        "Qualified Montessori catalog: records hiện là synthetic/provisional và production_eligible=false.",
    ])
    doc.add_page_break()

    # 3
    add_heading(doc, "3 Hiện trạng hệ thống", 1)
    add_body(doc, "Các nhãn trong bảng dưới đây phản ánh snapshot hiện tại, không phải cam kết tiến độ. Những workstream đã pass test offline vẫn cần integration evidence trước khi được xem là production capability.")
    maturity_rows = [
        ("Governance / context / security", "Validators và context harness chạy local", "Boundary/local"),
        ("Backend", "FastAPI health/live-fixture route, typed contracts, P1 admission/compiler và provider guard", "Boundary/fixture"),
        ("Mobile", "Android React Native fixture flow, optional live-fixture call, auth port, Pixi protocol boundary", "Fixture UI"),
        ("P1 Montessori", "100 MVP activities, 20 golden activities, deterministic rules, ExperienceSpec, Gate B integrity", "Offline; provisional"),
        ("P2 understanding", "Media validation, ASR/VLM contracts, fixtures và Lightning synthetic route", "Offline/dev"),
        ("P3 renderer", "PixiJS + GSAP package, closed Motion DSL, provenance và fallback", "Standalone package"),
        ("P4 learning media", "Cache-first resolver, reviewed identity check, typed fallback/block, replay", "Offline/in-memory"),
        ("Runtime integration", "FEAT-015 fixture flow và FEAT-016 in-memory session/job semantics", "Test harness"),
        ("Production infra / release", "Compose và Android skeleton; chưa wire auth/DB/storage/worker/signing", "Chưa wire"),
    ]
    add_table(doc, ["Vùng hệ thống", "Bằng chứng hiện có", "Maturity"], maturity_rows, widths=[1.75, 4.15, 1.0], font_size=8.7, status_col=2)
    add_figure(doc, matrix, "Hình 1. Ma trận hiện trạng theo vùng hệ thống", width=6.9)

    add_heading(doc, "3.1 Repository boundary", 2)
    add_bullets(doc, [
        "apps/mobile: presentation và device adapters; không giữ business truth.",
        "backend/src/sketch2life: Python modular monolith gồm domain, application, contracts, interfaces và infrastructure.",
        "packages/contracts: language-neutral versioned schemas và fixtures.",
        "packages/domain-montessori và packages/domain-experience: domain/spec boundary và catalog fixtures.",
        "packages/art-renderer: PixiJS/GSAP original-art renderer độc lập.",
        "features/: plan, approval, context, evidence và fixture isolation.",
    ])

    # 4
    add_heading(doc, "4 Kiến trúc và topology", 1)
    add_body(doc, "Kiến trúc được tổ chức theo dependency direction: interface adapters đi vào application use cases và ports; application phụ thuộc domain; infrastructure implement ports; contracts được version hóa ở boundary. Modular monolith là boundary đầu tiên, chỉ tách service khi có evidence về scale, ownership hoặc reliability.")
    add_figure(doc, topology, "Hình 2. Topology boundary hiện tại và mục tiêu đã được chấp nhận", width=6.9)
    add_body(doc, "Đường mobile chỉ đi qua backend HTTPS. Mobile không gọi Lightning, Runpod, S3 hoặc database trực tiếp. Firebase chỉ cung cấp Authentication; backend mới là nơi verify identity và quyết định authorization.")
    add_heading(doc, "4.1 Baseline platform", 2)
    platform_rows = [
        ("Client", "Bare React Native + TypeScript; Android-only; com.sketch2life.mobile"),
        ("Android", "minSdk 29; targetSdk 36; compileSdk 37"),
        ("Backend", "Python 3.12+; FastAPI + Pydantic boundary"),
        ("Async progress", "Bounded HTTP polling; bắt đầu khoảng 2 giây, backoff tối đa 10 giây, dừng ở terminal state"),
        ("Identity", "Firebase Authentication cho parent/guide; backend verify token"),
        ("Data", "PostgreSQL + S3-compatible storage + Redis/RQ là target adapters"),
        ("AI", "Lightning dev/fixture; Runpod production target qua AiGateway"),
    ]
    add_table(doc, ["Hạng mục", "Baseline"], platform_rows, widths=[1.45, 5.45], font_size=9.4)
    doc.add_page_break()

    # 5 flows
    add_heading(doc, "5 Luồng nghiệp vụ chính", 1)
    add_figure(doc, flow, "Hình 3. Luồng end-to-end từ capture đến feedback", width=6.9)
    flow_rows = [
        ("Capture", "Nhận drawing và narration trong supervised session; current capture UI/device path chưa hoàn chỉnh.", "Planned"),
        ("Media validation", "Đọc bounded snapshot, hash source, inspect PNG/WAV, trả PASS hoặc RECAPTURE với reason ổn định.", "Implemented / fixture"),
        ("Understanding", "ASR + vision tạo proposal; live route hiện dừng trước Gate A.", "Fixture/dev"),
        ("Gate A", "Adult xác nhận hoặc sửa meaning; semantic anchors phải có claim provenance.", "Fixture/runtime"),
        ("P1", "Deterministic safety, age/readiness, prerequisite, material, supervision và policy chạy trước model selector.", "Offline"),
        ("Gate B", "Adult/guide duyệt activity identity/version và learning-objective identity/version đồng thời.", "Fixture/runtime"),
        ("Experience", "P3 original-art animation và P4 reviewed media/cache/fallback; hai artifact giữ riêng.", "Standalone/offline"),
        ("Handoff + feedback", "Ready session phải dẫn đến off-screen activity và feedback; failure media không được xóa handoff.", "Fixture/runtime"),
    ]
    add_table(doc, ["Bước", "Behavior và rule", "Maturity"], flow_rows, widths=[1.35, 4.6, 0.95], font_size=8.9, status_col=2)

    add_heading(doc, "5.1 Live understanding dev boundary", 2)
    add_body(doc, "Route hiện có: POST /v1/live-understanding. Mobile synthetic client gửi mode=live-lightning, fixture_id, session_id, expected_session_version và request_id tùy chọn. Route parse request với extra=forbid, reject provider khác lightning_dev, reject expected version khác 1 bằng STALE_SESSION_VERSION, load synthetic fixture, chạy media validation rồi tạo ASR/vision proposal.")
    add_body(doc, "Route này chưa tạo backend session aggregate, chưa persist artifact, chưa gọi P1/Gate B, chưa queue job và chưa verify Firebase trong router. Vì vậy đây là dev smoke boundary, không phải public production API.", bold_prefix="Route này chưa")

    # 6 state/contracts
    add_heading(doc, "6 State machine và version semantics", 1)
    add_figure(doc, state, "Hình 4. Session state model trong FEAT-016 runtime fixture", width=6.9)
    state_rows = [
        ("CREATED", "Valid media -> GATE_A_PENDING; invalid media -> MEDIA_RECAPTURE", "Session init"),
        ("GATE_A_PENDING", "Chờ adult confirmation/correction", "No bypass by navigation"),
        ("UNDERSTANDING_PROPOSED", "Proposal đủ context -> Gate B; thiếu context -> CONTEXT_REQUIRED", "Application handler"),
        ("GATE_B_PENDING", "Chờ activity/objective identity/version approval", "Exact refs required"),
        ("EXPERIENCE_READY", "Compile immutable ExperienceSpec và prepare media/renderer", "Fit + continuity"),
        ("HANDOFF_READY", "ActivityHandoff READY hoặc typed blocked state", "Handoff preserved"),
        ("FEEDBACK_RECORDED", "Feedback đã được ghi nhận; terminal state", "No overwrite"),
    ]
    add_table(doc, ["State", "Transition rule", "Control"], state_rows, widths=[1.8, 3.7, 1.4], font_size=9.0)
    add_bullets(doc, [
        "Cùng command_id replay -> trả lại snapshot tương ứng, không tăng version lần hai.",
        "Expected session version cũ -> reject STALE_SESSION_VERSION.",
        "Job completion/cancellation lệch version -> reject; terminal job không bị ghi đè.",
        "Renderer sequence không tăng đơn điệu -> bỏ qua; Gate A/B không được bypass bằng local flag.",
    ])

    add_heading(doc, "6.1 Contract và provenance envelope", 2)
    contract_rows = [
        ("Identity", "contract_name, contract_version, session_id, expected_session_version", "Backend contracts authoritative"),
        ("Artifact", "artifact_id, artifact_version, source_artifact_ids[], created_at", "Original immutable; derivative traceable"),
        ("Provenance", "model, config, adapter version, actor, reason, hash/version", "Required where relevant"),
        ("Media", "SourceMediaReferenceV1, validation decision, SHA-256, bounded signals", "Hash required when AVAILABLE"),
        ("P1 / Gate B", "ActivityTemplateV1, ExperienceSpecV1, IntegrationGateDecisionV1, ActivityHandoffV1", "Exact identity/version/hash"),
        ("Learning media", "cache key, reviewed asset identity, fallback_type, generation_called", "Cache HIT -> READY; non-READY has reason"),
    ]
    add_table(doc, ["Nhóm", "Field / contract", "Rule"], contract_rows, widths=[1.2, 3.6, 2.1], font_size=8.9)
    doc.add_page_break()

    # 7 requirements
    add_heading(doc, "7 Functional requirements", 1)
    add_body(doc, "Các requirement dưới đây được tách thành capability đã có evidence, capability fixture/offline, và target cần integration. Reviewer nên cập nhật cột status khi có approval hoặc evidence mới.")
    fr_rows = [
        ("FR-001", "Hệ thống phải mở supervised session cho child và giữ session version để chống stale command.", "Planned / fixture model", "FEAT-016 contracts; backend session chưa persist"),
        ("FR-002", "Hệ thống phải validate image/audio trước AI; invalid media phải trả recapture reason và không gọi AI.", "Implemented / fixture", "media_validation.py; FEAT-015 invalid_media_recap"),
        ("FR-003", "Hệ thống phải lưu original child media immutable và gắn SHA-256 khi source AVAILABLE.", "Contract rule", "CURRENT_SYSTEM_STATE 3, 8.2"),
        ("FR-004", "Hệ thống phải tạo ASR/vision proposal từ source đã PASS và trả typed failure khi provider/schema lỗi.", "Fixture/dev", "understanding contracts; live-understanding"),
        ("FR-005", "Hệ thống phải yêu cầu Gate A xác nhận/correct semantic meaning trước P1.", "Fixture/runtime", "fixtureFlow; P1ExperienceCompiler"),
        ("FR-006", "P1 phải chạy hard rules về safety, age/readiness, prerequisite, material, supervision và policy trước model selection.", "Offline implemented", "p1_experience.py; ADR/product invariants"),
        ("FR-007", "Gate B phải khóa đồng thời activity identity/version và learning-objective identity/version.", "Fixture/runtime", "IntegrationGateDecisionV1; FEAT-015"),
        ("FR-008", "ExperienceSpec phải immutable, có canonical hash, fit evaluation và strict continuity.", "Offline implemented", "ExperienceSpecV1; P1 compiler"),
        ("FR-009", "Original-art animation phải dùng source artifact/provenance và fallback whole-drawing khi extraction/mask lỗi.", "Standalone package", "packages/art-renderer/src"),
        ("FR-010", "Learning media phải resolve exact cache key và reviewed identity; miss/stale/corrupt/unsafe trả typed fallback/block.", "Offline/in-memory", "learning_media_resolver.py"),
        ("FR-011", "Media/provider failure không được loại bỏ off-screen activity handoff.", "Fixture evidence", "FEAT-015 cache_miss_fallback, blocked_learning_media"),
        ("FR-012", "Ready session phải đi qua Activity Bridge, start activity và feedback trước terminal completion.", "Fixture/runtime", "fixtureFlow sequence v4->v6"),
        ("FR-013", "Backend phải verify Firebase ID token và map role/relationship trước authorization.", "Planned", "ADR-0005; auth adapter chưa wire"),
        ("FR-014", "Mobile phải chỉ gọi backend HTTPS và poll backend job resource với bounded backoff.", "Accepted architecture", "ADR-0004; contracts integration"),
        ("FR-015", "Backend phải sở hữu storage, retention/deletion và artifact access; mobile không nhận S3/provider credentials.", "Planned / security", "AGENTS.md; ADR-0005"),
        ("FR-016", "Runpod production provider phải được nối sau benchmark/profile freeze qua AiGateway, không thay đổi domain/application.", "Planned", "ADR-0003/0005; benchmark gate"),
    ]
    add_table(doc, ["ID", "Yêu cầu", "Maturity", "Traceability"], fr_rows, widths=[0.6, 3.75, 1.05, 1.5], font_size=8.2, status_col=2)

    add_heading(doc, "8 Non-functional requirements", 1)
    nfr_rows = [
        ("NFR-001", "Clean architecture", "Domain không import UI/framework/ORM/queue/provider SDK; infrastructure implement ports.", "Architecture validators"),
        ("NFR-002", "Contract integrity", "Boundary schemas versioned; mobile types generated/derived, không tự định nghĩa business truth.", "CONTRACTS_AND_INTEGRATION"),
        ("NFR-003", "Provenance", "Derived/generated artifact phải tham chiếu source IDs, version, hash, config, actor/reason.", "Product invariants"),
        ("NFR-004", "Idempotency", "Duplicate command không tăng version; stale async result không mutate state mới.", "FEAT-016 runtime"),
        ("NFR-005", "Resilience", "Typed failure/fallback; media failure không phá handoff.", "FEAT-015 scenarios"),
        ("NFR-006", "Security boundary", "No .env/credentials/real child data; Firebase Auth-only; provider secret backend-only.", "AGENTS.md; ADR-0005"),
        ("NFR-007", "Progress transport", "Polling bắt đầu khoảng 2 giây, backoff tối đa 10 giây, dừng terminal/background; SSE/WebSocket cần ADR.", "CONTRACTS_AND_INTEGRATION"),
        ("NFR-008", "Platform", "Android-only; com.sketch2life.mobile; minSdk 29; target 36; compile 37.", "ADR-0004"),
        ("NFR-009", "Testability", "Mỗi workstream chạy bằng synthetic fixture/versioned contract; evidence lưu trong feature folder.", "ADR-0006; governance"),
        ("NFR-010", "Production eligibility", "Catalog/activity/objective/asset chỉ dùng production sau review/benchmark/approval; current records provisional.", "P1/Gate B rules"),
    ]
    add_table(doc, ["ID", "Quality area", "Requirement", "Evidence / source"], nfr_rows, widths=[0.7, 1.25, 3.5, 1.45], font_size=8.4)
    doc.add_page_break()

    # 9 security and acceptance
    add_heading(doc, "9 Security privacy và data handling", 1)
    sec_rows = [
        ("SEC-001", "Không commit .env, token, service account, signing key, seed account hoặc real child data.", "Mandatory"),
        ("SEC-002", "Firebase chỉ dùng Authentication; cấm Storage, Firestore và Realtime Database.", "Mandatory"),
        ("SEC-003", "Mobile không chứa endpoint/credential của S3, Lightning, Runpod hoặc database.", "Mandatory"),
        ("SEC-004", "Object storage do backend ports sở hữu; mobile dùng short-lived backend-controlled references.", "Target architecture"),
        ("SEC-005", "Handbook/workbook originals và rendered extracts ở local external reference, không publish.", "Mandatory"),
        ("SEC-006", "Evidence chỉ lưu hash/metadata bounded; không lưu raw image, prompt, model output, token, signed URL hoặc provider headers.", "Mandatory"),
        ("SEC-007", "Consent, least privilege, retention và deletion phải được thể hiện trong session/artifact design trước production.", "Planned"),
    ]
    add_table(doc, ["ID", "Security / privacy rule", "Maturity"], sec_rows, widths=[0.8, 4.8, 1.3], font_size=9.0, status_col=2)

    add_heading(doc, "10 Acceptance matrix", 1)
    add_body(doc, "Đây là các scenario hiện có trong fixture/evidence. Cột review result là phần nhóm có thể cập nhật khi nối production vertical slice.")
    acc_rows = [
        ("happy_cache_hit", "READY_FOR_OFFSCREEN_ACTIVITY; handoff=true", "[ ] Reconfirm in integrated runtime"),
        ("modality_conflict_requires_gate_a", "Giữ claims; GATE_A_REQUIRED", "[ ] Reconfirm"),
        ("no_eligible_activity_add_context", "CONTEXT_REQUIRED", "[ ] Reconfirm"),
        ("cache_miss_fallback", "Fallback typed; handoff preserved", "[ ] Reconfirm"),
        ("blocked_learning_media", "Media blocked; handoff preserved", "[ ] Reconfirm"),
        ("asset_integrity_failure", "Reject hash/manifest", "[ ] Reconfirm"),
        ("invalid_media_recap", "Recapture; no AI call", "[ ] Reconfirm"),
        ("stale_gate_or_completion", "Reject stale version", "[ ] Reconfirm"),
    ]
    add_table(doc, ["Scenario", "Expected result", "Review result"], acc_rows, widths=[2.05, 2.75, 2.1], font_size=8.9)

    add_heading(doc, "10.1 Verification checklist", 2)
    add_bullets(doc, [
        "[ ] Repository security validator passes before commit/push.",
        "[ ] Contract registry V1 is frozen and duplicate ASR/vision generations are reconciled.",
        "[ ] Firebase auth verification and role authorization have fixture evidence.",
        "[ ] Session/artifact/job repository, storage adapter, queue worker and polling endpoint have one vertical-slice evidence.",
        "[ ] Pixi bridge event registry matches fixture lifecycle and full mobile playback is demonstrated.",
        "[ ] Learning media review, fallback, retention and deletion are verified without real child data.",
        "[ ] Android debug APK and later signed release artifacts are proven on approved API levels.",
    ])
    doc.add_page_break()

    # 11 traceability and open decisions
    add_heading(doc, "11 Traceability và roadmap", 1)
    trace_rows = [
        ("Product purpose / invariants", "docs/context/PROJECT_CONTEXT.md; docs/CURRENT_SYSTEM_STATE.md sections 1 and 3"),
        ("Current implementation snapshot", "docs/CURRENT_SYSTEM_STATE.md sections 4, 5, 6, 8, 11, 12"),
        ("Source authority boundary", "docs/context/SOURCE_REGISTER.md"),
        ("Clean architecture", "docs/architecture/OVERVIEW.md; BASE_PROJECT_STRUCTURE.md; DEPENDENCY_RULES.md"),
        ("Contract / polling rules", "docs/architecture/CONTRACTS_AND_INTEGRATION.md"),
        ("Android / private AI boundary", "docs/adr/ADR-0004-android-only-private-ai-boundary.md"),
        ("Auth / release / provider", "docs/adr/ADR-0005-auth-release-and-ai-provider-strategy.md"),
        ("Parallel workstream boundary", "docs/adr/ADR-0006-parallel-sprint-allocation.md"),
        ("Offline integration evidence", "features/FEAT-015-integration-readiness-review/"),
        ("Runtime semantics", "features/FEAT-016-runtime-integration/"),
        ("Current system documentation", "features/FEAT-019-current-system-documentation/"),
    ]
    add_table(doc, ["Trace target", "Repository source"], trace_rows, widths=[2.0, 4.9], font_size=8.9)
    add_heading(doc, "11.1 Dependency-driven next steps", 2)
    add_bullets(doc, [
        "Freeze shared contract registry V1, resolve ASR/vision generation split và reconcile Pixi bridge event registry.",
        "Install Android SDK and prove blank debug APK on API 29/36.",
        "Complete Firebase Authentication adapter/backend verification with fixtures.",
        "Create application-owned session/artifact/job/auth API vertical slice.",
        "Wire PostgreSQL, S3-compatible storage và Redis/RQ through ports/provenance rules.",
        "Connect P2 -> Gate A -> P1 -> Gate B -> P3/P4 -> handoff in one approved vertical slice.",
        "Benchmark Lightning dev, then benchmark/freeze Runpod production profile.",
        "Add privacy/retention/deletion, redacted observability, CI security/type/test gates and signed Android release.",
    ])

    add_heading(doc, "12 Open decisions và review form", 1)
    open_rows = [
        ("OD-001", "Contract registry V1: chọn generation ASR/vision nào làm shared authority?", "Owner + architecture review"),
        ("OD-002", "Pixi bridge event registry: reconcile protocol messages với fixture lifecycle.", "Renderer + mobile review"),
        ("OD-003", "Exact model profiles / revisions / compute profiles sau benchmark.", "AI evidence + owner approval"),
        ("OD-004", "PostgreSQL/S3/Redis/RQ migration và persistence boundary cho Integration Sprint.", "Separate integration allocation"),
        ("OD-005", "Consent, retention, deletion và privacy UX trước khi nhận real child data.", "Security/privacy review"),
        ("OD-006", "Catalog/activity/objective qualification để chuyển production_eligible=true.", "Montessori review"),
    ]
    add_table(doc, ["ID", "Câu hỏi cần quyết định", "Người / evidence"], open_rows, widths=[0.7, 4.6, 1.6], font_size=8.9)

    add_heading(doc, "12.1 Sign-off", 2)
    sign_rows = [
        ("Product owner", "Tên: ______________________________", "Ngày: __________", "[ ] Approved  [ ] Changes requested"),
        ("Architecture reviewer", "Tên: ______________________________", "Ngày: __________", "[ ] Approved  [ ] Changes requested"),
        ("Security reviewer", "Tên: ______________________________", "Ngày: __________", "[ ] Approved  [ ] Changes requested"),
        ("Montessori reviewer", "Tên: ______________________________", "Ngày: __________", "[ ] Approved  [ ] Changes requested"),
    ]
    add_table(doc, ["Vai trò", "Tên / chữ ký", "Ngày", "Kết quả"], sign_rows, widths=[1.3, 2.35, 1.0, 2.25], font_size=8.8)
    add_body(doc, "Ghi chú review: ____________________________________________________________________________________")
    add_body(doc, "________________________________________________________________________________________________")
    add_body(doc, "________________________________________________________________________________________________")

    # Appendix
    doc.add_page_break()
    add_heading(doc, "Phụ lục A Form cập nhật requirement", 1)
    add_body(doc, "Dùng mẫu này khi thêm hoặc thay đổi requirement. Mọi thay đổi behavior phải đi kèm source/contract, approval, evidence và status cập nhật.")
    req_form_rows = [
        ("Requirement ID", "FR-____ / NFR-____ / SEC-____"),
        ("Tên requirement", "____________________________________________________________"),
        ("Mô tả", "____________________________________________________________\n____________________________________________________________"),
        ("Actor / trigger", "____________________________________________________________"),
        ("Precondition", "____________________________________________________________"),
        ("Expected result", "____________________________________________________________\n____________________________________________________________"),
        ("Status", "[ ] Implemented  [ ] Fixture-only  [ ] Accepted architecture  [ ] Planned  [ ] Open"),
        ("Contract / source", "____________________________________________________________"),
        ("Acceptance evidence", "____________________________________________________________"),
        ("Approval", "Approver: _____________________  Date: __________  Decision: __________"),
    ]
    add_table(doc, ["Trường", "Nội dung điền"], req_form_rows, widths=[1.65, 5.25], font_size=9.1)
    add_heading(doc, "Phụ lục B Nguyên tắc đọc tài liệu", 1)
    add_bullets(doc, [
        "Current/implemented không đồng nghĩa production-ready.",
        "Fixture-only là evidence có thể tái lập, không phải live provider hoặc production infrastructure.",
        "Accepted architecture là quyết định định hướng; implementation vẫn cần feature plan và approval riêng.",
        "Mọi derived artifact phải giữ provenance; original child media không bị thay thế âm thầm.",
        "Khi có mâu thuẫn, ưu tiên direct task, ADR/decision đã approved, source/evidence hiện tại, rồi mới đến external reference.",
    ])

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()
