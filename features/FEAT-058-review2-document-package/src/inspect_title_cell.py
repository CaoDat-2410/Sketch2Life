from docx import Document
from lxml import etree
p=r'features/FEAT-058-review2-document-package/artifacts/SWP490_FA26SE225_R2_Reports1-5_Tracking.docx'; d=Document(p); c=d.tables[1].cell(2,0)
print(etree.tostring(c._tc,pretty_print=True,encoding='unicode')[:5000])
