"""Refresh the local Rev11 report gallery with the approved Rev12 Draw.io set."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from PIL import Image


FEATURE = Path(__file__).resolve().parents[1]
PACKAGE_IN = FEATURE / "artifacts" / "submission_r2_rev11"
PACKAGE_OUT = FEATURE / "artifacts" / "submission_r2_rev12"
SUPPLIED_IMAGES = FEATURE / "evidence" / "render" / "rev12" / "supplied-diagrams-clean"
AUTHORED_IMAGES = FEATURE / "evidence" / "render" / "rev12" / "diagrams"

NAMES = {
    "brief": "SWP490_FA26SE225_R2_DesignBrief",
    "registration": "SWP490_FA26SE225_R2_Registration",
    "report": "SWP490_FA26SE225_R2_Report1-5_Tracking",
    "selfreview": "SWP490_FA26SE225_R2_SelfReview",
}
SOURCE_DIAGRAMS = [
    ("D01", "Business Context", "Context Diagram", "context_page01.png"),
    ("D02", "C4 System Context", "C4 System Context Diagram", "context_page02.png"),
    *[(f"D{index:02d}", "Use Case Overview" if index == 3 else f"Use Case Details — page {index - 2}", "UML Use Case Diagram", f"usecases_page{index - 2:02d}.png") for index in range(3, 15)],
    ("D15", "Main Session Sequence", "UML Sequence Diagram", "sequence_page01.png"),
    ("D16", "Detailed Session Sequence", "UML Sequence Diagram", "sequence_page02.png"),
    ("D17", "Screen Flow", "Screen Flow Diagram", "thieu_page01.png"),
    ("D18", "Class Diagram", "UML Class Diagram", "thieu_page02.png"),
    ("D19", "Logical ERD", "Logical Entity-Relationship Diagram", "thieu_page03.png"),
    ("D20", "Khoa ERD", "Entity-Relationship Diagram", "khoa_page01.png"),
    ("D21", "System Architecture", "System Architecture Diagram", "khoa_page02.png"),
    ("D22", "Package Diagram", "UML Package Diagram", "khoa_page03.png"),
]


def set_cell(cell, value: str) -> None:
    cell.text = value
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"


AUTHORED_TITLES = [
    "Session Lifecycle",
    "Session Pause and Recovery",
    "Session Terminal Outcomes",
    "Shared Learning Stage Lifecycle",
    "Group Progress Lifecycle",
    "Participant Admission Lifecycle",
    "Device Connection Lifecycle",
    "Device Grant Revocation",
    "Draft Synchronization Lifecycle",
    "Sketch Pipeline Lifecycle",
    "Child Suggestion Response Lifecycle",
    "Knowledge Editorial Lifecycle",
    "AI and Media Job Lifecycle",
    "Video Delivery Lifecycle",
    "Activity Execution Lifecycle",
    "Observation Lifecycle",
    "Content Publication Lifecycle",
    "Data Request Lifecycle",
    "Data Request Cancellation Outcomes",
    "Data Request External Purge Outcomes",
    "WF-01 Prepare and Admit",
    "WF-02 Collaborative Drawing and Recovery",
    "WF-03 Request and Review AI Sketch",
    "WF-04 Reviewed Learning and Session Close",
]


def set_paragraph(paragraph, value: str) -> None:
    paragraph.text = value
    for run in paragraph.runs:
        run.font.name = "Arial"


def replace_report_references(doc) -> None:
    replacements = {
        "Figures D41–D44": "Figures D43–D46",
        "Figures D41-D44": "Figures D43-D46",
        "Figures D23–D40": "Figures D23–D42",
        "Figures D23-D40": "Figures D23-D42",
        "in eighteen focused views": "in twenty focused views",
        "session and device companion views": "session, device and Data Request companion views",
        "Rev11 Draw.io master": "Rev12 Draw.io master",
        "Rev11 diagrams": "Rev12 diagrams",
        "Rev11 diagram set": "Rev12 diagram set",
        "Rev11. The official `review-2` tag": "Rev12. The official `review-2` tag",
        "this Rev11 dossier update is a documentation publication.": "this Rev12 dossier update is a documentation publication.",
        "this package revision is Rev11.": "this package revision is Rev12.",
        "the Review 2 dossier is Rev11.": "the Review 2 dossier is Rev12.",
    }
    # Apply longer, unambiguous replacements first to avoid cascading D43/D44 edits.
    for paragraph in doc.paragraphs:
        text = paragraph.text
        for old, new in replacements.items():
            text = text.replace(old, new)
        figure_refs = {
            "Figure D41.": "__FIGURE_REF_1__",
            "Figure D42.": "__FIGURE_REF_2__",
            "Figure D43.": "__FIGURE_REF_3__",
            "Figure D44.": "__FIGURE_REF_4__",
        }
        for old, token in figure_refs.items():
            text = text.replace(old, token)
        for index, value in enumerate(("Figure D43.", "Figure D44.", "Figure D45.", "Figure D46."), start=1):
            text = text.replace(f"__FIGURE_REF_{index}__", value)
        if text != paragraph.text:
            set_paragraph(paragraph, text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                for old, new in replacements.items():
                    text = text.replace(old, new)
                figure_refs = {
                    "Figure D41.": "__FIGURE_REF_1__",
                    "Figure D42.": "__FIGURE_REF_2__",
                    "Figure D43.": "__FIGURE_REF_3__",
                    "Figure D44.": "__FIGURE_REF_4__",
                }
                for old, token in figure_refs.items():
                    text = text.replace(old, token)
                for index, value in enumerate(("Figure D43.", "Figure D44.", "Figure D45.", "Figure D46."), start=1):
                    text = text.replace(f"__FIGURE_REF_{index}__", value)
                if text != cell.text:
                    set_cell(cell, text)


def add_before(anchor, paragraph):
    anchor.addprevious(paragraph._p)
    return paragraph


def add_gallery(doc) -> None:
    body = doc._body._element
    gallery_heading = next(
        p for p in doc.paragraphs
        if p.text.strip() == "Appendix D. Editable Draw.io Diagram Gallery"
    )
    appendix_e = next(
        p for p in doc.paragraphs
        if p.text.strip() == "Appendix E. Current Delivery, Proposed Next Work and Unknowns"
    )
    h_xml, e_xml = gallery_heading._p, appendix_e._p
    children = list(body)
    start, end = children.index(h_xml), children.index(e_xml)
    for element in children[start + 1 : end]:
        body.remove(element)

    intro = doc.add_paragraph(
        "Figures D01–D22 reproduce the five supplied Draw.io files as derived, editable pages. "
        "Figures D23–D42 present the fifteen SRS v3.3/B19.4 lifecycle registries in twenty focused "
        "state-machine views, including session, device and Data Request companion views. Figures "
        "D43–D46 provide one activity diagram for each main workflow. All 46 pages are editable in "
        "the Rev12 Draw.io master and labeled by diagram type. These lifecycle and workflow figures "
        "express design intent; they do not prove runtime acceptance."
    )
    intro.paragraph_format.space_after = Pt(8)
    add_before(e_xml, intro)

    pages: list[tuple[str, str, str, Path]] = []
    for figure_id, title, diagram_type, filename in SOURCE_DIAGRAMS:
        pages.append((figure_id, title, diagram_type, SUPPLIED_IMAGES / filename))
    for index, title in enumerate(AUTHORED_TITLES, start=1):
        figure_id = f"D{22 + index:02d}"
        diagram_type = "State Machine Diagram" if index <= 20 else "Activity Diagram"
        pages.append((figure_id, title, diagram_type, AUTHORED_IMAGES / f"page-{index:02d}.png"))

    section = doc.sections[0]
    available_width = (section.page_width - section.left_margin - section.right_margin) / 914400
    available_height = (section.page_height - section.top_margin - section.bottom_margin) / 914400 - 1.0
    for figure_id, title, diagram_type, image_path in pages:
        if not image_path.is_file():
            raise FileNotFoundError(image_path)
        page_break = doc.add_paragraph()
        page_break.add_run().add_break(WD_BREAK.PAGE)
        add_before(e_xml, page_break)

        heading = doc.add_paragraph(f"Figure {figure_id}. {title}", style="Heading 3")
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_after = Pt(6)
        add_before(e_xml, heading)

        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(4)
        caption.add_run(f"Diagram type: {diagram_type}.").italic = True
        add_before(e_xml, caption)

        with Image.open(image_path) as image:
            width_px, height_px = image.size
        aspect = width_px / height_px
        width_in = min(available_width, available_height * aspect)
        height_in = width_in / aspect
        image_paragraph = doc.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_after = Pt(0)
        image_paragraph.add_run().add_picture(str(image_path), width=Inches(width_in), height=Inches(height_in))
        add_before(e_xml, image_paragraph)

    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    add_before(e_xml, page_break)


def main() -> None:
    PACKAGE_OUT.mkdir(parents=True, exist_ok=True)
    for stem in NAMES.values():
        shutil.copy2(PACKAGE_IN / f"{stem}.docx", PACKAGE_OUT / f"{stem}.docx")
        shutil.copy2(PACKAGE_IN / f"{stem}.pdf", PACKAGE_OUT / f"{stem}.pdf")

    report_path = PACKAGE_OUT / f"{NAMES['report']}.docx"
    doc = Document(report_path)
    replace_report_references(doc)
    add_gallery(doc)
    doc.core_properties.title = "Sketch2Life Review 2 Reports 1–5 and Project Tracking — Rev12"
    doc.core_properties.subject = "Consolidated report with the current state and Rev12 editable Draw.io gallery"
    doc.save(report_path)
    print(f"Prepared four DOCX/PDF dossier pairs in {PACKAGE_OUT}; refresh the report PDF from Word before QA.")


if __name__ == "__main__":
    main()
