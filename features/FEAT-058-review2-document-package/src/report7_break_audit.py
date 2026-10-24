from docx import Document
from docx.oxml.ns import qn
p=r'.runtime/fe058_docx_work/Report7-reference.docx'; d=Document(p)
for i,p in enumerate(d.paragraphs):
 t=p.text
 if t or p._p.xpath('.//w:br') or p._p.xpath('./w:pPr/w:pageBreakBefore'):
  print(i,repr(t[:100]),'breaks',len(p._p.xpath('.//w:br')),'before',len(p._p.xpath('./w:pPr/w:pageBreakBefore')))
for si,s in enumerate(d.sections):
 print('section',si,'header', [p.text for p in s.header.paragraphs], 'footer', [p.text for p in s.footer.paragraphs])
