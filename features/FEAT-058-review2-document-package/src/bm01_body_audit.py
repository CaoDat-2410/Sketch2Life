from docx import Document
for p in [r'.runtime/fe058_docx_work/BM01-reference.docx']:
 d=Document(p); print('FILE',p)
 for i,x in enumerate(d.paragraphs): print('P',i,x.style.name,repr(x.text))
 for ti,t in enumerate(d.tables):
  print('TABLE',ti,'rows',len(t.rows),'cols',len(t.columns))
  for ri,row in enumerate(t.rows): print(' R',ri,' | '.join(' / '.join(c.text.replace('\n','; ') for c in row.cells)))
