"""Render and record page-level QA for the four local Rev12 PDFs."""

from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from PIL import Image, ImageDraw, ImageFont


FEATURE = Path(__file__).resolve().parents[1]
PACKAGE = FEATURE / "artifacts" / "submission_r2_rev12"
RENDERS = FEATURE / "evidence" / "render" / "rev12" / "documents"
METRICS = FEATURE / "evidence" / "metrics" / "rev12-package-qa.json"
DOCUMENTS = [
    "SWP490_FA26SE225_R2_DesignBrief",
    "SWP490_FA26SE225_R2_Registration",
    "SWP490_FA26SE225_R2_Report1-5_Tracking",
    "SWP490_FA26SE225_R2_SelfReview",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def render_document(stem: str) -> dict:
    docx_path = PACKAGE / f"{stem}.docx"
    pdf_path = PACKAGE / f"{stem}.pdf"
    if not docx_path.is_file() or not pdf_path.is_file():
        raise FileNotFoundError(f"Missing required output for {stem}")
    with zipfile.ZipFile(docx_path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise ValueError(f"DOCX archive has a corrupt member: {bad_member}")
    parsed_doc = Document(docx_path)
    pdf = pdfium.PdfDocument(pdf_path)
    out_dir = RENDERS / stem
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("page-*.png"):
        old.unlink()
    thumbnails: list[Image.Image] = []
    page_records = []
    for index in range(len(pdf)):
        page = pdf[index]
        text = page.get_textpage().get_text_range().replace("\x00", "")
        image = page.render(scale=1.55).to_pil().convert("RGB")
        image.save(out_dir / f"page-{index + 1:03d}.png", optimize=True)
        thumbnail = image.copy()
        thumbnail.thumbnail((510, 710))
        canvas = Image.new("RGB", (530, 750), "#e7ebf1")
        canvas.paste(thumbnail, ((530 - thumbnail.width) // 2, 10))
        ImageDraw.Draw(canvas).text((12, 726), f"{stem} | page {index + 1}", fill="#17212e")
        thumbnails.append(canvas)
        page_records.append({
            "page": index + 1,
            "text_chars": len(text.strip()),
            "has_figure_caption": "Figure D" in text,
            "has_diagram_type_caption": "Diagram type:" in text,
            "has_source_caption": "Source:" in text and "Figure D" in text,
            "text_preview": (
                "[redacted: local-only Registration+BM03 content]"
                if stem.endswith("Registration")
                else " ".join(text.split())[:220]
            ),
        })
    for start in range(0, len(thumbnails), 4):
        sheet = Image.new("RGB", (1060, 1500), "white")
        for offset, image in enumerate(thumbnails[start : start + 4]):
            sheet.paste(image, ((offset % 2) * 530, (offset // 2) * 750))
        sheet.save(out_dir / f"contact-{start // 4 + 1:02d}.png", optimize=True)
    for start in range(0, len(pdf), 20):
        overview = Image.new("RGB", (880, 1575), "white")
        for offset, page_index in enumerate(range(start, min(start + 20, len(pdf)))):
            image = Image.open(out_dir / f"page-{page_index + 1:03d}.png").convert("RGB")
            image.thumbnail((210, 292))
            canvas = Image.new("RGB", (220, 315), "#e7ebf1")
            canvas.paste(image, ((220 - image.width) // 2, 4))
            ImageDraw.Draw(canvas).text((5, 298), f"Page {page_index + 1}", fill="#17212e")
            overview.paste(canvas, ((offset % 4) * 220, (offset // 4) * 315))
        overview.save(out_dir / f"overview-{start // 20 + 1:02d}.png", optimize=True)
    return {
        "name": stem,
        "docx_bytes": docx_path.stat().st_size,
        "docx_sha256": sha256(docx_path),
        "docx_zip_integrity": "pass",
        "paragraph_count": len(parsed_doc.paragraphs),
        "table_count": len(parsed_doc.tables),
        "pdf_bytes": pdf_path.stat().st_size,
        "pdf_sha256": sha256(pdf_path),
        "pdf_page_count": len(pdf),
        "pages_with_no_extractable_text": [item["page"] for item in page_records if item["text_chars"] == 0],
        "figure_captions": sum(item["has_figure_caption"] for item in page_records),
        "typed_figure_captions": sum(item["has_diagram_type_caption"] for item in page_records),
        "figure_pages_with_source_caption": [item["page"] for item in page_records if item["has_source_caption"]],
        "page_records": page_records,
    }


def main() -> None:
    RENDERS.mkdir(parents=True, exist_ok=True)
    records = [render_document(stem) for stem in DOCUMENTS]
    report_record = next(item for item in records if item["name"].endswith("Tracking"))
    result = {
        "revision": "Rev12",
        "dossier_count": len(records),
        "docx_pdf_pairs": len(records),
        "diagram_count_in_report": 46,
        "documents": records,
    }
    METRICS.parent.mkdir(parents=True, exist_ok=True)
    METRICS.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "dossiers": len(records),
        "page_counts": {item["name"]: item["pdf_page_count"] for item in records},
        "figure_captions": report_record["figure_captions"],
        "typed_figure_captions": report_record["typed_figure_captions"],
        "blank_pages": {item["name"]: item["pages_with_no_extractable_text"] for item in records},
        "source_caption_pages": {item["name"]: item["figure_pages_with_source_caption"] for item in records},
        "metrics": str(METRICS),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
