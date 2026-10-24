from __future__ import annotations

import math
import re
import shutil
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Cm
from report2_data import SCHEDULE_ROWS


ROOT = Path(__file__).resolve().parents[3]
FEATURE = ROOT / "features" / "FEAT-058-review2-document-package"
ARTIFACTS = FEATURE / "artifacts"
DIAGRAM_SOURCES = ARTIFACTS / "diagrams"
WORK = ROOT / ".runtime" / "fe058_docx_work"
PICTURES = WORK / "generated-diagrams"
REPORT_TEMPLATE = WORK / "Report7-reference.docx"
BM_TEMPLATE = WORK / "BM01-reference.docx"
REPORT_OUT = ARTIFACTS / "SWP490_FA26SE225_R2_Reports1-5_Tracking.docx"
BM_OUT = ARTIFACTS / "SWP490_FA26SE225_R2_DesignBrief.docx"
SRS_PATH = ROOT / "features" / "FEAT-039-collaborative-learning-scope" / "artifacts" / "Collaborative_Learning_SRS_v3.1.md"

NAVY = "14243D"
NAVY_2 = "203755"
TEAL = "197C78"
ORANGE = "E68B35"
PALE = "EAF2F4"
PALE_ORANGE = "FBF1E7"
PALE_BLUE = "EEF2F8"
GRAY = "566170"
BLACK = "111111"
WHITE = "FFFFFF"
FONT_REG = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD = r"C:\Windows\Fonts\arialbd.ttf"


def font(size: int, bold: bool = False):
    path = FONT_BOLD if bold else FONT_REG
    try:
        return ImageFont.truetype(path, size=size)
    except OSError:
        return ImageFont.load_default()


def write_mmd(name: str, source: str) -> Path:
    DIAGRAM_SOURCES.mkdir(parents=True, exist_ok=True)
    path = DIAGRAM_SOURCES / f"{name}.mmd"
    path.write_text(source.strip() + "\n", encoding="utf-8")
    return path


def wrap_lines(draw: ImageDraw.ImageDraw, text: str, fnt, max_width: int):
    words = str(text).split()
    lines, line = [], ""
    for word in words:
        trial = word if not line else line + " " + word
        if draw.textbbox((0, 0), trial, font=fnt)[2] <= max_width:
            line = trial
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines or [""]


def label_box(draw, rect, label, fill=PALE_BLUE, outline=NAVY_2, fnt=None, color=BLACK, radius=22):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=radius, fill="#" + fill, outline="#" + outline, width=3)
    fnt = fnt or font(25, True)
    lines = wrap_lines(draw, label, fnt, x2 - x1 - 34)
    height = sum(draw.textbbox((0, 0), s, font=fnt)[3] for s in lines) + 6 * (len(lines) - 1)
    y = y1 + (y2 - y1 - height) / 2
    for line in lines:
        box = draw.textbbox((0, 0), line, font=fnt)
        draw.text((x1 + (x2 - x1 - (box[2] - box[0])) / 2, y), line, font=fnt, fill="#" + color)
        y += box[3] + 6


def arrow(draw, start, end, label="", color=TEAL, width=4, dash=False, label_width=210, font_size=18):
    x1, y1 = start
    x2, y2 = end
    if dash:
        steps = max(1, int(math.hypot(x2 - x1, y2 - y1) / 18))
        for i in range(0, steps, 2):
            a, b = i / steps, min((i + 1) / steps, 1)
            draw.line((x1 + (x2 - x1) * a, y1 + (y2 - y1) * a, x1 + (x2 - x1) * b, y1 + (y2 - y1) * b), fill="#" + color, width=width)
    else:
        draw.line((x1, y1, x2, y2), fill="#" + color, width=width)
    theta = math.atan2(y2 - y1, x2 - x1)
    head = 16
    p1 = (x2, y2)
    p2 = (x2 - head * math.cos(theta - 0.55), y2 - head * math.sin(theta - 0.55))
    p3 = (x2 - head * math.cos(theta + 0.55), y2 - head * math.sin(theta + 0.55))
    draw.polygon([p1, p2, p3], fill="#" + color)
    if label:
        fnt = font(font_size)
        bx = (x1 + x2) / 2
        by = (y1 + y2) / 2 - 17
        half = label_width / 2
        draw.rounded_rectangle((bx - half, by - 4, bx + half, by + 24), radius=8, fill="#FFFFFF")
        draw.text((bx - half + 5, by), label, font=fnt, fill="#" + color)


def base_canvas(title: str, subtitle: str, width=1700, height=980):
    im = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, width, 128), fill="#" + NAVY)
    d.text((60, 25), title, font=font(40, True), fill="white")
    d.text((62, 79), subtitle, font=font(20), fill="#D9E3EF")
    d.line((60, height - 62, width - 60, height - 62), fill="#D9E3EF", width=2)
    d.text((60, height - 48), "Sketch2Life | Review 2 design view | Solid = present source or confirmed behavior; dashed = proposed/TBD", font=font(17), fill="#" + GRAY)
    return im, d


def flow_image(filename, title, subtitle, nodes, edges, width=1700, height=980):
    im, d = base_canvas(title, subtitle, width, height)
    for e in edges:
        a, b, label = e[:3]
        dash = bool(e[3]) if len(e) > 3 else False
        n1, n2 = nodes[a], nodes[b]
        x1 = (n1[0] + n1[2]) / 2
        y1 = (n1[1] + n1[3]) / 2
        x2 = (n2[0] + n2[2]) / 2
        y2 = (n2[1] + n2[3]) / 2
        dx, dy = x2 - x1, y2 - y1
        scale = max(abs(dx) / max(1, (n1[2] - n1[0]) / 2), abs(dy) / max(1, (n1[3] - n1[1]) / 2), 1)
        start = (x1 + dx / scale, y1 + dy / scale)
        end = (x2 - dx / scale, y2 - dy / scale)
        arrow(d, start, end, label, dash=dash)
    for key, spec in nodes.items():
        x1, y1, x2, y2, label, *style = spec
        fill = style[0] if style else PALE_BLUE
        label_box(d, (x1, y1, x2, y2), label, fill=fill, fnt=font(23, True))
    path = PICTURES / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path, quality=95)
    return path


def swimlane_image(filename, title, subtitle, lanes, activities, edges, width=1700, height=1080):
    im, d = base_canvas(title, subtitle, width, height)
    left, top, right, bottom = 220, 170, width - 60, height - 95
    lane_w = (right - left) // len(lanes)
    for i, name in enumerate(lanes):
        x1 = left + i * lane_w
        x2 = x1 + lane_w
        fill = "F7F9FC" if i % 2 == 0 else "F1F5F7"
        d.rectangle((x1, top, x2, bottom), fill="#" + fill, outline="#D8E0E8", width=2)
        d.rectangle((x1, top, x2, top + 54), fill="#" + NAVY_2)
        d.text((x1 + 20, top + 13), name, font=font(21, True), fill="white")
    for key, x, y, label, fill in activities:
        w = lane_w - 64
        label_box(d, (x, y, x + w, y + 88), label, fill=fill, fnt=font(20, True))
    for edge in edges:
        a, b, label = edge[:3]
        dash = bool(edge[3]) if len(edge) > 3 else False
        aa = next(x for x in activities if x[0] == a)
        bb = next(x for x in activities if x[0] == b)
        ar = (aa[1] + lane_w - 64, aa[2] + 44)
        br = (bb[1], bb[2] + 44)
        if br[0] <= ar[0]:
            ar = (aa[1] + (lane_w - 64) // 2, aa[2] + 88)
            br = (bb[1] + (lane_w - 64) // 2, bb[2])
        arrow(d, ar, br, label, dash=dash)
    path = PICTURES / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def sequence_image(filename, title, subtitle, actors, messages):
    width, height = 1700, max(1260, 300 + len(messages) * 64)
    im, d = base_canvas(title, subtitle, width, height)
    n = len(actors)
    xs = [130 + i * (1440 // max(1, n - 1)) for i in range(n)]
    top = 200
    for x, label in zip(xs, actors):
        label_box(d, (x - 110, top, x + 110, top + 74), label, fill=PALE_BLUE, fnt=font(19, True))
        d.line((x, top + 74, x, height - 115), fill="#C6D0DC", width=2)
    y = top + 125
    for idx, (src, dst, txt, dashed) in enumerate(messages, 1):
        x1, x2 = xs[src], xs[dst]
        if x1 == x2:
            arrow(d, (x1, y), (x1 + 80, y + 34), txt, dash=dashed)
        else:
            arrow(d, (x1, y), (x2, y), f"{idx}. {txt}", dash=dashed, label_width=250, font_size=16)
        y += 64
    path = PICTURES / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def activity_flow_image(filename, title, subtitle):
    width, height = 1700, 1040
    im, d = base_canvas(title, subtitle, width, height)
    boxes = {
        "a": (70, 205, 500, 325, "1  Child joins with QR / session code", PALE_BLUE),
        "b": (635, 205, 1065, 325, "2  Teacher verifies roster and admits", PALE),
        "c": (1200, 205, 1630, 325, "3  Backend stores session membership", PALE_ORANGE),
        "d": (1200, 435, 1630, 555, "4  Child draws and requests support", PALE_BLUE),
        "e": (635, 435, 1065, 555, "5  Backend validates, saves and prepares result", PALE_ORANGE),
        "f": (70, 435, 500, 555, "6  Teacher reviews exact AI suggestion", PALE),
        "g": (70, 665, 500, 785, "7  Child views, hides or declines", PALE_BLUE),
        "h": (635, 665, 1065, 785, "8  Teacher approves media and guides activity", PALE),
        "i": (1200, 665, 1630, 785, "9  Backend saves outcome and portfolio", PALE_ORANGE),
    }
    for a,b in [("a","b"),("b","c"),("c","d"),("d","e"),("e","f"),("f","g"),("g","h"),("h","i")]:
        x1,y1,x2,y2,*_=boxes[a]; X1,Y1,X2,Y2,*_=boxes[b]
        if y1 == Y1:
            start=(x2 if x2 < X1 else x1, (y1+y2)/2)
            end=(X1 if x2 < X1 else X2, (Y1+Y2)/2)
        elif x1 == X1:
            start=((x1+x2)/2,y2); end=((X1+X2)/2,Y1)
        else:
            start=(x1,(y1+y2)/2); end=(X2,(Y1+Y2)/2)
        arrow(d,start,end,width=4)
    for rect in boxes.values():
        x1,y1,x2,y2,label,fill=rect
        label_box(d,(x1,y1,x2,y2),label,fill=fill,fnt=font(21,True))
    path=PICTURES/filename
    path.parent.mkdir(parents=True,exist_ok=True)
    im.save(path)
    return path


def state_image(filename, title, subtitle, states, transitions):
    width = 1700
    height = max(790, 310 + max(len(states) * 72, len(transitions) * 58))
    im, d = base_canvas(title, subtitle, width, height)
    d.text((78, 176), "States", font=font(21, True), fill="#" + NAVY)
    d.text((655, 176), "Transitions", font=font(21, True), fill="#" + NAVY)
    row_y = 210
    for i, (key, label, fill) in enumerate(states):
        y = row_y + i * 72
        label_box(d, (70, y, 590, y + 60), label, fill=fill, fnt=font(19, True), radius=14)
    for i, (a, b, label, dash) in enumerate(transitions):
        y = row_y + i * 58
        fill = "F8FAFC" if i % 2 == 0 else "EEF2F8"
        d.rounded_rectangle((650, y, 1628, y + 48), radius=10, fill="#" + fill, outline="#D8E0E8", width=2)
        marker = "- -" if dash else "→"
        transition_text = f"{a}  {marker}  {b}   |   {label}"
        d.text((672, y + 12), transition_text, font=font(19), fill="#" + BLACK)
    path = PICTURES / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def make_diagrams():
    PICTURES.mkdir(parents=True, exist_ok=True)
    outputs = {}
    outputs["D01"] = flow_image("D01_context.png", "System context", "Confirmed classroom actors and external boundaries around the application", {
        "child": (60, 340, 320, 470, "Child\nAndroid app", PALE_BLUE),
        "teacher": (60, 550, 320, 680, "Teacher\nDesktop", PALE_BLUE),
        "admin": (60, 760, 320, 890, "Super Admin\nConsole", PALE_BLUE),
        "api": (610, 450, 970, 650, "Sketch2Life\nFastAPI application", PALE),
        "auth": (1270, 180, 1600, 300, "Adult identity\nFirebase Auth only", PALE_ORANGE),
        "modal": (1270, 345, 1600, 465, "Modal AI hosting\nselected; config TBD", PALE_ORANGE),
        "data": (1270, 510, 1600, 650, "Database, queue,\nartifact storage\nconfiguration TBD", PALE_ORANGE),
        "class": (1270, 730, 1600, 850, "School classroom\nconsent process", PALE_BLUE),
    }, [("child","api","HTTPS"),("teacher","api","HTTPS"),("admin","api","scoped HTTPS"),("api","auth","adult identity"),("api","modal","backend only",True),("api","data","durable adapters",True),("class","teacher","teacher-led use")])
    write_mmd("D01_context", """flowchart LR
  Child[Child Android] -->|HTTPS| API[Sketch2Life FastAPI]
  Teacher[Teacher Desktop] -->|HTTPS| API
  Admin[Super Admin Console] -->|Scoped HTTPS| API
  API -->|Adult identity only| Firebase[Firebase Authentication]
  API -. Backend-only AI calls .-> Modal[Modal selected; deployment TBD]
  API -. Proposed durable adapters .-> Data[(Database, queue, artifact storage TBD)]
  School[School-mediated consent] --> Teacher
""")
    outputs["D02"] = flow_image("D02_use_cases.png", "Use case overview", "Target functions by actor; adult permission details remain subject to contract implementation", {
        "child": (60, 230, 360, 370, "Child\nAndroid", PALE_BLUE),
        "teacher": (60, 440, 360, 580, "Teacher\nDesktop", PALE_BLUE),
        "admin": (60, 650, 360, 790, "Super Admin", PALE_BLUE),
        "join": (540, 190, 900, 290, "Join admitted\nsession", PALE),
        "draw": (980, 190, 1340, 290, "Draw and recover\ncanvas", PALE),
        "ai": (540, 360, 900, 460, "Request / review\nAI sketch", PALE),
        "media": (980, 360, 1340, 460, "Approve video and\nactivity", PALE),
        "learn": (540, 540, 900, 640, "Reflect and record\nprogress", PALE),
        "manage": (980, 540, 1340, 640, "Manage class,\ncontent and policy", PALE),
        "privacy": (760, 740, 1120, 840, "Consent, export,\ndelete and audit", PALE_ORANGE),
    }, [("child","join",""),("child","draw",""),("child","ai",""),("child","learn",""),("teacher","join",""),("teacher","draw",""),("teacher","ai","review gate"),("teacher","media",""),("teacher","learn",""),("teacher","manage",""),("admin","manage",""),("admin","privacy","scoped")])
    write_mmd("D02_use_cases", """flowchart LR
  Child[Child] --> UC1[Join admitted session]
  Child --> UC2[Draw, recover, and respond to a reviewed suggestion]
  Child --> UC3[Preview, reflect, and join an off-screen activity]
  Teacher[Teacher] --> UC4[Manage class, groups, and session pace]
  Teacher --> UC5[Review AI sketch, knowledge, and video]
  Teacher --> UC6[Present artwork and record observations]
  Admin[Super Admin] --> UC7[Manage accounts, content, policy, and operations]
  Admin --> UC8[Handle scoped consent, audit, export, and delete]
""")
    outputs["D03"] = sequence_image("D03_sequence.png", "Sequence for join and durable drawing", "Target durable classroom path; production UI-to-API integration is not established", ["Child UI","Teacher UI","FastAPI","Domain / Ports","PostgreSQL","Worker / Modal"], [
        (1,2,"create session/code",False),(0,2,"submit join",False),(2,1,"show pending",False),(1,2,"admit child",False),(2,3,"authorize/validate",False),(3,4,"persist membership",True),(4,3,"commit result",False),(0,2,"draw op + contributor",False),(2,3,"scope/version/key check",False),(3,4,"persist accepted op",True),(4,3,"durable receipt/sequence",False),(3,5,"enqueue explicit AI",True),(5,3,"typed result",True),(2,0,"accepted status/receipt",False)
    ])
    write_mmd("D03_sequence", """sequenceDiagram
  participant C as Child UI
  participant T as Teacher UI
  participant API as FastAPI
  participant APP as Application and domain
  participant DB as PostgreSQL
  participant W as AI worker / Modal
  T->>API: Configure session and issue join code
  C->>API: Request admission
  T->>API: Admit participant
  API->>APP: Authorize and validate command
  APP->>DB: Persist membership and session state
  C->>API: Submit drawing operation with selected contributor context
  API->>APP: Check scope, epoch and idempotency
  APP->>DB: Commit accepted operation and receipt
  DB-->>APP: Durable sequence / receipt
  APP-->>C: Accepted result
  APP-->>W: Optional AI job through backend-owned port (proposed)
""")
    outputs["D04"] = flow_image("D04_screen_flow.png", "Screen flow by product surface", "Three distinct surfaces; screens are prototypes and fixture/service integration is incomplete", {
        "c1": (55, 190, 480, 300, "Child: Join / Lobby", PALE_BLUE),
        "c2": (55, 380, 480, 490, "Draw / AI suggestion", PALE_BLUE),
        "c3": (55, 570, 480, 680, "Gallery / Knowledge", PALE_BLUE),
        "c4": (55, 760, 480, 870, "Activity / Reflection", PALE_BLUE),
        "t1": (630, 190, 1055, 300, "Teacher: Dashboard / Class", PALE),
        "t2": (630, 380, 1055, 490, "Session setup / Lobby", PALE),
        "t3": (630, 570, 1055, 680, "Monitor / AI review", PALE),
        "t4": (630, 760, 1055, 870, "Gallery / Media / Records", PALE),
        "a1": (1220, 190, 1650, 300, "Admin: Overview / Accounts", PALE_ORANGE),
        "a2": (1220, 380, 1650, 490, "Content / AI policy", PALE_ORANGE),
        "a3": (1220, 570, 1650, 680, "Audit / Reports / Requests", PALE_ORANGE),
        "end": (1220, 760, 1650, 870, "Role and purpose checks\napply at every API action", PALE),
    }, [("c1","c2","admitted"),("c2","c3","approved only"),("c3","c4","teacher guided"),("t1","t2","configure"),("t2","t3","start / monitor"),("t3","t4","approve / present"),("a1","a2","manage"),("a2","a3","audit"),("t4","end","scoped access")])
    write_mmd("D04_screen_flow", """flowchart LR
  subgraph Child_Android
    C1[Join / Lobby] --> C2[Draw]
    C2 --> C3[Reviewed AI suggestion]
    C3 --> C4[Gallery / Knowledge / Video]
    C4 --> C5[Activity / Reflection]
  end
  subgraph Teacher_Desktop
    T1[Dashboard / Class] --> T2[Session setup / Lobby]
    T2 --> T3[Group monitoring / AI review]
    T3 --> T4[Gallery / Media / Assessment]
  end
  subgraph Super_Admin
    A1[Overview / Accounts] --> A2[Content / AI policy]
    A2 --> A3[Audit / Reports / Data requests]
  end
""")
    outputs["D05"] = flow_image("D05_class_model.png", "Classroom domain class view", "Logical target model; entities and multiplicities are proposed, not a migrated schema", {
        "school": (70, 190, 390, 300, "OrganizationScope\n(id, policyRef)", PALE_ORANGE),
        "class": (510, 190, 830, 300, "Class\n(id, schoolRef)", PALE),
        "session": (950, 190, 1270, 300, "ClassroomSession\n(lifecycle, version)", PALE),
        "group": (1350, 190, 1640, 300, "SessionGroup\n(progress)", PALE),
        "student": (70, 470, 390, 580, "StudentProfile\n(managed; no login)", PALE_BLUE),
        "member": (510, 470, 830, 580, "ChildParticipant\n(session admission)", PALE_BLUE),
        "canvas": (950, 470, 1270, 580, "CanvasDocument\n(scope / epoch)", PALE),
        "op": (1350, 470, 1640, 580, "CanvasOperation\n(principal + contributor)", PALE),
        "snapshot": (270, 735, 620, 845, "CanvasSnapshot\n(hash / provenance)", PALE),
        "job": (720, 735, 1070, 845, "AI Job / Proposal\n(versioned result)", PALE_ORANGE),
        "review": (1170, 735, 1520, 845, "TeacherReview\n(content version + decision)", PALE),
    }, [("school","class","1 to many",True),("class","session","1 to many",True),("session","group","1 to many",True),("student","member","logical link",True),("session","member","1 to many",True),("member","group","membership",True),("session","canvas","1 to many",True),("canvas","op","1 to many",True),("canvas","snapshot","checkpoint",True),("snapshot","job","input lineage",True),("job","review","human gate",True)])
    write_mmd("D05_class_model", """classDiagram
  class ClassroomSession { id; classRef; lifecycle; version }
  class SessionGroup { id; sessionRef; progress }
  class ChildParticipant { id; studentRef; admissionState }
  class CanvasDocument { id; sessionRef; scope; epoch }
  class CanvasOperation { id; documentRef; principalRef; contributorRef; serverSeq }
  class CanvasSnapshot { id; documentRef; hash; provenance }
  class AIJob { id; inputVersion; state; resultDisposition }
  class TeacherReview { id; contentVersion; reviewer; decision }
  ClassroomSession "1" --> "many" SessionGroup
  ClassroomSession "1" --> "many" ChildParticipant
  ClassroomSession "1" --> "many" CanvasDocument
  CanvasDocument "1" --> "many" CanvasOperation
  CanvasDocument "1" --> "many" CanvasSnapshot
  CanvasSnapshot --> AIJob : logical input reference
  AIJob --> TeacherReview : reviewed result
""")
    outputs["D06"] = activity_flow_image("D06_activity.png", "Whole-system classroom activity", "Actor-ordered flow; Teacher controls admission, review, media, activity and assessment gates")
    write_mmd("D06_activity", """flowchart LR
  subgraph Child
    C1[Join session] --> C2[Draw and collaborate]
    C2 --> C3[Request optional AI support]
    C3 --> C4[View / hide / decline reviewed suggestion]
    C4 --> C5[Reflect and join off-screen activity]
  end
  subgraph Teacher
    T1[Check roster and admit] --> T2[Control groups and session pace]
    T2 --> T3[Review each suggestion]
    T3 --> T4[Approve exact knowledge or video]
    T4 --> T5[Record observation and session outcome]
  end
  subgraph Backend
    B1[Authorize and persist accepted operation] --> B2[Run bounded AI job through provider port]
    B2 --> B3[Return typed result to review queue]
    B3 --> B4[Serve approved version only]
    B4 --> B5[Close and persist records]
  end
""")
    outputs["D07"] = flow_image("D07_erd.png", "Current persistence ERD excerpt", "Solid arrows are physical foreign keys in migrations; cross-module references are opaque IDs", {
        "rev": (50, 190, 430, 330, "core_aggregate_revisions\nPK aggregate_id", PALE),
        "op": (520, 190, 900, 330, "core_canvas_operations\nPK operation_id\nFK canvas_id", PALE),
        "receipt": (1000, 190, 1380, 330, "core_idempotency_receipts\nPK org/scope/key\nFK operation_id", PALE),
        "check": (50, 470, 430, 610, "core_canvas_checkpoints\nFK canvas_id\nFK through_operation_id", PALE),
        "outbox": (520, 470, 900, 610, "core_outbox_events\nFK aggregate_id", PALE),
        "consumer": (1000, 470, 1380, 610, "core_outbox_consumer_receipts\nFK event_id", PALE),
        "job": (50, 730, 430, 870, "ai_jobs.job\nPK job_id\nopaque session_id", PALE_ORANGE),
        "attempt": (520, 730, 900, 870, "ai_jobs.attempt\nPK attempt_id\nFK job_id", PALE_ORANGE),
        "completion": (1000, 730, 1380, 870, "ai_jobs.completion_receipt\nFK job_id + attempt_id", PALE_ORANGE),
        "logical": (1440, 475, 1650, 620, "Session / class / child\nIDs may be logical refs\nNo cross-schema FK", PALE_BLUE),
    }, [("rev","op","canvas_id FK"),("op","receipt","operation_id FK"),("rev","check","canvas_id FK"),("op","check","through_operation_id FK"),("rev","outbox","aggregate_id FK"),("outbox","consumer","event_id FK"),("job","attempt","job_id FK"),("job","completion","job_id FK"),("attempt","completion","attempt_id FK"),("job","logical","opaque session_id",True)])
    write_mmd("D07_erd", """erDiagram
  CORE_AGGREGATE_REVISIONS ||--o{ CORE_CANVAS_OPERATIONS : canvas_id_fk
  CORE_CANVAS_OPERATIONS ||--o| CORE_IDEMPOTENCY_RECEIPTS : operation_id_fk
  CORE_AGGREGATE_REVISIONS ||--o{ CORE_CANVAS_CHECKPOINTS : canvas_id_fk
  CORE_CANVAS_OPERATIONS ||--o{ CORE_CANVAS_CHECKPOINTS : through_operation_id_fk
  CORE_AGGREGATE_REVISIONS ||--o{ CORE_OUTBOX_EVENTS : aggregate_id_fk
  CORE_OUTBOX_EVENTS ||--o{ CORE_OUTBOX_CONSUMER_RECEIPTS : event_id_fk
  AI_JOBS_JOB ||--o{ AI_JOBS_ATTEMPT : job_id_fk
  AI_JOBS_JOB ||--o{ AI_JOBS_COMMAND_RECEIPT : job_id_fk
  AI_JOBS_JOB ||--o{ AI_JOBS_COMPLETION_RECEIPT : job_id_fk
  AI_JOBS_ATTEMPT ||--o{ AI_JOBS_COMPLETION_RECEIPT : attempt_id_fk
  AI_JOBS_JOB }o..o{ CLASS_SESSION : opaque_session_id_no_fk
""")
    outputs["D08"] = flow_image("D08_architecture.png", "Target system architecture", "Durable classroom runtime is the target entrypoint; current UI clients remain fixture-backed", {
        "child": (40, 210, 380, 340, "Child Android\nExpo / React Native", PALE_BLUE),
        "teacher": (40, 410, 380, 540, "Teacher Desktop\nReact / Vite prototype", PALE_BLUE),
        "admin": (40, 610, 380, 740, "Super Admin\nReact / Vite prototype", PALE_BLUE),
        "api": (520, 390, 870, 570, "FastAPI\nHTTP interfaces", PALE),
        "app": (1010, 390, 1360, 570, "Application use cases\nports + domain rules", PALE),
        "db": (1490, 180, 1650, 300, "PostgreSQL\ncurrent modules", PALE_ORANGE),
        "queue": (1490, 350, 1650, 470, "Workers / queue\nconfig TBD", PALE_ORANGE),
        "modal": (1490, 540, 1650, 660, "Modal AI\nselected / not deployed", PALE_ORANGE),
        "demo": (520, 750, 870, 870, "Separate demo\nentrypoints", "FBF1E7"),
    }, [("child","api","current UI stub",True),("teacher","api","current UI fixture",True),("admin","api","current UI fixture",True),("api","app","router composition"),("app","db","ports / adapters",True),("app","queue","durable jobs proposed",True),("queue","modal","backend-only provider",True),("demo","api","not main runtime",True)])
    write_mmd("D08_architecture", """flowchart LR
  Child[Child Android UI] -. stub client today .-> API[FastAPI entrypoint]
  Teacher[Teacher Desktop UI] -. fixture service today .-> API
  Admin[Super Admin UI] -. fixture service today .-> API
  API --> App[Application use cases and ports]
  App --> Domain[Framework-independent domain rules]
  App -. proposed .-> DB[(Durable persistence adapters)]
  App -. proposed .-> Queue[Workers and queue]
  Queue -. selected provider; config TBD .-> Modal[Modal AI]
  Demo[Separate legacy/demo entrypoints] -. isolated from target entrypoint .-> API
""")
    outputs["D09"] = flow_image("D09_package.png", "Package and dependency view", "Dependencies point inward; infrastructure implements application-owned interfaces", {
        "surface": (50, 230, 420, 390, "apps/child-ui\napps/teacher-ui\napps/admin-ui", PALE_BLUE),
        "http": (500, 230, 870, 390, "interfaces/http\nrouters / composition", PALE),
        "app": (950, 230, 1320, 390, "application\nuse cases / ports", PALE),
        "domain": (1400, 230, 1650, 390, "domain\nentities / policies", PALE),
        "infra": (950, 520, 1320, 680, "infrastructure\npersistence / providers\nworkers / storage", PALE_ORANGE),
        "contracts": (500, 520, 870, 680, "versioned contracts\ncommands / events", PALE),
        "demo": (50, 740, 420, 860, "demo entrypoints\nseparate composition", PALE_ORANGE),
    }, [("surface","http","HTTP / app boundary",True),("http","app","calls use cases"),("app","domain","depends inward"),("infra","app","implements ports",True),("contracts","app","versioned boundary"),("demo","http","separate router set",True)])
    write_mmd("D09_package", """flowchart LR
  UI[apps / Child, Teacher, Admin] --> HTTP[interfaces / HTTP and composition]
  HTTP --> APP[application / use cases and ports]
  APP --> DOMAIN[domain / entities and rules]
  INFRA[infrastructure / SQL, storage, workers, providers] -. implements .-> APP
  CONTRACTS[versioned cross-feature contracts] --> APP
  DEMO[separate demo composition] --> HTTP
""")
    outputs["D10"] = state_image("D10_state_session.png", "Session lifecycle", "Product state view; detailed persisted enum and close policy require contract approval", [
        ("draft","Draft",PALE_BLUE),("lobby","Lobby",PALE),("active","Active",PALE),("paused","Paused",PALE_ORANGE),("completed","Completed",PALE),("early","Ended early",PALE_ORANGE)], [
        ("draft","lobby","Teacher configures",False),("lobby","active","Teacher starts",False),("active","paused","Teacher pauses",False),("paused","active","Teacher resumes",False),("active","completed","Teacher closes",False),("active","early","Teacher ends",False),("paused","early","Teacher ends",False)])
    write_mmd("D10_state_session", """stateDiagram-v2
  [*] --> Draft
  Draft --> Lobby: teacher configures
  Lobby --> Active: explicit teacher start
  Active --> Paused: teacher pause
  Paused --> Active: teacher resume
  Active --> Completed: teacher close
  Active --> EndedEarly: teacher ends early
  Paused --> EndedEarly: teacher ends early
  Completed --> [*]
  EndedEarly --> [*]
""")
    outputs["D11"] = state_image("D11_state_membership.png", "Participant admission and membership", "Admission and group membership are separate from child login identity", [
        ("pending","Pending\nrequest",PALE_BLUE),("admitted","Admitted",PALE),("assigned","Group\nassigned",PALE),("active","Active\nparticipant",PALE),("moved","Moved / left",PALE_ORANGE),("revoked","Revoked",PALE_ORANGE)], [
        ("pending","admitted","teacher admits",False),("admitted","assigned","teacher assigns",False),("assigned","active","session active",False),("active","moved","move / leave",False),("moved","active","teacher re-admits",True),("active","revoked","teacher / policy",True)])
    write_mmd("D11_state_membership", """stateDiagram-v2
  [*] --> PendingRequest
  PendingRequest --> Admitted: teacher verifies roster
  Admitted --> GroupAssigned: teacher assigns group
  GroupAssigned --> Active: session starts
  Active --> MovedOrLeft: teacher move or child leaves
  Active --> Revoked: permission / enrollment / consent policy
  MovedOrLeft --> Active: re-admitted by teacher
  Revoked --> [*]
""")
    outputs["D12"] = state_image("D12_state_ai_job.png", "AI job lifecycle", "States match the current ai_jobs migration; disposition is a separate accepted / stale / rejected value", [
        ("queued","queued",PALE_BLUE),("running","running",PALE),("retry","failed\nretryable",PALE_ORANGE),("succeeded","succeeded",PALE),("exhausted","failed\nexhausted",PALE_ORANGE),("unsafe","blocked\nunsafe",PALE_ORANGE),("cancelled","cancelled",PALE_ORANGE)], [
        ("queued","running","worker claim",False),("running","succeeded","typed completion",False),("running","retry","transient failure",False),("retry","running","eligible retry",True),("retry","exhausted","budget exhausted",False),("running","unsafe","moderation",True),("queued","cancelled","cancel request",True),("running","cancelled","cancel request",True)])
    write_mmd("D12_state_ai_job", """stateDiagram-v2
  [*] --> queued
  queued --> running: worker claim
  running --> succeeded: completed
  running --> failed_retryable: transient failure
  failed_retryable --> running: retry eligible
  failed_retryable --> failed_exhausted: retry budget exhausted
  running --> blocked_unsafe: safety policy
  queued --> cancelled: cancellation
  running --> cancelled: cancellation
  note right of succeeded: result disposition is accepted, stale, or rejected
""")
    outputs["D13"] = state_image("D13_state_video.png", "Knowledge and video review", "Teacher approves the exact version; retry / skip / end after exhausted failure is owner-confirmed", [
        ("draft","Draft /\nselected",PALE_BLUE),("generating","Generating",PALE),("ready","Ready for\nreview",PALE),("approved","Approved\nexact version",PALE),("playing","Playing",PALE),("failed","Failed",PALE_ORANGE),("skip","Skipped by\nTeacher",PALE_ORANGE),("ended","Ended early",PALE_ORANGE)], [
        ("draft","generating","request generation",True),("generating","ready","completed",False),("ready","approved","teacher approves",False),("approved","playing","session playback",False),("generating","failed","exhausted retry",False),("failed","generating","teacher retry",False),("failed","skip","teacher skip",False),("failed","ended","teacher ends session",False)])
    write_mmd("D13_state_video", """stateDiagram-v2
  [*] --> Draft: proposed
  Draft --> Generating: teacher requests
  Generating --> ReadyForReview: complete
  ReadyForReview --> Approved: teacher approves exact version
  Approved --> Playing: teacher starts playback
  Generating --> Failed: retries exhausted
  Failed --> Generating: teacher retries
  Failed --> SkippedByTeacher: teacher explicit skip
  Failed --> EndedEarly: teacher ends session
  note right of ReadyForReview: approval is version-specific
""")
    return outputs


def set_run(run, size=10.5, bold=None, color=BLACK, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def no_split_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_cell(cell, text, *, header=False, size=9.3, fill=None, color=BLACK, bold=False, align=None):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    if fill:
        set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    set_run(r, size=size, bold=bold or header, color=color if not header else WHITE)


def set_cell_content(cell, text, *, size=9.2, bold=False, color=BLACK):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for idx, line in enumerate(str(text).split("\n")):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        set_run(r, size=size, bold=bold, color=color)


def style_table(table, headers, rows, widths=None, *, size=9.2, header_fill=NAVY, first_col_fill=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if len(table.rows) < len(rows) + 1:
        for _ in range(len(rows) + 1 - len(table.rows)):
            table.add_row()
    while len(table.rows) > len(rows) + 1:
        table._tbl.remove(table.rows[-1]._tr)
    for j, h in enumerate(headers):
        set_cell(table.rows[0].cells[j], h, header=True, size=size, fill=header_fill)
    set_repeat_table_header(table.rows[0])
    no_split_row(table.rows[0])
    for i, row in enumerate(rows, 1):
        for j, value in enumerate(row):
            fill = None
            if first_col_fill and j == 0:
                fill = first_col_fill
            elif i % 2 == 0:
                fill = "F6F7F8"
            set_cell(table.rows[i].cells[j], value, size=size, fill=fill, color=BLACK, bold=bool(first_col_fill and j == 0))
        no_split_row(table.rows[i])
    if widths:
        # Keep the grid, header row, and data rows on the same fixed layout.
        # Previously the first-row cells retained the wider default grid and
        # pushed tables beyond the page's left/right content bounds.
        widths = [float(width) for width in widths]
        max_width = 6.0  # report body is 6.268 in wide
        total = sum(widths)
        if total > max_width:
            widths = [width * max_width / total for width in widths]
        twips = [round(width * 1440) for width in widths]
        max_twips = round(max_width * 1440)
        if sum(twips) > max_twips:
            twips[-1] -= sum(twips) - max_twips
        for grid_col, width_twips in zip(table._tbl.tblGrid.gridCol_lst, twips):
            grid_col.set(qn("w:w"), str(width_twips))
        for row in table.rows:
            if len(row.cells) == len(twips):
                for cell, width_twips in zip(row.cells, twips):
                    cell.width = Inches(width_twips / 1440)
        tbl_pr = table._tbl.tblPr
        tbl_width = tbl_pr.find(qn("w:tblW"))
        if tbl_width is None:
            tbl_width = OxmlElement("w:tblW")
            tbl_pr.append(tbl_width)
        tbl_width.set(qn("w:type"), "dxa")
        tbl_width.set(qn("w:w"), str(sum(twips)))
        table_indent = tbl_pr.find(qn("w:tblInd"))
        if table_indent is not None:
            table_indent.set(qn("w:type"), "dxa")
            table_indent.set(qn("w:w"), "0")
    return table


def add_table(doc, headers, rows, widths=None, size=9.2, header_fill=NAVY, first_col_fill=None):
    table = doc.add_table(rows=1, cols=len(headers))
    return style_table(table, headers, rows, widths, size=size, header_fill=header_fill, first_col_fill=first_col_fill)


def add_para(doc, text, *, style=None, size=10.2, bold=False, color=BLACK, italic=False, after=4, before=0, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.03
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else (8 if level == 2 else 5))
    p.paragraph_format.space_after = Pt(5 if level == 1 else 3)
    for r in p.runs:
        set_run(r, size={1:16,2:12.5,3:11.5}.get(level,10.5), bold=True, color=NAVY if level == 1 else BLACK)
    return p


def add_bullet(doc, text, level=0, size=10.1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("•  " + text)
    set_run(r, size=size)
    return p


def add_number(doc, text, size=10.1):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run(r, size=size)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_run(r, size=9.2, italic=True, color=GRAY)
    return p


def add_figure(doc, image_path, caption, width=6.28):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(width))
    add_caption(doc, caption)


def add_page_field(p):
    r = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for e in (begin,instr,sep,txt,end): r._r.append(e)
    set_run(r, size=8, color=GRAY)


def add_toc_field(p):
    r = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "Update this table of contents in Word."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for e in (begin,instr,sep,text,end): r._r.append(e)


def configure_common(doc, *, margin=1.0, body_size=10.2):
    for sec in doc.sections:
        sec.top_margin = Inches(margin)
        sec.bottom_margin = Inches(margin)
        sec.left_margin = Inches(margin)
        sec.right_margin = Inches(margin)
        sec.header_distance = Inches(0.45)
        sec.footer_distance = Inches(0.45)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(body_size)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.03
    if "Title" not in [s.name for s in styles]:
        title_style = styles.add_style("Title", 1)
    else:
        title_style = styles["Title"]
    title_style.font.name = "Calibri"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    for key, size in (("Heading 1",16),("Heading 2",12.5),("Heading 3",11.5),("Heading 4",10.5)):
        try:
            st = styles[key]
        except KeyError:
            st = styles.add_style(key, 1)
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(NAVY if key == "Heading 1" else BLACK)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(10 if key == "Heading 1" else 6)
        st.paragraph_format.space_after = Pt(4)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.clear()
    rr = footer.add_run("Sketch2Life | Review 2  |  ")
    set_run(rr, size=8, color=GRAY)
    add_page_field(footer)
    return doc


def cell_replace(cell, text, size=10, bold=False, color=BLACK):
    cell.text = ""
    for i, line in enumerate(str(text).split("\n")):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        set_run(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, 90, 110, 90, 110)


def clear_after_cover(doc):
    body = doc._element.body
    children = list(body)
    keep = children[:9]
    sect = body.sectPr
    for element in children:
        if element is sect or element in keep:
            continue
        body.remove(element)


def fill_report_cover(doc):
    # Retain all institutional marks and template cover geometry.
    cover = doc.tables
    title_table = cover[1]
    cell_replace(title_table.cell(2,0), "Sketch2Life: Collaborative Creative Learning with AI for Children", size=18, bold=True, color=NAVY)
    title_table.cell(2,0).paragraphs[0].style = doc.styles["Title"]
    team = cover[2]
    cell_replace(team.cell(0,0), "Sketch2Life", size=12, bold=True, color=NAVY)
    cell_replace(team.cell(0,1), "Sketch2Life", size=12, bold=True, color=NAVY)
    members = [
        "Đỗ Cao Đạt  |  SE193713  |  Leader, Fullstack",
        "Phạm Tấn Hiếu  |  SE184252  |  BE2",
        "Vũ Chí Bảo  |  SE182733  |  FE",
        "Lý Anh Khoa  |  SE182682  |  BE1",
    ]
    cell_replace(team.cell(1,1), "\n".join(members), size=9.5)
    cell_replace(team.cell(2,1), "Lâm Hữu Khánh Phương", size=10)
    cell_replace(team.cell(3,1), "N/A", size=10)
    cell_replace(team.cell(4,1), "FA26SE225 (inferred)", size=10)
    # Fill the date paragraph in the last cover paragraph without touching the logo tables.
    for p in doc.paragraphs:
        if "Hanoi" in p.text or "month" in p.text:
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("Ho Chi Minh City, October 2026")
            set_run(r, size=10.5, color=NAVY)
    clear_after_cover(doc)


def add_report_frontmatter(doc):
    doc.add_page_break()
    add_heading(doc, "Table of Contents", 1)
    toc_p = doc.add_paragraph()
    toc_p.paragraph_format.space_after = Pt(10)
    add_toc_field(toc_p)
    add_heading(doc, "Acknowledgement", 2)
    add_para(doc, "The team acknowledges the academic supervision of Lâm Hữu Khánh Phương and the work contributed by the four registered members. This report presents the current source evidence and the target classroom specification for Review 2.", size=10)
    add_heading(doc, "Definitions and Acronyms", 2)
    rows = [
        ("AI","Artificial Intelligence"),("API","Application Programming Interface"),("ADR","Architecture Decision Record"),("BE","Backend"),("CI","Continuous Integration"),("ERD","Entity Relationship Diagram"),("FR","Functional Requirement"),("NFR","Non-Functional Requirement"),("QR","Quick Response code"),("SRS","Software Requirements Specification"),("SDD","Software Design Description"),("TBD","To be determined"),("UC","Use Case"),("UI","User Interface"),("UAT","User Acceptance Testing"),("Modal","Selected AI hosting platform; deployment configuration TBD")
    ]
    add_table(doc,["Acronym","Definition"],rows,[1.15,5.1],size=9.3,header_fill=NAVY)
    doc.add_page_break()


def add_section_intro(doc, text):
    p = add_para(doc, text, size=10.1)
    p.paragraph_format.keep_together = True


def add_req_catalog(doc):
    # Compact English requirements catalog aligned to canonical SRS v3.1 stable identifiers.
    reqs = [
    ("FR001","Enforce Child, Teacher and Super Admin permissions by resource scope","Confirmed"),("FR002","Admit children by QR/session code or managed profile after Teacher verification","Confirmed"),("FR003","Support one or multiple children using a shared device with session/group context","Confirmed"),("FR004","Let Teachers manage assigned classes, enrollments and child profiles","Confirmed"),("FR005","Verify adult identity and issue/revoke scoped participant capabilities","Proposed; Firebase Auth boundary"),("FR006","Apply purpose-specific consent before policy-limited collection or AI processing","Confirmed; detailed procedure TBD"),
    ("FR007","Configure topic, age range, drawing tools, groups and session preset","Confirmed"),("FR008","Create, edit and manage session groups and members","Confirmed"),("FR009","Track group progress independently; one ready group does not advance the class","Confirmed"),("FR010","Allow Teacher pause, resume, intervention, draft, stage advance and early end","Confirmed"),("FR011","Close a session with artwork, process, comments and portfolio outcome; expose save/recovery state","Confirmed"),("FR012","Move or release a child while retaining recorded artwork and contribution history","Confirmed"),("FR013","Return to a stage already completed after video/off-screen work","TBD"),("FR065","Offer automatic grouping in Phase 2 with Teacher review/override; criteria TBD","Confirmed capability; details TBD"),
    ("FR014","Provide personal and shared canvases with scoped drawing regions","Confirmed"),("FR015","Adapt tools by age band and let Teacher configure allowed tools","Confirmed"),("FR016","Show participants/help requests and protect, lock, unlock or restore canvas regions","Confirmed"),("FR017","Persist canvas state and recover after reconnect without losing accepted contributions","Confirmed"),("FR018","Provide author-aware operation history and scoped undo","Proposed"),("FR019","Keep AI sketch content separate from child strokes and provide an unoverlaid view","Mixed: separation confirmed; view proposed"),("FR020","Import image, sticker, text or starter template","TBD"),("FR021","Replay drawing process and support advanced layers for older children","Confirmed capability; details proposed"),("FR066","On shared device, record selected contributor for new strokes without relabeling old strokes","Owner-confirmed principle; contract proposed"),
    ("FR022","Analyze an authorized canvas snapshot with topic, age, request and Teacher feedback context","Confirmed"),("FR023","Return candidate objects and uncertainty; preserve child intent and Teacher correction","Confirmed"),("FR024","Accept child help request or Teacher-triggered assistance and allow cancellation","Confirmed"),("FR025","Offer more than one assistance level or presentation mode","Confirmed capability; taxonomy proposed"),("FR026","Require Teacher review of each pilot sketch before it reaches the child","Owner-confirmed"),("FR027","Let child view, hide or decline a reviewed sketch and continue drawing","Confirmed"),("FR028","Never replace child strokes or infer lack of knowledge from a pause","Confirmed"),("FR029","Block unsafe output and report typed failures to Teacher","Confirmed"),("FR030","Deduplicate and bound AI requests; reject results for stale snapshots","Proposed"),
    ("FR031","Let Teacher select, order, annotate and present artwork within class scope","Confirmed"),("FR032","Let child preview and explain permitted artwork","Confirmed"),("FR033","Version artwork sharing and enforce class/session scope","Proposed detail"),
    ("FR034","Create video or select a reviewed library item for class or child context","Confirmed"),("FR035","Keep factual knowledge age-suitable and distinguish creative story detail from fact","Confirmed"),("FR036","Require Teacher approval of the exact knowledge/video version before playback","Confirmed"),("FR037","Show truthful video states including Generating, Ready for Review, Approved, Playing and Failed","Confirmed baseline states"),("FR038","Wait while video generation is active; do not auto-fallback or skip a required stage","Confirmed"),("FR039","After retry exhaustion, require Teacher to choose retry, skip or end and record disposition","Owner-confirmed; retry budget TBD"),("FR040","Keep sources, script version and approval invalidation traceable across edits","Proposed"),
    ("FR041","Recommend library activities; let Teacher choose or adapt for class/group","Confirmed"),("FR042","Describe activity objectives, age range, materials, steps, duration, safety and observation","Proposed detail"),("FR043","Have Teacher confirm activity feasibility and materials before use","Proposed"),("FR044","Save activity, reflection and off-screen evidence with session outcome","Confirmed"),
    ("FR045","Record age-aware observations and progress over time","Confirmed"),("FR046","Avoid ranking children and compare progress with the child's own history","Confirmed"),("FR047","Keep session, art, contribution, content, activity, reflection and comment history","Confirmed"),("FR048","Export reports by permission and separate direct evidence, Teacher judgment and AI suggestion","Mixed: permission confirmed; taxonomy proposed"),("FR049","Use descriptive criteria for exploration, cooperation, expression, creativity and skills","Proposed"),("FR050","Do not treat unobserved behavior as lack of ability or infer psychology/intelligence from drawings","Repository constraint; report detail proposed"),
    ("FR051","Provide Teacher class dashboard, group status, help queue and controls","Confirmed"),("FR052","Support reviewed presets/automation with immediate Teacher override","Confirmed"),("FR053","Version preset settings for age, topic, tools, AI and session stages","Proposed detail"),("FR054","Prioritize help requests and audit automated action/override","Proposed"),
    ("FR055","Manage adult accounts, roles, school/class scopes and Teacher assignments","Confirmed"),("FR056","Author/review/publish/recall learning content and version AI policy","Confirmed"),("FR057","Monitor service health, authorized reports, consent, audit and incidents","Confirmed"),("FR058","Limit Admin access to child content by purpose and log access","Proposed access detail"),
    ("FR059","Apply consent, retention, export/delete, access control and audit rules","Confirmed"),("FR060","Recover session and artwork from saved drafts/checkpoints with safe feedback","Confirmed"),("FR061","Correct recognition, block unsafe content and never show pending review","Confirmed"),("FR062","Prevent recalled content from use in new sessions","Confirmed"),("FR063","Define offline drawing and queued approval behavior","Proposed"),("FR064","Define backup/provider-copy purge and shared-canvas deletion rules","Proposed; data policy TBD")]
    add_table(doc,["ID","Requirement summary","Authority / refinement"],reqs,[0.62,4.55,1.08],size=8.25,header_fill=NAVY)
    return reqs


def add_report_chapter_one(doc):
    add_heading(doc,"I. Project Introduction",1)
    add_section_intro(doc,"Sketch2Life is a collaborative classroom learning system for children ages 3–12. The product connects individual and shared drawing to Teacher-led knowledge, reviewed AI suggestions, class activities and descriptive reflection. The registration v1.2 is the scope authority; the codebase is evidence of implementation only and does not reduce or prove the full target scope.")
    add_heading(doc,"1. Overview",2)
    add_heading(doc,"1.1 Project Information",3)
    add_table(doc,["Field","Value"],[("Project","Sketch2Life: Collaborative Creative Learning with AI for Children"),("Abbreviation","Sketch2Life"),("Project reference","FA26SE225 (inferred from the registration filename; source form field is blank)"),("Class","TBD; blank in the supplied registration"),("Registration period","September 2026 to March 2027"),("Document snapshot","dev at 0a7ce768c153f2104cff930569c3eda6e484532c plus the FEAT-059 cleanup working-tree change set; not a release tag"),("Review tag","review-2: TBD; not created"),("Repository","https://github.com/CaoDat-2410/Sketch2Life")],[1.55,4.7],size=9.4)
    add_heading(doc,"1.2 Project Team",3)
    add_table(doc,["Member","Code","Role / ownership"],[("Đỗ Cao Đạt","SE193713","Leader; fullstack and cross-surface acceptance"),("Phạm Tấn Hiếu","SE184252","BE2; AI and learning-media services"),("Vũ Chí Bảo","SE182733","FE; Child Android surface"),("Lý Anh Khoa","SE182682","BE1; classroom API and persistence")],[1.8,1.05,3.4],size=9.3)
    add_para(doc,"Supervisor: Lâm Hữu Khánh Phương. The registration’s Class field is empty; no class value is inferred.",size=9.4)
    add_heading(doc,"2. Product Background",2)
    add_para(doc,"The registered problem is that classroom drawing can capture children’s ideas but often leaves collaboration, reliable contribution history, age-appropriate guidance, and follow-up activity outside one coherent session. Sketch2Life links those actions under Teacher control. Drawing similarity is not the learning goal. AI may help explore a child’s idea, but it cannot replace the child’s work or act as an assessment authority.")
    add_heading(doc,"3. Existing Systems and Current Baseline",2)
    add_para(doc,"No named external competitor or classroom system was evidenced in the supplied registration. The report therefore describes the current project baseline rather than making an unsupported comparative market claim. The codebase contains a FastAPI backend, a durable classroom persistence slice, legacy/demo compositions and three separate UI applications. Current UI services rely on fixtures or stubs; branch presence and component tests do not show end-to-end integration.")
    add_table(doc,["Surface / component","Current source evidence","Current limitation"],[("Child","apps/child-ui; Expo SDK 55, React Native 0.83.10","Prototype with local fixtures/stub clients; API integration not established"),("Teacher","apps/teacher-ui; React + Vite","Screens and component tests exist; service calls use fixtures or stubs"),("Super Admin","apps/admin-ui; React + Vite","Prototype services; routes are not a complete authorization boundary"),("Backend","FastAPI application; classroom runtime, demo composition, SQL persistence slices","Durable BE1 slice coexists with demo routes and independent module migration heads"),("AI hosting","Modal selected by owner","Application, model, GPU, endpoint, scaling and latency are TBD; no product deployment claimed")],[1.25,2.45,2.55],size=8.7)
    add_heading(doc,"4. Business Opportunity",2)
    add_para(doc,"The product aims to support one-school classroom pilots for children from 36 through 155 completed months, grouped into experience bands 3–5, 6–8 and 9–12. The target capacity is one class of up to 40 children; it is a planning target, not a measured load result. The opportunity is a teacher-managed flow that joins expressive drawing, safe review, collaborative participation and practical off-screen learning while preserving consent and provenance.")
    add_heading(doc,"5. Software Product Vision",2)
    add_para(doc,"Children explore and explain ideas through drawing and shared classroom activity; Teachers guide each group, control the pace, review AI and media, and record observations; the system preserves accepted work and access history so the classroom can continue safely after interruption.")
    add_heading(doc,"6. Project Scope and Limitations",2)
    add_table(doc,["In scope","Explicit limitation / not established"],[("Child Android app; Teacher Desktop; Super Admin Console","The supplied code does not yet demonstrate an integrated release across these surfaces."),("Managed child profiles, class/group/session management and QR/session-code admission","A child does not have a separate login account. Detailed capability/identity contracts remain proposed."),("Personal/shared drawing, AI suggestion review, gallery, knowledge/video, off-screen activities, reflection and progress","Target behavior exceeds current fixture-backed UI and deployed service evidence."),("Consent, purpose-scoped access, audit, retention and recovery","School process, jurisdiction and data-class-specific retention require further sign-off; session/artwork default is 90 days after session end."),("Parent portal, payments, public social features, ranking and diagnosis","Not in the supplied registration scope; do not infer these features from legacy code.")],[3.1,3.15],size=8.7)


def add_report_chapter_two(doc):
    add_heading(doc,"II. Project Management Plan",1)
    add_section_intro(doc,"The team plans work around four product work packages and keeps that dependency view separate from individual sprint allocation. The 15-week relative plan and preliminary person-day estimates are proposals requiring group and supervisor confirmation; calendar anchors remain TBD within the registered September 2026–March 2027 period.")
    add_heading(doc,"1. Overview",2)
    add_heading(doc,"1.1 Scope and Estimation",3)
    add_table(doc,["Work package","Primary outputs","Effort / status"],[("WP1 Product, classroom learning and evaluation","Approved requirements, age-band interaction, safety review, test protocol","Planned; estimates to be agreed with supervisor"),("WP2 Backend and classroom services","Adult identity, class/profile/consent, sessions/groups, permissions, persistence/recovery","In progress in source slices; not yet verified end to end"),("WP3 AI and learning media","Vision, separate sketch suggestions, review gates, knowledge/video and provenance","Source modules and contracts exist; Modal runtime deployment TBD"),("WP4 Product surfaces and integration","Child Android, Teacher Desktop, Super Admin workflows and cross-surface acceptance","Prototype surfaces exist; fixture/stub integration gap remains")],[1.6,3.6,1.05],size=8.5)
    add_heading(doc,"1.2 Project Objectives",3)
    add_table(doc,["Objective","Target measure","Evidence required"],[("Preserve child agency and Teacher decision gates","100% of test cases block pending/unapproved AI or video from Child playback","Negative authorization tests and version-specific approval cases"),("Support classroom collaboration and recovery","40 synthetic clients for the pilot target; accepted operation is retrievable after restart","Load profile, durable-ack and restart/reconnect results; not yet measured"),("Deliver scoped product surfaces","Child, Teacher and Admin complete assigned workflows against versioned backend contracts","Staging end-to-end run; not yet available"),("Protect child data","Consent/purpose and retention policy checked before protected processing; 90-day default for session/artwork","Policy tests and purge evidence; other data classes TBD")],[2.15,2.5,1.6],size=8.5)
    add_heading(doc,"1.3 Project Risks",3)
    add_table(doc,["Risk","Impact","Control / owner"],[("Fixture-backed UI surfaces do not yet consume a shared versioned contract","Integration defects may appear after screen completion","Publish contract, map API owners, run cross-surface acceptance; Đạt"),("Nine independent Alembic heads lack one verified deployment order","A fresh environment may not reach the same schema as individual feature tests","Agree schema ownership and migration orchestration before staging; Khoa"),("Modal is selected without deployment/model/GPU/latency details","AI availability, cost and response time remain unbounded","Complete provider profile and acceptance test before product claims; Hiếu")],[2.15,2.1,2.0],size=8.4)
    add_heading(doc,"2. Management Approach",2)
    add_para(doc,"Use short, reviewed feature increments with an explicit feature plan, acceptance criteria, approval and evidence folder. Domain rules stay framework-independent; interfaces depend inward; persistence and provider integrations implement application-owned ports. Cross-feature interaction uses versioned contracts. A code change is not accepted as a product claim until the named tests and evidence exist.")
    add_heading(doc,"2.1 Project Process",3)
    add_figure(doc, PICTURES / "D06_activity.png", "Figure 1. End-to-end classroom activity and Teacher review gates.", width=6.25)
    add_para(doc,"The dependency roadmap identifies what must exist before a workflow can be integrated. Parallel team assignments remain an independent planning view; do not assume that one member owns the entire backend, infrastructure, E2E suite or complete Android app by default.",size=9.7)
    add_heading(doc,"2.2 Quality Management",3)
    for item in ["Every feature has measurable acceptance criteria and feature-local evidence.","Run the relevant unit, contract and integration suite in CI; record failures and scope rather than reporting a global pass.","Keep migration and API changes reviewed against source contracts and actual database constraints.","Use synthetic fixtures in development. Do not commit credentials, real child data or provider secrets.","Run the repository security validator before each commit or push; no commit or push is made by this document task."]:
        add_bullet(doc,item,size=9.8)
    add_heading(doc,"2.3 Training Plan",3)
    add_table(doc,["Topic","Participants","Evidence"],[("Child-centered age-adaptive interaction and classroom facilitation","All members; supervisor review","Reviewed screen/workflow acceptance"),("FastAPI, API contracts, migrations and recovery","BE owners and integrator","Contract test and migration review"),("Provider safety, job states, review provenance and Modal deployment","BE2 and integrator","Provider configuration record and no-secret deployment evidence"),("Privacy, school consent, audit and retention","All members","Approved school procedure and policy test cases")],[2.4,1.6,2.25],size=8.6)
    add_heading(doc,"3. Project Deliverables",2)
    add_table(doc,["Deliverable","Status / evidence"],[("Reports 1–5 and Project Tracking consolidated report","This document; editable DOCX, diagrams in repository"),("BM01 Design Brief","Separate document with sections 0–8 and Appendix A"),("Versioned SRS/SDD and editable diagrams","SRS v3.1 is the canonical scope source; Review 2 diagram source is under FEAT-058 artifacts/diagrams"),("Child / Teacher / Admin applications and backend","Prototype sources on local dev baseline; integration and deployment remain open"),("Test evidence","FEAT-059 local cleanup checks; no full system/E2E or pilot report"),("Release tag and deployment profile","review-2 tag TBD; Modal and environment parameters TBD")],[2.65,3.6],size=8.7)
    add_heading(doc,"3.1 Project Schedule and Time Plan",3)
    add_para(doc,"Relative weeks follow the SEP490 sample sequence. Exact calendar start/end dates and official due dates must be confirmed; the standalone Report 2 contains dependencies and output gates. Status labels describe current evidence, not forecast completion.",size=9.1)
    schedule_rows=[(row[0],f"{row[1]}: {row[4]}",row[3],row[5]) for row in SCHEDULE_ROWS]
    add_table(doc,["Relative window","Work / output gate","Lead","Current status"],schedule_rows,[0.85,2.45,1.15,1.8],size=7.6)
    add_heading(doc,"4. Responsibility Assignments",2)
    add_table(doc,["Member","Owned parts (2–3)","Main workflow"],[("Vũ Chí Bảo (SE182733)","Child Android screens; drawing/join interactions; child-surface UI acceptance","WF-02 Draw and recover collaborative canvas"),("Đỗ Cao Đạt (SE193713)","Cross-surface integration; requirements/traceability; acceptance coordination","WF-04 Teacher-led content and learning close"),("Phạm Tấn Hiếu (SE184252)","AI job/provider boundary; sketch review and learning media; safety/result provenance","WF-03 Review AI sketch suggestion"),("Lý Anh Khoa (SE182682)","Classroom HTTP/API; persistence and migrations; session/recovery contracts","WF-01 Prepare and admit classroom session")],[1.5,3.4,1.35],size=8.5)
    add_heading(doc,"5. Project Communications",2)
    add_para(doc,"Use the project repository for versioned requirements, plans, architecture decisions, tests and review evidence. Use the team’s agreed meeting channel for sprint coordination and record decisions in repository ADR/feature records. Escalate scope, consent, child-safety and deployment decisions to the supervisor/owner before implementation. The specific meeting cadence is not evidenced and remains for the team to set.")
    add_heading(doc,"6. Configuration Management",2)
    add_heading(doc,"6.1 Document Management",3)
    add_para(doc,"Keep the official registration and supplied school templates unchanged. Review 2 documents are versioned as editable DOCX; source diagrams remain Mermaid text in the feature package. Record source hashes, output hashes, document version and rendering evidence in FEAT-058. The Review 2 tag has not been created.")
    add_heading(doc,"6.2 Source Code Management",3)
    add_para(doc,"Source repository: https://github.com/CaoDat-2410/Sketch2Life. Document source snapshot: local dev at 0a7ce768c153f2104cff930569c3eda6e484532c plus the uncommitted FEAT-059 cleanup worktree manifest. The source snapshot is not a release. Feature/approval records and automated security checks govern commits; secrets remain in ignored local files or a runtime secret manager.")
    add_heading(doc,"6.3 Tools and Infrastructure",3)
    add_table(doc,["Area","Current / selected","Target or open setting"],[("Child client","Expo SDK 55; React Native 0.83.10","Android confirmed; device matrix TBD"),("Teacher/Admin web","React 19; Vite 8; TypeScript","Desktop browser; support matrix TBD"),("Backend","FastAPI; SQLAlchemy; Alembic","Durable classroom ports; migration order TBD"),("Adult identity","Firebase Authentication only","Storage, Firestore and Realtime Database prohibited"),("AI hosting","Modal selected","App/model/GPU/endpoint/region/scale/timeout/budget TBD"),("DevOps / AI tools","GitHub and CI scripts; Codex assisted cleanup/docs","Tag review-2 and staging/prod/rollback TBD; humans own decisions")],[1.2,2.35,2.8],size=8.15)


def add_report_chapter_three(doc, diagrams):
    add_heading(doc,"III. Software Requirements Specification",1)
    add_section_intro(doc,"This SRS expresses the registered product target and owner-confirmed refinements. It does not certify that the current branch implements each item. CONFIRMED describes scope authority; PROPOSED and TBD identify refinements that need approval. The canonical Vietnamese SRS v3.1 remains the full source of detailed rules.")
    add_heading(doc,"1. Product Overview",2)
    add_para(doc,"The system serves one-school classroom pilots with children aged 36–155 completed months and experience bands 3–5, 6–8 and 9–12. A Teacher prepares a class session, admits children using managed profiles and QR/session code, controls group progress, and approves content. Children draw individually or collaboratively, may request an AI suggestion and may hide or decline it. The session can include gallery presentation, reviewed knowledge/video, an off-screen activity, reflection and descriptive assessment. A child has no separate login account.")
    add_heading(doc,"1.1 Status and assumptions",3)
    add_table(doc,["Status label","Meaning in this report"],[("Confirmed / owner-confirmed","Included product behavior from the registration or direct owner answer"),("Proposed","A design refinement; do not present as frozen contract or implemented behavior"),("TBD","Requires owner/school/provider decision before dependent implementation"),("Source implemented","Code exists; production, durable use and successful provider inference are not implied"),("Fixture/demo","Local, synthetic, mock or process-local behavior only")],[1.8,4.45],size=9)
    add_heading(doc,"2. User Requirements",2)
    add_heading(doc,"2.1 Actors and authorization",3)
    add_table(doc,["Actor","Goals","Authorization boundary"],[("Child participant","Join, draw, request/view/hide/decline reviewed suggestion, share, reflect","Teacher-managed profile; session-bound participant capability is a proposed contract, not a child login"),("Teacher","Manage assigned classes, admit participants, direct pace, review AI/media, record observation","Assigned school/class/session scope; final decision-maker for sessions and review"),("Super Admin","Manage adult accounts, policy, content, consent workflow, audit and operations","Purpose-scoped and audited access; role alone does not imply access to all raw child content"),("School / legal representative","Provide consent and classroom context","School-mediated process; evidence, verifier and purpose record are governed by local policy"),("AI / media provider","Process backend-authorized jobs","No direct mobile credentials; provider output is untrusted until application policy and Teacher review")],[1.2,2.5,2.55],size=8.7)
    add_figure(doc, diagrams["D02"], "Figure 2. Actor/use-case overview. Detailed use cases are indexed by stable UC identifiers in the canonical SRS v3.1.", width=6.25)
    add_heading(doc,"2.2 Use case catalog",3)
    ucs = [
        ("UC01","Prepare a class session and admit children","Teacher, Child","Persisted session; invalid class/code or duplicate admission is rejected."),("UC02","Draw together on a shared canvas","Child, Teacher","Accepted strokes remain in checkpoint; unauthorized region operations do not change peer work."),("UC03","Take turns on one tablet","Child, Teacher","New stroke records selected contributor; old strokes are not relabeled."),("UC04","Let groups progress independently","Child, Teacher","A ready group does not advance the whole class."),("UC05","Request and use a reviewed sketch","Child, Teacher, AI service","Pending, unsafe or stale content is unavailable; child may hide/decline."),("UC06","Correct a misunderstood drawing","Child, Teacher","Child intent/correction remains distinct from AI candidate."),("UC07","Present artwork and play approved knowledge/video","Child, Teacher","Only the exact approved version is played in the allowed scope."),("UC08","Handle slow or failed video","Teacher, Child, media worker","Wait during generation; after exhausted retries Teacher explicitly retries, skips or ends."),("UC09","Run activity and reflection","Child, Teacher","Class/group activity and observation are recorded without assigning peer art to one child."),("UC10","Reconnect after interruption","Child, Teacher, backend","Checkpoint plus accepted operations restore without duplicate operation."),("UC11","Override a preset","Teacher","Policy change is visible and audited; automation cannot bypass review."),("UC12","Handle consent, export or deletion","Teacher / authorized Admin","Purpose-scoped request has traceable disposition; role alone does not expose all records."),("UC13","Propose Phase 2 group assignment","Teacher","Proposal is reviewable and editable; algorithm cannot change class stage or erase history.")]
    add_table(doc,["ID","Use case","Actors","Main result / exception"],ucs,[0.55,1.75,1.2,2.75],size=8.25)
    add_heading(doc,"2.3 Detailed use case descriptions",3)
    detailed = [
        ("UC01 Prepare and admit","Precondition: Teacher is authorized for the class and consent policy allows the session. Main flow: Teacher selects topic/age/tools, assigns groups, opens lobby, issues QR/code, checks incoming roster requests and admits participants. Outcome: active session membership is persisted. Exception: expired code, wrong class or duplicate request is rejected and audited; session does not start without explicit Teacher action."),
        ("UC02 Collaborative drawing","Precondition: admitted participant and canvas permission exist. Main flow: Child selects allowed canvas/tool, submits a versioned operation, backend validates scope/epoch/idempotency, persists accepted operation and returns durable receipt. Outcome: peers can resync from checkpoint and ordered operations. Exception: stale epoch or forbidden region is rejected; local tentative draft is not represented as accepted."),
        ("UC05 Review an AI sketch","Precondition: authorized snapshot and policy permit analysis. Main flow: Child/Teacher requests support, backend captures input version, worker returns typed candidate, Teacher reviews exact sketch, then Child may view, hide or decline. Outcome: source/approval lineage is retained. Exception: unsafe/stale/failure remains unavailable; no automatic display or replacement of child strokes."),
        ("UC08 Resolve failed video","Precondition: video generation is pending for an approved session purpose. Main flow: show truthful progress and wait; after retry budget is exhausted, present explicit Teacher choices. Outcome: Teacher selects retry, skip with disposition or end session. Exception: no auto-fallback, no silent skip and no “video complete” state for a skipped item."),
        ("UC10 Recover a session","Precondition: previously accepted state and checkpoint exist. Main flow: reconnecting client authenticates, fetches current scope/version, loads checkpoint and missing ordered operations, deduplicates receipts and resumes permitted work. Outcome: accepted work and contributor history remain. Exception: revoked or stale capability is denied; local unaccepted draft is retained as a conflict draft, not replayed automatically."),
        ("UC12 Handle data request","Precondition: authorized requester, purpose and scope are established. Main flow: record request, determine data classes/copies under approved retention policy, execute export/delete, record result and verify purge. Outcome: auditable disposition. Exception: unresolved legal hold or policy conflict is escalated; authorization is not inferred from a broad role alone.")]
    for title, body in detailed:
        add_heading(doc,title,4)
        add_para(doc,body,size=9.4)
    add_heading(doc,"3. Functional Requirements",2)
    add_heading(doc,"3.1 System functional overview",3)
    add_figure(doc, diagrams["D04"], "Figure 3. Screen flow separated by Child Android, Teacher Desktop and Super Admin surfaces.", width=6.25)
    add_heading(doc,"3.1.1 Screen inventory",4)
    screens = [
        ("Child Android","Join / session lobby; drawing canvas; optional AI suggestion view; gallery; reviewed knowledge/video; activity instructions; reflection/session close","Participant admission and approved content only; no child login"),
        ("Teacher Desktop","Dashboard/class; profile/enrollment; session setup/lobby; group monitoring; canvas/help; AI review; gallery; content/video review; activity selection; observation/portfolio","Assigned class/school; explicit session control and version-specific review"),
        ("Super Admin Console","Account/role/scope overview; content and policy review; service/job monitoring; audit and data requests","Purpose-scoped, least-privilege and audited"),
    ]
    add_table(doc,["Surface","Screens / tasks","Authorization notes"],screens,[1.25,3.55,1.45],size=8.55)
    add_heading(doc,"3.1.2 Screen authorization",4)
    add_para(doc,"Every command, subscription, query, media playback request and data export is checked at the application boundary for actor, resource, school/class/session scope, purpose and current policy version. Authentication does not by itself grant resource access. Firebase is restricted to adult Authentication in the repository; Firebase Storage, Firestore and Realtime Database are prohibited.",size=9.5)
    add_heading(doc,"3.1.3 Non-screen functions",4)
    for item in ["Adult authentication and scoped capability lifecycle; class/enrollment/consent lookup.","Session and group command validation, durable state transition, checkpoint/reconnect and audit.","Canvas operation validation, idempotency, contribution attribution and provenance.","AI job dispatch, moderation, stale-result rejection, Teacher review gate and child response.","Knowledge/video versioning, truthful progress and explicit failed-video disposition.","Activity safety/readiness check, observation/reflection/portfolio traceability.","Retention, export/delete, audit query, redacted telemetry and incident handling."]:
        add_bullet(doc,item,size=9.2)
    add_heading(doc,"3.1.4 Logical ERD",4)
    add_para(doc,"Logical target entities include OrganizationScope, AdultAccount, Class, StudentProfile, TeacherAssignment, Enrollment, ConsentRecord, ClassroomSession, SessionGroup, ChildParticipant, GroupMembership, DeviceConnection, CanvasDocument, CanvasOperation, CanvasSnapshot, TeacherReview, AI job/proposal, KnowledgeBundle, ReviewedVideo, ActivityDefinition/Assignment, ObservationRecord, ReflectionRecord, PortfolioEntry and AuditEvent. These are specification concepts; most are not all present as migrated classroom tables. Logical links are not declared SQL foreign keys unless the migration explicitly defines one.",size=9.2)
    add_heading(doc,"3.2 Requirement catalog",3)
    add_para(doc,"The following English summaries preserve the stable FR identifiers from canonical SRS v3.1. Refinement status is copied at the requirement level; “confirmed” is not an implementation status.",size=9.4)
    reqs = add_req_catalog(doc)
    add_heading(doc,"3.3 Candidate non-functional requirements",3)
    nfrs = [
        ("NFR01","Age-adaptive interaction","Use age bands 3–5, 6–8 and 9–12 with adult-confirmed age context; verify screen/tool checklist per band."),("NFR02","Authorization and privacy","Reject 100% of unauthorized scope/purpose negative cases; verify no cross-class read or write."),("NFR03","AI safety and review","Zero pending, unsafe or stale AI/media versions served to Child in the negative suite; model-quality thresholds TBD."),("NFR04","Provenance","Every accepted original and derivative in test fixtures retains source/hash/version and review decision."),("NFR05","Recovery","100% of acknowledged operations in the recovery suite are restored once after restart/reconnect; exact RPO/RTO TBD."),("NFR06","Collaboration performance","Candidate p95 durable drawing acknowledgement ≤500 ms and peer update ≤1 s at 40 synthetic children; measure before adoption, no benchmark exists."),("NFR07","Data minimization/retention","Session/artwork default 90 days after session end; other data classes, copies and legal exceptions have separate policies."),("NFR08","Operational telemetry","Operational logs redact child content and credentials; audit records remain separately queryable."),("NFR09","Bounded work","Request size, retry, rate, concurrency and job budget have explicit limits before production; thresholds TBD."),("NFR10","Versioned contracts","Cross-feature commands/events carry contract version; stale result tests reject mismatched context."),("NFR11","Development data","Development checks use synthetic fixtures and repository security validation; no real child data or secrets committed."),("NFR12","Session pacing","Teacher controls shared stage and screen/off-screen transitions; age-specific time thresholds TBD.")]
    add_table(doc,["ID","Attribute","Requirement / verification"],nfrs,[0.6,1.45,4.2],size=8.15)
    add_heading(doc,"4. Requirement Appendix",2)
    add_heading(doc,"4.1 Business rules",3)
    rules = [("BR01","Children are 36–155 completed months inclusive; experience bands are 3–5, 6–8 and 9–12."),("BR02","Teacher retains final control of session stages, AI suggestions, video and activities."),("BR03","Group progress is independent; group readiness does not advance the whole class."),("BR04","Child intent is distinct from AI candidate; a child may hide/decline support."),("BR05","Pending/unsafe/stale content is never served for Child use."),("BR06","After video retry exhaustion, Teacher explicitly chooses retry/skip/end; no auto fallback."),("BR07","No ranking, diagnosis, intelligence/personality inference or treating unobserved behavior as deficiency."),("BR08","School-mediated consent and purpose-scoped access apply; session/artwork default retention is 90 days after session end."),("BR09","Firebase is Authentication-only; mobile clients do not hold AI provider/storage credentials."),("BR10","Preserve original art and record source/version/reviewer provenance for derived content.")]
    add_table(doc,["ID","Rule"],rules,[0.7,5.55],size=8.8)
    add_heading(doc,"4.2 Error and status behavior",3)
    add_para(doc,"Errors distinguish unauthorized, invalid state/version, stale AI output, unsafe content, retryable provider failure, exhausted retry, persistence failure and retained local draft. The caller never receives a success state for a rejected/uncommitted operation. Detailed error codes and HTTP mappings must be frozen in the versioned API contracts.",size=9.4)


def add_report_chapter_four(doc, diagrams):
    add_heading(doc,"IV. Software Design Description",1)
    add_section_intro(doc,"The accepted target composition is a classroom runtime backed by application-owned persistence and provider ports, with legacy Parent/Guide and process-local demos behind separate demo entrypoints. This section documents design direction; it does not claim the runtime was refactored or the Modal service deployed as part of the document work.")
    add_heading(doc,"1. System Design",2)
    add_heading(doc,"1.1 System overview and architecture",3)
    add_figure(doc, diagrams["D01"], "Figure 4. System context and external boundaries. Modal is selected; provider deployment and durable infrastructure settings remain TBD.", width=6.25)
    add_figure(doc, diagrams["D08"], "Figure 5. Current-to-target composition. Dashed connectors are target integrations or TBD adapters.", width=6.25)
    add_para(doc,"Current source exposes a FastAPI app with BE1 classroom composition plus local/demo routers. The FEAT-059 cleanup extracted router and demo composition while preserving route/method order and app state behavior. The product target uses the durable classroom entrypoint as the primary product path and isolates old Parent/Guide/process-local demos. Child, Teacher and Admin applications remain separate packages. Their current fixture/stub services are not an integrated client contract.")
    add_para(doc,"For AI requests, the client calls the application backend only. Application policy validates actor, purpose, content version and request budget before a worker calls a provider adapter. Modal is selected; application name, model, GPU, endpoint, region, scaling, timeout, secret source, cost and measured latency are TBD. Lightning remains a research/test lane until a separately approved migration.",size=9.7)
    add_heading(doc,"1.2 Package diagram",3)
    add_figure(doc, diagrams["D09"], "Figure 6. Package dependencies. Infrastructure implements ports declared inward by the application.", width=6.25)
    add_para(doc,"Backend packages are organized around interface adapters, application use cases and ports, domain policies, infrastructure adapters, and feature-owned contracts. Routers validate/translate transport input but should not own domain decisions or reach into another feature’s database tables.",size=9.5)
    add_heading(doc,"2. Database Design",2)
    add_figure(doc, diagrams["D07"], "Figure 7. Physical ERD excerpt from the BE1 core and AI jobs migration branches. Solid arrows are actual FKs; dashed links are opaque references.", width=6.25)
    add_para(doc,"The BE1 core migration defines durable aggregate revisions, canvas operations, idempotency receipts, checkpoints, asset metadata, audit entries, an outbox and consumer receipts. The AI jobs branch defines durable job, attempt, command/semantic request, completion receipt, cancellation notice and outbox records. Actual foreign keys remain within each module boundary; session_id, organization_id, participant_id, target_id, result_refs and provider references are opaque identifiers unless a migration defines a physical FK.",size=9.4)
    add_para(doc,"At the inspected baseline, Alembic reports nine independent heads: BE1 core, AI jobs, learning content, AI policy, knowledge/video, video review, activity execution, preset automation and video failure. Do not infer one production migration order. A staging migration plan and module ownership decision are required before release.",size=9.4)
    add_table(doc,["Schema / migration head","Purpose shown by source","Relationship note"],[("be1_core / 20261014_0004","Canvas operation receipts/checkpoints, aggregate revisions, audit and outbox","FKs within core migration; session and participant IDs are opaque"),("ai_jobs / 20261015_ai_0001","Durable provider job and attempt lifecycle","FKs within ai_jobs; no FK into BE1 core"),("learning_content / 20261017_lc_0001","Content publication records","Independent head"),("ai_policy / 20261018_ap_0009","Policy, resource/copy governance and resolution records","Independent head"),("knowledge_video / 20261018_kv_0001; video_review / 20261018_vr_0001","Knowledge/video lifecycle and review slices","Independent heads; deployment ordering TBD"),("activity_execution / 20261019_ae_0001; preset_automation / 20261019_pa_0001; video_failure / 20261019_vf_0001","Activity, preset and failed-video disposition slices","Independent heads; do not assume cross-schema FK")],[1.7,2.55,2.0],size=8.15)
    add_heading(doc,"3. Detailed Design",2)
    add_heading(doc,"3.1 Classroom admission and persistent drawing",3)
    add_figure(doc, diagrams["D03"], "Figure 8. Sequence for Teacher admission and durable drawing operation.", width=6.25)
    add_heading(doc,"3.1.1 Class and state model",4)
    add_figure(doc, diagrams["D05"], "Figure 9. Logical classroom model. All target entity relationships are proposed until contracts and migrations are approved.", width=6.25)
    add_para(doc,"A verified security principal and a selected child contributor are separate values. New strokes record the selected contributor for that turn; changing the device selection does not rewrite prior operation attribution. Durable acceptance means a transaction has stored the operation and returned an idempotent receipt; transport receipt or optimistic UI alone is not durable acceptance.",size=9.3)
    add_heading(doc,"3.1.2 Session lifecycle",4)
    add_figure(doc, diagrams["D10"], "Figure 10. Session lifecycle target view; exact wire enum is a contract decision.", width=6.25)
    add_heading(doc,"3.1.3 Participant admission lifecycle",4)
    add_figure(doc, diagrams["D11"], "Figure 11. Participant admission and group assignment. A managed child profile is not a login account.", width=6.25)
    add_heading(doc,"3.2 AI assistance and reviewed media",3)
    add_para(doc,"AI and video work is asynchronous. An AI job records input version, policy snapshot, purpose, retry state, and typed outcome. A result can be accepted, stale or rejected; the job’s successful worker execution does not itself mean product acceptance. Teacher approval is version-specific. Any edit creates a new version and requires a new review before it is shown to a Child.",size=9.5)
    add_heading(doc,"3.2.1 AI job class and state model",4)
    add_figure(doc, diagrams["D12"], "Figure 12. AI job states match the current ai_jobs migration; results retain a separate disposition.", width=6.25)
    add_heading(doc,"3.2.2 Video review and playback",4)
    add_figure(doc, diagrams["D13"], "Figure 13. Video lifecycle with explicit Teacher choice after retry exhaustion.", width=6.25)
    add_heading(doc,"3.2.3 Data boundaries and provenance",4)
    add_table(doc,["Data class","Owner / source of truth","Protection and lineage"],[("Session and group state","Classroom application aggregate","Versioned commands, scoped reads, durable close/recovery"),("Drawing operation and checkpoint","Canvas persistence module","Operation ID, epoch, contributor context, sequence, hash/checkpoint"),("AI/video job and attempt","AI job / media module","Backend-only provider call, input/config version, typed failure and disposition"),("AI suggestion or approved media","Versioned artifact reference plus review record","Source snapshot, policy, reviewer, exact content hash/version"),("Consent and audit","Privacy/application policy records","Purpose, scope, decision, actor and retention; do not log raw child media/secrets")],[1.45,2.05,2.75],size=8.3)


def add_report_chapter_five(doc):
    h = add_heading(doc,"V. Software Testing Documentation",1)
    h.paragraph_format.page_break_before = True
    add_section_intro(doc,"Testing is designed to demonstrate requirement behavior and system boundaries. The current evidence covers selected local UI checks and backend health/session/composition contracts after the FEAT-059 cleanup. It is not an end-to-end acceptance run, deployment test, load benchmark or classroom pilot.")
    add_heading(doc,"1. Scope of Testing",2)
    add_para(doc,"In scope for planned verification: child/teacher/admin screens; session admission and group control; drawing persistence and reconnect; permission boundaries; AI/version/review gates; video generation/failure disposition; consent, retention and audit; accessibility and measurable latency/recovery objectives. Out of scope for current receipt: Modal GPU inference, staging deployment, Android device matrix, real child study, production load, and full cross-surface E2E.",size=9.7)
    add_heading(doc,"2. Test Strategy",2)
    add_heading(doc,"2.1 Testing Types",3)
    add_table(doc,["Type","Purpose","Technique / evidence"],[("Unit","Domain policy, validation and state transitions","Pure input/output cases, boundary and negative cases"),("UI component","Visible state, choices and navigation","React Testing Library / Jest or Vitest with synthetic fixtures"),("Contract","Request/response/event compatibility","Schema validation, route/method snapshot and version compatibility"),("Integration","Persistence, outbox, workers and adapters","Disposable PostgreSQL/queue; no real provider in ordinary CI"),("System / acceptance","Full actor workflow and requirement acceptance","Staging cross-surface scenarios with test accounts and synthetic data"),("Performance / recovery","P95 latency, load, restart and reconnect","Named load profile, 40 synthetic children target, durable receipt and restore assertions")],[1.25,2.0,3.0],size=8.3)
    add_heading(doc,"2.2 Test Levels",3)
    add_table(doc,["Level","Owner","Entry / acceptance gate"],[("Unit / UI","Feature owner","Approved criteria and local fixtures; deterministic run"),("API / contract","Khoa + Hiếu; Đạt coordinates","Versioned contract accepted; auth, errors and migration behavior verified"),("Integrated workflow","Đạt with all surface owners","One staging build connects Child, Teacher, Admin and backend for the named workflow"),("System / acceptance","Whole team and supervisor","BM01 mandatory feature cases pass; no unreviewed AI/media is available to Child"),("Pilot evaluation","Supervisor-approved protocol","School consent and child data governance approved before recruitment")],[1.4,1.6,3.25],size=8.4)
    add_heading(doc,"2.3 Supporting Tools",3)
    add_para(doc,"Current repository receipts use Jest/React Native Testing Library for Child, Vitest/Testing Library for Teacher and Admin, pytest for backend contracts, TypeScript builds, Expo web export, Oxlint and repository architecture/harness/security validators. PostgreSQL, queue and provider-backed environments remain separate acceptance prerequisites. No real Modal request or credentialed provider call is included in the receipt.",size=9.4)
    add_heading(doc,"3. Test Plan",2)
    add_heading(doc,"3.1 Human Resources",3)
    add_table(doc,["Owner","Verification responsibility"],[("Bảo (FE)","Child screen interaction, age-band and drawing view cases"),("Đạt (Fullstack)","Cross-surface contract, end-to-end workflow and acceptance evidence coordination"),("Hiếu (BE2)","AI job lifecycle, review gate, media failure, provider contract"),("Khoa (BE1)","Session, authorization, persistence, migration and recovery contract")],[1.5,4.75],size=8.8)
    add_heading(doc,"3.2 Test Environment",3)
    add_para(doc,"Local feature checks use synthetic fixtures. The CI environment must pin runtime dependencies and service versions. Staging needs a disposable database with a reviewed multi-head migration plan, isolated storage/queue, a provider sandbox or deterministic adapter, adult test accounts, generated QR/session codes and no real child data. Modal endpoint, GPU, model, secret provisioning and scale settings remain TBD.",size=9.5)
    add_heading(doc,"3.3 Test Milestones",3)
    add_table(doc,["Relative window","Gate","Status"],[("W1–W5", "Requirements, scope, design and contracts reviewed", "Drafts exist; approvals and migration order open"),("W6–W11", "Three feature iterations with owner tests and CI", "In progress; cross-surface E2E not proven"),("W12–W13", "Staging E2E, privacy/security, recovery and capacity verification", "Planned; environment not ready"),("W14–W15", "User guides, rollback plan, UAT and final package", "Planned; calendar dates TBD")],[1.15,3.6,1.5],size=8.2)
    add_heading(doc,"4. Test Cases",2)
    test_cases = [
        ("TC-01","Teacher creates session; valid admitted child joins","Session membership persisted; invalid/expired code rejected"),("TC-02","Two clients draw; duplicate request replays idempotency key","One operation is stored; same receipt returned; conflicting payload rejected"),("TC-03","Shared tablet changes active child between strokes","New stroke uses selected contributor; earlier attribution is unchanged"),("TC-04","One group finishes before another","Class stage remains under Teacher control"),("TC-05","AI suggestion pending, stale, unsafe, approved, hidden and declined","Only exact approved version is accessible; child choice is stored"),("TC-06","Video exceeds retry budget","Teacher chooses retry, skip or end; no automatic fallback or false success"),("TC-07","Restart after durable ack, then reconnect","Acknowledged operation and session state restore once"),("TC-08","Wrong-class actor requests art/portfolio/export","Request is denied and audited; no content is returned"),("TC-09","Retention clock reaches 90 days after session end","Session/artwork purge is verified; other data classes use separate policy"),("TC-10","Provider timeout / stale completion","Typed failure returned; stale job result is not served"),("TC-11","Teacher assigns off-screen activity and records group observation","Selected activity/reflection is saved; observation stays descriptive without ranking or false individual attribution")]
    add_table(doc,["ID","Scenario","Pass condition"],test_cases,[0.65,2.6,3.0],size=8.3)
    add_heading(doc,"5. Test Reports",2)
    add_table(doc,["Receipt","Result","Scope and limit"],[("Child UI check: FEAT-059 evidence/raw/ACTIVE_UI_CHECKS.txt","20 tests passed; typecheck, Expo web export and lint command completed","Local Child surface and fixtures; lint retains one warning; no API E2E"),("Teacher UI check: FEAT-059 evidence/raw/ACTIVE_UI_CHECKS.txt","2 tests passed; build, lint and test commands completed","Local React/Vite surface"),("Admin UI check: FEAT-059 evidence/raw/ACTIVE_UI_CHECKS.txt","2 tests passed; build, lint and test commands completed","Local React/Vite surface; not a production authorization assessment"),("Backend core suite: FEAT-059 evidence/raw/BACKEND_CORE_TESTS.txt","9 tests passed","Health, durable session HTTP API and app composition; no full provider or multi-module DB test"),("Post-checks","Architecture, harness, skeleton, security validators and diff check passed","Feature-local receipts; do not establish deployment or full product readiness")],[1.8,1.4,3.05],size=8.1)
    add_para(doc,"The 24 UI tests and 9 backend tests were run for FEAT-059 code cleanup on the current uncommitted working tree. The route/method mapping stayed at 76 paths and 83 operations in the recorded composition check. The checks are historical, bounded receipts for this source snapshot; no tests were run specifically for document authoring. No staging, E2E, live Modal/GPU, capacity, classroom or learning-outcome result is claimed.",size=9.2)


def add_report_tracking(doc):
    add_heading(doc,"VI. Project Tracking",1)
    add_section_intro(doc,"This tracking view reports the document baseline and current integration gaps. It does not replace an approved sprint allocation or convert proposed milestones into commitments.")
    add_heading(doc,"1. Source and Branch Snapshot",2)
    branch_rows = [
        ("local dev","0a7ce768c153f2104cff930569c3eda6e484532c","Review 2 documentation baseline; not pushed"),
        ("origin/FE_UI_Child","ca3779869dfc1385cbbc7b8ca730f780386af7b2","Included in local dev; Child Expo prototype and fixture services"),
        ("origin/FE_UI_Teacher","a5c7fa84c9d05f66f6278954d2fc6b5bb4fc8f62","Included in local dev; Teacher UI prototype"),
        ("origin/FE_UI_Admin","2bb2ab4bb7d1d09b93e566e9b519cd006984ca85","Included in local dev; Admin UI prototype"),
        ("origin/feature/be2-backend-ai-integration","aa41ed6f481a6d4a797609e588137c94a7c778d7","Integrated into local dev; provider deployment not established"),
        ("origin/feature/feat048-srs-diagram-pack","3d8acb30e9a93339b9fc559523b77675421abd40","Audited source branch; not used as the code baseline"),
        ("origin/feature/feat018-p2t5-evaluation-harness","c345822d6b5e29367d7813781970252a8bfcd92b","Separate; not merged or cherry-picked per owner direction"),
    ]
    add_table(doc,["Ref","Source SHA","Disposition"],branch_rows,[2.25,2.4,1.6],size=8.15)
    add_heading(doc,"2. Work Status",2)
    add_table(doc,["Work item","Status","Evidence / next gate"],[("Requested UI/BE2 branch integration","Merged into local dev before cleanup and documentation","Merge records in FEAT-058 notes; no push"),("FEAT-059 structure cleanup","Complete for approved structural scope; changes are uncommitted","24 UI + 9 backend tests and validators recorded; source manifest lists changed files"),("Child/Teacher/Admin backend integration","Open","Replace fixture/stub services with versioned contracts; prove four main workflows end to end"),("Modal product runtime","Selected; configuration TBD","Record model/license, GPU, endpoint/auth, region, scale, timeout, retries, secrets and measurements"),("Schema deploy order","Open","Nine independent Alembic heads; agree module order and clean-database bootstrap"),("review-2 tag and class field","TBD","Tag not created; class is blank in registration; FA26SE225 inferred from filename"),("Staging / pilot / E2E / load testing","Not run","Requires approved deployment, consent/synthetic test data and test protocol")],[1.75,1.45,3.05],size=8.2)
    add_heading(doc,"3. Open Decisions and Gaps",2)
    for item in ["Confirm official class and project code; the class is blank and FA26SE225 is inferred from the filename.","Freeze versioned contracts, cross-surface ownership and the production UI baseline.","Approve migration order for nine independent Alembic heads before staging.","Complete the Modal runtime profile and measurements; settle retention, consent and export/delete controls before pilot."]:
        add_bullet(doc,item,size=9.4)
    add_heading(doc,"4. Release and Change Record",3)
    add_para(doc,"Snapshot: local dev@0a7ce768c153f2104cff930569c3eda6e484532c plus FEAT-059’s uncommitted 72-file manifest (SHA256 87e84a4faedf0572ac0e88b21238b38eb7cd02c9ebc92aea5f4a7931528165ef). No commit, push or review-2 tag is claimed; originals are preserved and Report6 is excluded.",size=9.2)


def build_report(diagrams):
    doc = Document(str(REPORT_TEMPLATE))
    configure_common(doc, margin=1.0, body_size=10.2)
    fill_report_cover(doc)
    doc.core_properties.title = "Sketch2Life Reports 1 to 5 and Project Tracking"
    doc.core_properties.subject = "Review 2 consolidated capstone report"
    doc.core_properties.author = "Sketch2Life Team"
    doc.core_properties.keywords = "Sketch2Life, SRS, SDD, testing, project tracking"
    add_report_frontmatter(doc)
    add_report_chapter_one(doc)
    doc.add_page_break()
    add_report_chapter_two(doc)
    doc.add_page_break()
    add_report_chapter_three(doc, diagrams)
    doc.add_page_break()
    add_report_chapter_four(doc, diagrams)
    add_report_chapter_five(doc)
    add_report_tracking(doc)
    doc.save(str(REPORT_OUT))
    return REPORT_OUT


def bm_title(doc, text, level=2):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5 if level == 2 else 3)
    p.paragraph_format.space_after = Pt(2)
    for r in p.runs:
        set_run(r, size=12.5 if level == 2 else 11, bold=True, color=NAVY)
    return p


def bm_table(doc, headers, rows, widths, size=11, first_col=False):
    t = add_table(doc, headers, rows, widths=widths, size=size, header_fill=NAVY, first_col_fill=NAVY if first_col else None)
    for ri, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    is_header = ri == 0
                    is_first_label = first_col and c._tc is row.cells[0]._tc
                    set_run(r,size=size,bold=r.bold,color=(WHITE if is_header or is_first_label else BLACK))
    return t


def build_brief():
    doc = Document(str(BM_TEMPLATE))
    # Keep the authoritative template as the source; replace body content, not template file.
    body = doc._element.body
    sect = body.sectPr
    for el in list(body):
        if el is not sect:
            body.remove(el)
    configure_common(doc, margin=0.7874, body_size=11)
    for sec in doc.sections:
        sec.page_width = Inches(8.2677)
        sec.page_height = Inches(11.6929)
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2)
        sec.right_margin = Cm(2)
    doc.core_properties.title = "Sketch2Life Review 2 Design Brief"
    doc.core_properties.subject = "BM01 Design Brief"
    doc.core_properties.author = "Sketch2Life Team"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("SWP490  |  CAPSTONE PROJECT  |  REVIEW 2  |  BM01")
    set_run(r,size=9,bold=True,color=TEAL)
    p = doc.add_paragraph()
    p.style = "Title"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Design Brief")
    set_run(r,size=22,bold=True,color=BLACK)
    add_para(doc,"Sketch2Life: Collaborative Creative Learning with AI for Children",size=11,bold=True,color=NAVY,after=3)

    bm_title(doc,"0. Header")
    bm_table(doc,["Field","Value"],[("Project code","FA26SE225 (inferred from registration filename; confirm)"),("Class","TBD; blank in registration"),("Team","Sketch2Life | Đạt SE193713, Hiếu SE184252, Bảo SE182733, Khoa SE182682"),("Supervisor","Lâm Hữu Khánh Phương"),("Repository / tag","https://github.com/CaoDat-2410/Sketch2Life | review-2 tag TBD"),("Version / date","Review 2 draft v1.0 | 23 October 2026 | Ho Chi Minh City")],[1.55,4.7],size=11,first_col=True)

    bm_title(doc,"1. Scope")
    bm_title(doc,"1.1 Registration Changes",3)
    bm_table(doc,["Section","Current registration","Proposed Review 2 wording","Reason"],[("Product direction","Child drawing prototype","Collaborative teacher-led classroom session for ages 36–155 months; personal/shared drawing, reviewed AI help, gallery, knowledge/video, off-screen activity and reflection","Aligns to supplied registration v1.2 and SRS v3.1"),("Actors / surfaces","Earlier app scaffolds","Child Android, Teacher Desktop and Super Admin Console; managed child profiles and QR/session-code admission","Separate prototype surfaces from target service integration")],[0.85,1.65,2.45,1.3],size=11)
    bm_title(doc,"1.2 Out of Scope",3)
    add_para(doc,"Parent portal, payment/credit, public social features, leaderboard or diagnostic inference are not part of the supplied scope. Current prototypes, provider adapters or legacy demo routes do not add those requirements.",size=10.3,after=3)
    bm_title(doc,"2. System at a Glance")
    at_glance = "Sketch2Life supports a Teacher-led classroom where children ages 3–12 join through managed profiles and a QR or session code, draw alone or together, and explore reviewed learning content and an off-screen activity. Teachers control group progress, review every AI sketch and exact video version, and record descriptive observations. Children may hide or decline a reviewed sketch. The product target is one school and one class of up to 40 children; this is not a measured capacity result. Child, Teacher and Admin UIs currently use fixtures or stub services, and Modal is selected without a deployed model or endpoint. The architecture diagram is in the consolidated SDD."
    add_para(doc,at_glance,size=10.5,after=4)

    bm_title(doc,"3. Key Design Decisions")
    decisions = [
        ("D1","Persistent classroom runtime is the product path; isolate legacy/demo entrypoints and integrate by versioned contracts.","Durable classroom path vs process-local demo","Đỗ Cao Đạt | FR007–017 | D08"),
        ("D2","Teacher-managed profile; QR/code join follows roster check. No separate child login.","Child login vs roster admission","Vũ Chí Bảo | FR002–004 | UC01"),
        ("D3","Keep verified principal separate from selected contributor; turn changes never rewrite earlier strokes.","Touch identity vs explicit turn attribution","Lý Anh Khoa | FR003/066 | D05"),
        ("D4","Teacher approves each exact AI sketch; Child can view, hide or decline.","Auto-display vs human review","Phạm Tấn Hiếu | FR026–029 | D12"),
        ("D5","Modal is selected for AI hosting; deployment/model/GPU/endpoint/region/scale remain TBD.","Modal vs research/test lane","Phạm Tấn Hiếu | FR022–030 | D01/D08"),
    ]
    bm_table(doc,["ID","Selected decision and reason","Alternative considered","Owner / trace"],decisions,[0.4,2.7,1.65,1.5],size=11)

    bm_title(doc,"4. Top Risks")
    risks = [
        ("R1","UI clients use fixtures/stubs while API contracts and staging integration remain open; end-to-end workflow may fail at integration.","Integrate one workflow at a time against versioned contracts; require a staging E2E receipt.","Đạt"),
        ("R2","Nine independent Alembic heads have no verified single deployment sequence; fresh environment may diverge.","Agree schema ownership and bootstrap order; test an empty disposable DB before staging.","Khoa"),
        ("R3","Modal is selected but model/GPU/endpoint/region/scale/timeout/cost are TBD; speed and capacity are unknown.","Record deployment profile and run bounded provider acceptance before any product claim.","Hiếu"),
    ]
    risk_rows = [(a, b, "Mitigation: " + c + " Owner: " + d) for a,b,c,d in risks]
    bm_table(doc,["ID","Risk and impact","Control / owner"],risk_rows,[0.4,2.85,3.0],size=11,first_col=True)
    bm_title(doc,"5. Known Gaps")
    gaps = [
        ("G1","Registration code/class fields are blank; FA26SE225 is inferred from filename","Confirm official class/project code","Đạt / supervisor","Before final submission"),
        ("G2","Child, Teacher and Admin prototypes do not prove API integration","Versioned contract + four workflow E2E","All owners; Đạt coordinates","Before staging acceptance"),
        ("G3","Database has 9 independent migration heads","Approve module migration/bootstrap order","Khoa","Before deploy"),
        ("G4","Modal deployment has no model, GPU, endpoint or test receipt","Complete provider profile and measured evaluation","Hiếu","Before AI acceptance"),
    ]
    gap_rows = [(a, b, c + " Owner: " + d + "; gate: " + e) for a,b,c,d,e in gaps]
    bm_table(doc,["ID","Known gap","Action / owner / gate"],gap_rows,[0.45,2.75,3.05],size=11)

    doc.add_page_break()
    bm_title(doc,"6. Proposed Mandatory Features")
    features = [
        ("MF-01","Managed profile and session admission","Teacher, Child","Valid roster/code joins; expired/wrong-class/duplicate request is rejected","TC-01","Planned"),
        ("MF-02","Teacher-controlled session and independent group progress","Teacher, Child","Teacher start/pause/resume/end persists; one ready group cannot advance class","TC-04","Planned"),
        ("MF-03","Collaborative drawing, explicit contributor and recovery","Child, Teacher","Accepted op persists once; shared-device turn does not relabel earlier strokes; restart restores ack","TC-02/03/07","Planned"),
        ("MF-04","Teacher-reviewed AI sketch with child choice","Child, Teacher, AI service","Pending/stale/unsafe sketch is blocked; approved version can be hidden or declined","TC-05","Planned"),
        ("MF-05","Reviewed knowledge/video and explicit failure choice","Teacher, Child, media worker","Exact approved version plays; after retry exhaustion Teacher chooses retry/skip/end","TC-06/10","Planned"),
        ("MF-06","Off-screen activity, reflection and descriptive record","Teacher, Child","Selected class/group activity and Teacher observation save without ranking or false individual attribution","TC-11","Planned"),
    ]
    bm_table(doc,["ID","Feature","Actor","Acceptance criteria (pass/fail)","Test","Status"],features,[0.5,1.4,0.75,2.35,0.8,0.45],size=11)
    add_para(doc,"Status note: these are target acceptance cases; “Planned” means no full product E2E result is recorded. FEAT-059 local cleanup tests are not evidence that every mandatory feature passes.",size=9.4,italic=True,color=GRAY,after=4)
    bm_title(doc,"7. Non-Functional Requirements")
    nfrs = [
        ("NFR-01","Durable response","p95 accepted drawing ack ≤500 ms at 40 synthetic clients","Load test; candidate target, not measured"),
        ("NFR-02","Peer update","p95 accepted update visible to connected peers ≤1 s at 40 synthetic clients","Timed multi-client test; target needs approval"),
        ("NFR-03","Authorization","100% of unauthorized class/purpose negative cases denied; zero content leaked","API/UI boundary negative suite"),
        ("NFR-04","Recovery","100% of acknowledged ops restored once after restart/reconnect; RPO/RTO TBD","Restart/replay suite with unique op IDs"),
        ("NFR-05","Review gate","Zero pending/stale/unsafe/unapproved versions served to Child","Negative version/policy test suite"),
        ("NFR-06","Retention","Session/artwork purged by 90 days after session end; other data classes separate","Clocked retention test and purge audit"),
    ]
    bm_table(doc,["ID","Quality attribute","Metric / threshold","Measurement method"],nfrs,[0.55,1.2,2.7,1.65],size=11)
    add_para(doc,"NFR-01/02 are candidate thresholds for review, not guarantees or observed results. The target pilot of 40 children is an input to a future synthetic load profile.",size=9.2,italic=True,color=GRAY)

    bm_title(doc,"8. Ownership Map")
    owners = [
        ("Vũ Chí Bảo\nSE182733","FE","Child Android screens; join/drawing interaction; UI acceptance","WF-02","D2","FR014–021"),
        ("Đỗ Cao Đạt\nSE193713","Leader / Fullstack","Cross-surface integration; requirements/traceability; acceptance coordination","WF-04","D1","FR031–048"),
        ("Phạm Tấn Hiếu\nSE184252","BE2","AI job/provider boundary; sketch review; knowledge/video safety","WF-03","D4/D5","FR022–030, FR034–040"),
        ("Lý Anh Khoa\nSE182682","BE1","Classroom API; persistence/migrations; session/recovery contract","WF-01","D3","FR001–017"),
    ]
    bm_table(doc,["Member","Role","Owned parts (2–3)","Main workflow","Key decisions","Related requirements"],owners,[0.95,0.7,2.2,0.75,0.75,0.9],size=11)
    add_para(doc,"Each person owns one main workflow and 2–3 named parts. The team will review cross-layer dependencies using versioned contracts; the role labels do not imply that one person owns the whole backend, infrastructure or complete mobile application.",size=10.1)

    doc.add_page_break()
    bm_title(doc,"Appendix A. Defense 1.1 Targets",1)
    bm_title(doc,"A1. Main Workflows",3)
    workflows = [
        ("WF-01","Prepare and admit class session","Teacher + Child + Backend","Teacher configures class/topic/groups; Child joins; server admits and stores membership","Session: Lobby → Active; membership admitted","Khoa"),
        ("WF-02","Draw and recover shared canvas","Child + Teacher + Backend","Children draw from multiple devices or take turns; Teacher pauses/restores; server stores checkpoint","Canvas revision increments; checkpoint references accepted op","Bảo"),
        ("WF-03","Request and review AI sketch","Child + Teacher + AI service","Child requests help; worker prepares suggestion; Teacher reviews; Child views/hides/declines","Job queued → reviewed proposal; child response recorded","Hiếu"),
        ("WF-04","Review knowledge/video and close learning session","Teacher + Child + media service","Teacher selects/reviews exact version; waits or handles failure; class completes activity/reflection","Video approved/played or explicit skip/end; session outcome saved","Đạt"),
    ]
    bm_table(doc,["ID","Workflow","Actors","Trigger / business outcome","Entity state change","Presenter"],workflows,[0.45,1.15,1.2,1.75,1.15,0.55],size=11)
    add_para(doc,"Every workflow includes at least two actors or three mandatory features and changes a main entity state. The scenarios below are executable test plans; they have not been run against staging.",size=9.4)
    bm_title(doc,"A2. Demo Scenarios",3)
    demos = [
        ("DS-01","WF-01 main","Teacher creates a session for a synthetic class, shows QR/code, admits the correct managed profile and starts the session.","Admitted participant and active session are visible to both surfaces.","Khoa"),
        ("DS-02","WF-01 exception","Child submits expired code or wrong-class request.","Request is denied; no membership is created; denial is auditable.","Khoa"),
        ("DS-03","WF-02 main","Two synthetic clients draw; one shared tablet changes selected contributor; Teacher pauses and resumes.","Unique accepted operations survive reconnect; old contributor labels remain unchanged.","Bảo"),
        ("DS-04","WF-02 exception","Client retries same op after lost acknowledgement, then sends conflicting payload with same idempotency key.","Original receipt replays once; conflicting payload is rejected; canvas is not duplicated.","Bảo"),
        ("DS-05","WF-03 main","Child asks for a sketch; Teacher approves exact version; Child hides it and continues drawing.","Review and child choice are recorded; original strokes remain unchanged.","Hiếu"),
        ("DS-06","WF-03 exception","Worker returns stale or unsafe result, or Teacher has not reviewed it.","Result never appears in Child view; typed failure/review state remains visible to Teacher.","Hiếu"),
        ("DS-07","WF-04 main","Teacher approves a reviewed video, plays it, runs a class activity and records an observation.","Exact approved media plays; activity and observation appear in session outcome.","Đạt"),
        ("DS-08","WF-04 exception","Video retries exhaust; Teacher chooses skip; class ends early.","Skip disposition is recorded; video is not marked played/successful; saved work is retained.","Đạt"),
    ]
    bm_table(doc,["ID","Workflow / path","Scenario","Expected result","Presenter"],demos,[0.5,0.9,2.4,1.95,0.5],size=11)
    bm_title(doc,"A3. Functional Test Case Target",3)
    arithmetic = [("Mandatory features",6,3,18),("Main workflows",4,1,4),("Main entity state transitions",4,1,4),("Agreed NFRs",4,1,4),("Minimum committed total", "", "", 30),("Automated in CI minimum", "50% of total", "", 15)]
    bm_table(doc,["Source","Count","Minimum per item","Subtotal / target"],arithmetic,[2.0,1.1,1.55,1.7],size=11)
    add_para(doc,"Minimum target: 30 committed functional test cases; at least 15 must run automatically in CI. The four initial NFR cases are NFR-03 through NFR-06. The two latency thresholds NFR-01 and NFR-02 remain candidate targets pending owner approval.",size=11)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal":
            for run in paragraph.runs:
                run.font.name = "Calibri"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
                run.font.size = Pt(11)
    doc.save(str(BM_OUT))
    return BM_OUT


def main():
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    diagrams = make_diagrams()
    report = build_report(diagrams)
    brief = build_brief()
    print(f"report={report}")
    print(f"brief={brief}")
    print(f"diagram_sources={len(list(DIAGRAM_SOURCES.glob('*.mmd')))}")


if __name__ == "__main__":
    main()
