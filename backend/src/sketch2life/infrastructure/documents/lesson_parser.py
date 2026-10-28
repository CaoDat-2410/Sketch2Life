"""Bounded, deterministic text extraction for PDF, DOCX and PPTX lesson files."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterable
from typing import BinaryIO

from sketch2life.application.ports.lesson_documents import (
    LessonDocumentExtraction,
    LessonDocumentRejected,
)

MAX_ARCHIVE_ENTRIES = 2048
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_EXTRACTED_TEXT_BYTES = 2 * 1024 * 1024
PARSER_NAME = "sketch2life-lesson-document-parser"
PARSER_VERSION = "1.0.0"

_MEDIA = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


class DeterministicLessonDocumentParser:
    def extract(self, file_obj: BinaryIO, *, file_extension: str) -> LessonDocumentExtraction:
        extension = file_extension.lower()
        media_type = _MEDIA.get(extension)
        if media_type is None:
            raise LessonDocumentRejected("UNSUPPORTED_MEDIA_TYPE")
        file_obj.seek(0)
        signature = file_obj.read(8)
        file_obj.seek(0)
        if extension == ".pdf" and not signature.startswith(b"%PDF-"):
            raise LessonDocumentRejected("INVALID_SOURCE_DOCUMENT")
        if extension in {".docx", ".pptx"}:
            _validate_office_archive(file_obj, extension)

        try:
            if extension == ".pdf":
                pages, coverage, text_parts = _extract_pdf(file_obj)
            elif extension == ".docx":
                pages, coverage, text_parts = _extract_docx(file_obj)
            else:
                pages, coverage, text_parts = _extract_pptx(file_obj)
        except LessonDocumentRejected:
            raise
        except Exception:
            raise LessonDocumentRejected("INVALID_SOURCE_DOCUMENT") from None

        text = "\n".join(part for part in text_parts if part)
        if len(text.encode("utf-8")) > MAX_EXTRACTED_TEXT_BYTES:
            raise LessonDocumentRejected("DOCUMENT_TEXT_LIMIT_EXCEEDED")
        fields, warnings = _extract_allowlisted_fields(text)
        if not text.strip():
            raise LessonDocumentRejected("OCR_NOT_SUPPORTED")
        return LessonDocumentExtraction(
            detected_media_type=media_type,
            parser_name=PARSER_NAME,
            parser_version=PARSER_VERSION,
            page_or_slide_count=pages,
            page_coverage=tuple(coverage),
            extracted_fields=fields,
            warnings=tuple(warnings),
        )


def _validate_office_archive(file_obj: BinaryIO, extension: str) -> None:
    file_obj.seek(0)
    try:
        with zipfile.ZipFile(file_obj) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > MAX_ARCHIVE_ENTRIES:
                raise LessonDocumentRejected("INVALID_SOURCE_DOCUMENT")
            expanded = 0
            names = {item.filename for item in entries}
            for item in entries:
                if item.flag_bits & 0x1:
                    raise LessonDocumentRejected("ENCRYPTED_DOCUMENT_UNSUPPORTED")
                if item.filename.startswith(("/", "\\")) or ".." in item.filename.split("/"):
                    raise LessonDocumentRejected("INVALID_SOURCE_DOCUMENT")
                expanded += item.file_size
                if expanded > MAX_UNCOMPRESSED_BYTES:
                    raise LessonDocumentRejected("DOCUMENT_EXPANSION_LIMIT_EXCEEDED")
                ratio = item.file_size / max(item.compress_size, 1)
                if ratio > MAX_COMPRESSION_RATIO:
                    raise LessonDocumentRejected("DOCUMENT_EXPANSION_LIMIT_EXCEEDED")
            expected = (
                {"[Content_Types].xml", "word/document.xml"}
                if extension == ".docx"
                else {"[Content_Types].xml", "ppt/presentation.xml"}
            )
            if not expected.issubset(names):
                raise LessonDocumentRejected("INVALID_SOURCE_DOCUMENT")
    except LessonDocumentRejected:
        raise
    except (zipfile.BadZipFile, OSError):
        raise LessonDocumentRejected("INVALID_SOURCE_DOCUMENT") from None
    finally:
        file_obj.seek(0)


def _extract_pdf(file_obj: BinaryIO) -> tuple[int, list[int], list[str]]:
    from pypdf import PdfReader

    reader = PdfReader(file_obj, strict=True)
    if reader.is_encrypted:
        raise LessonDocumentRejected("ENCRYPTED_DOCUMENT_UNSUPPORTED")
    if not 1 <= len(reader.pages) <= 100:
        raise LessonDocumentRejected("DOCUMENT_PAGE_LIMIT_EXCEEDED")
    parts: list[str] = []
    coverage: list[int] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            coverage.append(index)
            parts.append(text)
    return len(reader.pages), coverage, parts


def _extract_docx(file_obj: BinaryIO) -> tuple[None, list[int], list[str]]:
    from docx import Document

    document = Document(file_obj)
    values: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(value for value in values if value.strip())
    return None, [1] if text.strip() else [], [text]


def _extract_pptx(file_obj: BinaryIO) -> tuple[int, list[int], list[str]]:
    from pptx import Presentation

    presentation = Presentation(file_obj)
    if not 1 <= len(presentation.slides) <= 100:
        raise LessonDocumentRejected("DOCUMENT_PAGE_LIMIT_EXCEEDED")
    parts: list[str] = []
    coverage: list[int] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
        text = "\n".join(line for line in lines if line.strip())
        if text.strip():
            coverage.append(index)
            parts.append(text)
    return len(presentation.slides), coverage, parts


_FIELD_ALIASES = {
    "title": {"title", "tên bài", "tiêu đề", "tên giáo án"},
    "topic": {"topic", "chủ đề", "đề tài"},
    "drawing_subject": {"drawing subject", "chủ thể vẽ", "vẽ", "đối tượng vẽ"},
    "drawing_steps_text": {"drawing steps", "các bước vẽ", "hướng dẫn vẽ"},
    "subject_learning_notes": {"subject notes", "kiến thức chủ đề", "ghi chú kiến thức"},
    "objectives": {"objectives", "mục tiêu", "mục tiêu học tập"},
    "script_direction": {"script direction", "định hướng kịch bản"},
}


def _extract_allowlisted_fields(text: str) -> tuple[dict[str, object], list[str]]:
    fields: dict[str, object] = {}
    current: str | None = None
    collected: dict[str, list[str]] = {}
    for line in text.splitlines():
        match = re.match(r"^\s*([^:：]{1,80})\s*[:：]\s*(.*)$", line)
        if match:
            label, value = match.groups()
            normalized = label.casefold().strip()
            current = next(
                (key for key, aliases in _FIELD_ALIASES.items() if normalized in aliases),
                None,
            )
            if current is not None and value.strip():
                collected.setdefault(current, []).append(value.strip())
            continue
        if current is not None and line.strip():
            collected.setdefault(current, []).append(line.strip())
    for key, values in collected.items():
        joined = "\n".join(values).strip()
        if key == "objectives":
            objectives = [
                re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", value).strip()
                for value in re.split(r"[\n;]+", joined)
            ]
            fields[key] = [value for value in objectives if value][:12]
        else:
            fields[key] = joined[:5000]
    warnings = [f"MISSING_FIELD_{name.upper()}" for name in _FIELD_ALIASES if name not in fields]
    if "activities" not in fields:
        warnings.append("ACTIVITIES_REQUIRE_TEACHER_REVIEW")
    return fields, warnings


__all__ = ["DeterministicLessonDocumentParser"]
